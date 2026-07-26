from pathlib import Path

from openpyxl import Workbook

from app.importer import read_tcga_cdr_rows, tcga_cdr_endpoint_rows


def test_tcga_cdr_tsv_parser_builds_endpoint_rows(tmp_path: Path) -> None:
    path = tmp_path / "cdr.tsv"
    path.write_text(
        "\t".join(["bcr_patient_barcode", "type", "OS", "OS.time", "PFI", "PFI.time"]) + "\n"
        + "\t".join(["TCGA-AA-0001", "KIRC", "1", "100", "0", "80"]) + "\n"
        + "\t".join(["TCGA-AA-0002", "TCGA-BRCA", "0", "200", "", ""]) + "\n",
        encoding="utf-8",
    )

    rows = read_tcga_cdr_rows(path)
    endpoints = tcga_cdr_endpoint_rows(rows)

    assert {(item.patient_id, item.cohort, item.endpoint, item.event, item.time_days) for item in endpoints} == {
        ("TCGA-AA-0001", "TCGA-KIRC", "OS", 1, 100.0),
        ("TCGA-AA-0001", "TCGA-KIRC", "PFI", 0, 80.0),
        ("TCGA-AA-0002", "TCGA-BRCA", "OS", 0, 200.0),
    }


def test_tcga_cdr_xlsx_parser_uses_first_sheet(tmp_path: Path) -> None:
    path = tmp_path / "cdr.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["bcr_patient_barcode", "type", "DSS", "DSS.time"])
    sheet.append(["TCGA-AA-0003", "ACC", 1, 50])
    workbook.save(path)

    rows = read_tcga_cdr_rows(path)
    endpoints = tcga_cdr_endpoint_rows(rows)

    assert len(endpoints) == 1
    assert endpoints[0].patient_id == "TCGA-AA-0003"
    assert endpoints[0].cohort == "TCGA-ACC"
    assert endpoints[0].endpoint == "DSS"
    assert endpoints[0].event == 1
    assert endpoints[0].time_days == 50.0


def test_tcga_cdr_imports_explicit_competing_risk_status(tmp_path: Path) -> None:
    path = tmp_path / "cdr.xlsx"
    workbook = Workbook()
    base = workbook.active
    base.title = "TCGA-CDR"
    base.append(
        [
            "bcr_patient_barcode",
            "type",
            "OS",
            "OS.time",
            "DSS",
            "DSS.time",
            "DFI",
            "DFI.time",
        ]
    )
    base.append(["TCGA-AA-0004", "LIHC", 1, 400, 0, 400, 0, 400])
    extra = workbook.create_sheet("ExtraEndpoints")
    extra.append(
        [
            "bcr_patient_barcode",
            "type",
            "DSS_cr",
            "DSS.time.cr",
            "DFI.cr",
            "DFI.time.cr",
        ]
    )
    extra.append(["TCGA-AA-0004", "LIHC", 2, 400, 2, 400])
    workbook.save(path)

    endpoints = tcga_cdr_endpoint_rows(
        read_tcga_cdr_rows(path),
        competing_rows=read_tcga_cdr_rows(
            path,
            sheet_name="ExtraEndpoints",
        ),
    )
    by_endpoint = {endpoint.endpoint: endpoint for endpoint in endpoints}

    for endpoint in ("DSS", "DFI"):
        metadata = by_endpoint[endpoint].raw_metadata
        assert metadata["competing_risk_status"] == 2
        assert metadata["competing_event"] == 1
        assert metadata["source_competing_sheet"] == "ExtraEndpoints"
        assert metadata["competing_risk_coding"]["1"] == "event_of_interest"


def test_tcga_cdr_rejects_misaligned_competing_time(tmp_path: Path) -> None:
    path = tmp_path / "cdr.xlsx"
    workbook = Workbook()
    base = workbook.active
    base.title = "TCGA-CDR"
    base.append(["bcr_patient_barcode", "type", "DSS", "DSS.time"])
    base.append(["TCGA-AA-0005", "LIHC", 0, 400])
    extra = workbook.create_sheet("ExtraEndpoints")
    extra.append(
        ["bcr_patient_barcode", "type", "DSS_cr", "DSS.time.cr"]
    )
    extra.append(["TCGA-AA-0005", "LIHC", 2, 401])
    workbook.save(path)

    endpoints = tcga_cdr_endpoint_rows(
        read_tcga_cdr_rows(path),
        competing_rows=read_tcga_cdr_rows(
            path,
            sheet_name="ExtraEndpoints",
        ),
    )

    assert "competing_risk_status" not in endpoints[0].raw_metadata
