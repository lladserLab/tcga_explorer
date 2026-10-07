import csv
from datetime import datetime, timedelta, timezone
import gzip
import hashlib
import io
import json
import math
from pathlib import Path
import re
import tarfile
from types import SimpleNamespace
import zipfile
import xml.etree.ElementTree as ET

import pytest
from pydantic import ValidationError
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.main import latest_data_manifest
from app.models import (
    CancerType,
    DataManifest,
    DataSource,
    RepositoryDataset,
    RepositoryEndpointDefinition,
    RepositoryExpressionLayer,
    RepositoryRelease,
)
from app.pancancer_study_universes import (
    SourceKind,
    StudyUniverseCategory,
)
from app.multiverse import expand_multiverse_request, summarize_multiverse
from app.repository import importer
from app.repository.service import list_repository_datasets
from app.repository.adapters.cbioportal import (
    RETRYABLE_HTTP_CODES,
    apply_expression_transform,
    build_recipe_metadata,
    download as download_repository_source,
    download_with_http_ranges,
    duplicated_source_feature_symbols,
    expression_feature_id,
    fetch_provided_ensembl_hugo_mapping,
    fetch_provided_hugo_entrez_mapping,
    fetch_provided_hugo_optional_ensembl_mapping,
    merge_gdc_survival_into_patients,
    parse_endpoint,
    read_expression_header,
    resolve_source_file_spec,
    sample_passes_eligibility,
    sample_selection_rank,
    source_feature_ids_match,
)
from app.repository.adapters.cgga import (
    extract_exact_zip_member,
    materialize_cgga_expression,
    normalized_html_text,
    select_cgga_clinical_rows,
    validate_cgga_access_evidence,
    validate_selected_cohort,
    verify_pinned_archive,
)
from app.repository.adapters.geo import (
    attach_geo_curated_clinical_records,
    derive_date_interval_days,
    derive_geo_curated_clinical_fields,
    geo_sample_metadata_values,
    geo_sample_metadata_identifier,
    geo_sample_identifier,
    geo_source_filename,
    map_geo_event,
    materialize_ensembl_expression,
    materialize_entrez_expression,
    materialize_geo_additional_soft_sources,
    materialize_geo_curated_clinical_evidence,
    materialize_geo_count_matrix,
    materialize_geo_download,
    materialize_geo_expression_files,
    materialize_geo_featurecounts_tar,
    materialize_geo_normalized_files_tar,
    materialize_geo_recount3_matrix,
    materialize_geo_transcript_count_files_tar,
    materialize_geo_transcript_count_matrix,
    materialize_geo_wide_normalized_matrix,
    materialize_geo_xlsx_count_matrix,
    normalize_geo_identifier,
    normalized_geo_age,
    normalized_geo_categorical_value_map,
    normalized_event_value_map,
    parse_ncbi_gene_info_map,
    parse_geo_date,
    parse_geo_family_soft,
    read_geo_delimited_clinical_evidence,
    sample_passes_characteristic_filters,
    select_geo_patient_samples,
)
from app.repository.adapters.gdc import (
    gdc_case_passes_eligibility,
    gdc_star_tpm_iterator,
    materialize_gdc_star_tpm_matrix,
    select_gdc_expression_files,
)
from app.repository.adapters.figshare import (
    ensure_patient_fields_are_consistent,
    materialize_figshare_counts_matrix,
)
from app.repository.adapters.mendeley import (
    materialize_mendeley_counts_matrix,
    select_validation_rows,
    validate_mendeley_file_pin,
)
from app.repository.adapters.dryad import (
    materialize_dryad_counts_matrix,
    parse_anubis_challenge,
    read_dryad_clinical,
    solve_anubis_pow,
    validate_dryad_file_pin,
)
from app.repository.adapters.biostudies import (
    materialize_featurecounts_archives,
    materialize_processed_fpkm_matrix,
    parse_biostudies_sdrf,
    parse_uromol_sdrf,
    parse_uvm_clinical_text,
)
from app.repository.adapters import nottingham as nottingham_adapter
from app.repository.adapters.nottingham import (
    ICART1_2_LAYOUT,
    assert_consistent_clinical_core,
    canonicalize_ena_report,
    materialize_nottingham_expression,
    nottingham_quant_passes_qc,
    parse_ena_report,
    parse_gtf_exon_interval,
    parse_mashmap_intervals,
    parse_nottingham_sdrf,
    rank_participant_runs,
    read_fasta_sequences,
    salmon_quant_options,
    write_fasta_sequences,
)
from app.repository.adapters.icgc import (
    materialize_gene_reference,
    parse_object_listing,
    resolve_icgc_os,
    resolve_icgc_recurrence_interval,
    select_expression_analysis,
    selected_expression_values,
)
from app.repository.adapters import metaprism as metaprism_adapter
from app.repository.adapters.metaprism import (
    materialize_metaprism_expression,
    select_metaprism_samples,
)
from app.repository.adapters.pmc import (
    detect_creative_commons_license,
    materialize_geo_rpkm_matrix,
    materialize_xlsb_expression_matrix,
    materialize_xlsx_expression_matrix,
    normalize_curated_records,
    parse_pmc_pow_challenge,
    parse_gencode_gene_map,
    publication_expression_sample_ids,
    read_delimited_records,
    read_xlsb_sheet,
    read_xlsx_sheet,
    selected_sample_indexes,
    solve_pmc_pow_nonce,
    verify_pinned_supplement,
)
from app.repository.adapters.pdc import (
    materialize_pdc_clinical,
    materialize_pdc_expression,
)
from app.repository.adapters.zenodo import (
    materialize_direct_clinical_file,
    materialize_direct_patient_matrix,
    materialize_patient_level_counts,
    select_deepest_library_per_patient,
    validate_crossref_license,
    validate_zenodo_file_pin,
)
from app.repository.discovery import (
    _rna_seq_sample_count,
    discover_cbioportal_candidates,
)
from app.repository.storage import (
    read_float32le_row,
    safe_bundle_path,
    sha256_file,
    write_float32le_matrix,
)
from app.schemas import (
    AnalysisRequest,
    MultiverseAnalysisOut,
    MultiverseAnalysisRequest,
    SignatureGene,
)
from app.survival import (
    ClinicalOutcome,
    filter_sample_candidates,
    select_expression_complete_samples,
)


class EmptyFilters:
    sample_types: list[str] = []
    stages: list[str] = []
    grades: list[str] = []
    genders: list[str] = []
    races: list[str] = []
    age_min = None
    age_max = None
    max_time_days = None


def test_geo_age_divisor_converts_days_without_relabelling_them_as_years():
    assert float(normalized_geo_age("3652.5", divisor=365.25)) == 10.0
    assert normalized_geo_age("42", divisor=1.0) == "42"


def test_nottingham_ena_report_is_canonicalized_and_parsed(tmp_path):
    header = [
        "run_accession",
        "study_accession",
        "sample_accession",
        "experiment_accession",
        "library_layout",
        "read_count",
        "base_count",
        "fastq_ftp",
        "fastq_md5",
        "fastq_bytes",
        "submitted_ftp",
        "submitted_md5",
        "submitted_bytes",
    ]
    rows = [
        [
            "ERR2",
            "PRJEB1",
            "SAMEA2",
            "ERX2",
            "PAIRED",
            "20",
            "3000",
            "ftp.sra.ebi.ac.uk/ERR2_1.fastq.gz;"
            "ftp.sra.ebi.ac.uk/ERR2_2.fastq.gz",
            f"{'b' * 32};{'c' * 32}",
            "101;102",
            "",
            "",
            "",
        ],
        [
            "ERR1",
            "PRJEB1",
            "SAMEA1",
            "ERX1",
            "PAIRED",
            "10",
            "1000",
            "ftp.sra.ebi.ac.uk/ERR1_1.fastq.gz;"
            "ftp.sra.ebi.ac.uk/ERR1_2.fastq.gz",
            f"{'d' * 32};{'e' * 32}",
            "91;92",
            "ftp.sra.ebi.ac.uk/S1_R1.fastq.gz;"
            "ftp.sra.ebi.ac.uk/S1_R2.fastq.gz",
            f"{'a' * 32};{'f' * 32}",
            "81;82",
        ],
    ]
    raw = tmp_path / "raw.tsv"
    with raw.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)
    canonical = tmp_path / "canonical.tsv"

    canonicalize_ena_report(raw, canonical)
    runs = parse_ena_report(
        canonical,
        expected_study_accession="PRJEB1",
    )

    assert list(runs) == ["ERR1", "ERR2"]
    assert runs["ERR2"]["base_count"] == 3000
    assert runs["ERR1"]["fastq_files"][0]["url"].startswith("https://")
    assert runs["ERR1"]["fastq_source"] == "submitted"
    assert canonical.read_text(encoding="utf-8").splitlines()[1].startswith(
        "ERR1\t"
    )


def test_nottingham_icart1_2_sdrf_links_endpoints_to_ena(tmp_path):
    header = [
        "Source Name",
        "Comment[ENA_SAMPLE]",
        "Comment[BioSD_SAMPLE]",
        "Characteristics[organism]",
        "Comment[id assigned at time of rna-sequencing]",
        "Characteristics[disease]",
        "Characteristics[age]",
        "Characteristics[sex]",
        "Characteristics[organism status]",
        "Characteristics[survival time]",
        "Characteristics[distant metastasis status]",
        "Characteristics[distant metastasis time]",
        "Characteristics[disease staging]",
        "Comment[ENA_RUN]",
        "Comment[FASTQ_URI]",
        "Comment[SUBMITTED_FILE_NAME]",
    ]
    rows = []
    ena_runs = {}
    for index in range(1, 11):
        run = f"ERR{index}"
        sample = f"SAMEA{index}"
        status = (
            "died from breast cancer" if index in {1, 2} else "alive"
        )
        metastasis = "yes" if index in {1, 3, 4} else "no"
        fastq_files = []
        for mate in (1, 2):
            name = f"{run}_{mate}.fastq.gz"
            fastq_files.append(
                {
                    "url": f"https://ftp.sra.ebi.ac.uk/{name}",
                    "name": name,
                    "size": 100 + mate,
                    "md5": str(mate) * 32,
                }
            )
            rows.append(
                [
                    f"Sample {index}",
                    f"ERS{index}",
                    sample,
                    "Homo sapiens",
                    f"R38_{index}",
                    "triple-negative breast cancer",
                    str(30 + index),
                    "female",
                    status,
                    str(20 + index),
                    metastasis,
                    str(10 + index),
                    "2",
                    run,
                    f"ftp://ftp.sra.ebi.ac.uk/{name}",
                    name,
                ]
            )
        ena_runs[run] = {
            "run_accession": run,
            "study_accession": "PRJEB1",
            "sample_accession": sample,
            "experiment_accession": f"ERX{index}",
                "read_count": 1000 + index,
                "base_count": 10_000 + index,
                "fastq_files": fastq_files,
                "reported_fastq_name_sets": [
                    sorted(item["name"] for item in fastq_files)
                ],
            }
    sdrf = tmp_path / "study.sdrf.tsv"
    with sdrf.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)

    participants = parse_nottingham_sdrf(
        sdrf,
        layout=ICART1_2_LAYOUT,
        ena_runs=ena_runs,
    )

    assert len(participants) == 10
    assert sum(
        row["clinical"]["bcss_event"] for row in participants.values()
    ) == 2
    assert sum(
        row["clinical"]["dmfs_event"] for row in participants.values()
    ) == 3
    assert participants["Sample 1"]["runs"][0]["biosd_sample"] == "SAMEA1"


def test_nottingham_run_ranking_and_qc_are_prespecified():
    ranked = rank_participant_runs(
        [
            {"run_accession": "ERR3", "base_count": 200},
            {"run_accession": "ERR2", "base_count": 300},
            {"run_accession": "ERR1", "base_count": 300},
        ]
    )
    assert [row["run_accession"] for row in ranked] == [
        "ERR1",
        "ERR2",
        "ERR3",
    ]
    assert nottingham_quant_passes_qc(
        processed_fragments=2_000_000,
        mapped_fragments=1_500_000,
        mapping_rate=75.0,
    ) == (True, "")
    passed, reason = nottingham_quant_passes_qc(
        processed_fragments=2_000_000,
        mapped_fragments=900_000,
        mapping_rate=9.0,
    )
    assert passed is False
    assert "mapped fragments" in reason
    assert "mapping rate" in reason


def test_nottingham_cpp_salmon_options_are_supported_and_pinned():
    assert salmon_quant_options(threads=8, extra_options=[]) == [
        "-l",
        "A",
        "--seqBias",
        "--gcBias",
        "--posBias",
        "-p",
        "8",
    ]


def test_nottingham_partial_decoy_coordinates_are_deterministic(tmp_path):
    assert parse_gtf_exon_interval(
        "chr1\tsource\texon\t101\t250\t.\t+\t.\tgene_id \"G1\";\n",
        line_number=1,
    ) == ("chr1", 100, 250)
    assert (
        parse_gtf_exon_interval(
            "chr1\tsource\tgene\t101\t250\t.\t+\t.\tgene_id \"G1\";\n",
            line_number=2,
        )
        is None
    )

    mashmap_output = tmp_path / "mashmap.out"
    mashmap_output.write_text(
        "tx2 1000 0 500 + chr2 2000 200 700 95\n"
        "tx1 1000 0 500 + chr1 2000 300 800 90\n"
        "tx1 1000 0 500 + chr1 2000 100 600 91\n",
        encoding="utf-8",
    )
    assert parse_mashmap_intervals(
        mashmap_output,
        allowed_contigs={"chr1", "chr2"},
    ) == [
        ("chr1", 100, 600),
        ("chr1", 300, 800),
        ("chr2", 200, 700),
    ]


def test_nottingham_partial_decoy_fasta_round_trip(tmp_path):
    path = tmp_path / "decoys.fa"
    records = [
        ("chr1", "ACGT" * 30),
        ("chr2", "TGCA" * 3),
    ]
    write_fasta_sequences(path, records)

    assert list(read_fasta_sequences(path)) == records
    assert max(
        len(line)
        for line in path.read_text(encoding="ascii").splitlines()
        if not line.startswith(">")
    ) <= 80


def test_nottingham_gentrome_headers_are_the_tx2gene_authority(tmp_path):
    gentrome = io.BytesIO(
        b">ENST1.2 cdna gene:ENSG1.5 gene_symbol:GeneA description:x\n"
        b"ACGT\n"
        b">ENST2.1 cdna gene:ENSG2.3 gene_symbol:GeneB description:y\n"
        b"TGCA\n"
        b">chr1\n"
        b"NNNN\n"
    )
    output = tmp_path / "tx2gene.tsv"

    rows = nottingham_adapter.build_tx2gene_from_gentrome(
        gentrome,
        output,
        minimum_rows=2,
    )

    assert rows == 2
    assert output.read_text(encoding="utf-8").splitlines() == [
        "ENST1.2\tENSG1.5|GENEA",
        "ENST2.1\tENSG2.3|GENEB",
    ]


def test_nottingham_gtf_versions_are_preserved_in_tx2gene(tmp_path):
    annotation = tmp_path / "annotation.gtf"
    annotation.write_text(
        '1\ts\ttranscript\t1\t2\t.\t+\t.\tgene_id "ENSG1"; '
        'gene_version "5"; transcript_id "ENST1"; '
        'transcript_version "2"; gene_name "GeneA";\n',
        encoding="utf-8",
    )
    output = tmp_path / "tx2gene.tsv"

    rows = nottingham_adapter.build_tx2gene(
        annotation,
        output,
        minimum_rows=1,
    )

    assert rows == 1
    assert output.read_text(encoding="utf-8") == (
        "ENST1.2\tENSG1.5|GENEA\n"
    )


def test_nottingham_prebuilt_index_extraction_is_selective(tmp_path):
    required = {
        "complete_ref_lens.bin",
        "ctable.bin",
        "ctg_offsets.bin",
        "info.json",
        "mphf.bin",
        "pos.bin",
        "rank.bin",
        "refAccumLengths.bin",
        "reflengths.bin",
        "refseq.bin",
        "seq.bin",
        "versionInfo.json",
    }
    archive_path = tmp_path / "reference.tgz"
    with tarfile.open(archive_path, "w:gz") as archive:
        for name, payload in [
            (
                "default/gentrome.fa",
                b">ENST1.1 gene:ENSG1.1 gene_symbol:G1\nACGT\n",
            ),
            ("default/reference.gtf", b"must not be extracted\n"),
            ("default/gentrome.fa.unused", b"retained\n"),
            *[
                (f"default/{name}", b"{}\n")
                for name in sorted(required)
            ],
        ]:
            info = tarfile.TarInfo(name)
            info.size = len(payload)
            archive.addfile(info, io.BytesIO(payload))
    index_dir = tmp_path / "index"
    index_dir.mkdir()
    tx2gene = tmp_path / "tx2gene.tsv"

    files, rows = nottingham_adapter.extract_prebuilt_salmon_index(
        archive_path,
        index_dir,
        tx2gene_path=tx2gene,
        archive_root="default",
        gentrome_name="gentrome.fa",
        minimum_tx2gene_rows=1,
    )

    assert rows == 1
    assert required <= set(files)
    assert not (index_dir / "reference.gtf").exists()
    assert tx2gene.read_text(encoding="utf-8") == (
        "ENST1.1\tENSG1.1|G1\n"
    )


def test_nottingham_duplicate_blocks_require_consistent_endpoints():
    first = {
        "age": 52.0,
        "bcss_months": 80.0,
        "bcss_event": 0,
    }
    assert_consistent_clinical_core(
        first,
        dict(first),
        participant_id="1",
    )
    with pytest.raises(ValueError, match="inconsistent clinical"):
        assert_consistent_clinical_core(
            first,
            {**first, "bcss_months": 81.0},
            participant_id="1",
        )


def test_nottingham_expression_uses_selected_gene_tpm(tmp_path):
    selected = []
    for sample_index in range(2):
        quant = tmp_path / f"quant-{sample_index}.genes.sf"
        with quant.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(
                handle, delimiter="\t", lineterminator="\n"
            )
            writer.writerow(
                ["Name", "Length", "EffectiveLength", "TPM", "NumReads"]
            )
            for gene_index in range(10_001):
                writer.writerow(
                    [
                        f"ENSG{gene_index:011d}.1|GENE{gene_index}",
                        "1000",
                        "900",
                        str(gene_index + sample_index),
                        "10",
                    ]
                )
        selected.append(
            {
                "sample_id": f"SAMPLE-{sample_index}",
                "quant": {"quant_genes_path": quant},
            }
        )
    output = tmp_path / "expression.tsv"

    summary = materialize_nottingham_expression(selected, output)

    assert summary["source_expression_rows"] == 10_001
    assert output.read_text(encoding="utf-8").splitlines()[0] == (
        "GENE_ID\tSAMPLE-0\tSAMPLE-1"
    )


def test_nottingham_quant_cache_reuses_identical_fastq_content(tmp_path):
    fastq_identity = [
        {"md5": "a" * 32, "size": 101},
        {"md5": "b" * 32, "size": 202},
    ]
    fastq_input_signature = hashlib.sha256(
        json.dumps(
            fastq_identity,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    quant_signature = "c" * 64
    quant_dir = (
        tmp_path
        / "quant"
        / "by-input"
        / f"{fastq_input_signature}-{quant_signature}"
    )
    quant_genes = quant_dir / "salmon" / "quant.genes.sf"
    quant_genes.parent.mkdir(parents=True)
    quant_genes.write_text("Name\tTPM\nENSG1|GENE1\t1\n", encoding="utf-8")
    quant_sha256 = hashlib.sha256(quant_genes.read_bytes()).hexdigest()
    state = {
        "schema_version": "tcga-trace-nottingham-run-state-v1",
        "run_accession": "ERR-FIRST",
        "source_accession": "E-MTAB-FIRST",
        "quant_signature": quant_signature,
        "fastq_input_signature": fastq_input_signature,
        "salmon_version": "1.12.1",
        "salmon_binary_sha256": "d" * 64,
        "fastq_files": [],
        "processed_fragments": 2_000_000,
        "mapped_fragments": 1_500_000,
        "mapping_rate": 75.0,
        "qc_passed": True,
        "qc_reason": "",
        "quant_genes_relative_path": "salmon/quant.genes.sf",
        "quant_genes_sha256": quant_sha256,
    }
    (quant_dir / "run_state.json").write_text(
        json.dumps(state, sort_keys=True, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    second_run = {
        "run_accession": "ERR-SECOND",
        "fastq_files": [
            {
                **fastq_identity[0],
                "name": "second_R1.fastq.gz",
                "url": "https://example.test/second_R1.fastq.gz",
            },
            {
                **fastq_identity[1],
                "name": "second_R2.fastq.gz",
                "url": "https://example.test/second_R2.fastq.gz",
            },
        ],
    }

    result = nottingham_adapter.quantify_nottingham_run(
        second_run,
        cache_root=tmp_path,
        accession="E-MTAB-SECOND",
        salmon_binary="/not/used/salmon",
        salmon_version="1.12.1",
        salmon_binary_sha256="d" * 64,
        index_dir=tmp_path / "unused-index",
        tx2gene_path=tmp_path / "unused-tx2gene.tsv",
        quant_options=[],
        quant_signature=quant_signature,
        min_mapped_fragments=1_000_000,
        min_mapping_rate=10.0,
    )

    assert result["quant_reused"] is True
    assert result["run_accession"] == "ERR-FIRST"
    assert result["source_accession"] == "E-MTAB-FIRST"
    assert result["quant_genes_path"] == quant_genes


def test_geo_source_filename_supports_official_query_downloads():
    url = (
        "https://www.ncbi.nlm.nih.gov/geo/download/"
        "?acc=GSE1&format=file&file=counts.csv.gz"
    )
    assert geo_source_filename(
        url,
        configured_name="GSE1_counts.csv.gz",
        label="expression",
    ) == "GSE1_counts.csv.gz"
    assert geo_source_filename(
        "https://ftp.ncbi.nlm.nih.gov/GSE1_family.soft.gz",
        label="family SOFT",
    ) == "GSE1_family.soft.gz"
    with pytest.raises(ValueError, match="safe source filename"):
        geo_source_filename(
            url,
            configured_name="../counts.csv.gz",
            label="expression",
        )


def test_geo_individual_expression_files_are_pinned_and_deterministic(
    tmp_path,
    monkeypatch,
):
    payloads = {
        "https://example.test/GSM2.tsv.gz": gzip.compress(
            b"gene\tvalue\nB\t2\n",
            mtime=0,
        ),
        "https://example.test/GSM1.tsv.gz": gzip.compress(
            b"gene\tvalue\nA\t1\n",
            mtime=0,
        ),
    }
    file_specs = [
        {
            "name": url.rsplit("/", 1)[-1],
            "url": url,
            "size": len(payload),
            "sha256": hashlib.sha256(payload).hexdigest(),
        }
        for url, payload in reversed(list(payloads.items()))
    ]

    def fake_materialize(
        url,
        path,
        *,
        expected_size,
        cache_name=None,
    ):
        assert expected_size == len(payloads[url])
        assert cache_name == path.name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payloads[url])

    monkeypatch.setattr(
        "app.repository.adapters.geo.materialize_geo_download",
        fake_materialize,
    )
    first_archive, first_manifest, first_summary = (
        materialize_geo_expression_files(
            file_specs,
            tmp_path / "first",
        )
    )
    second_archive, second_manifest, second_summary = (
        materialize_geo_expression_files(
            list(reversed(file_specs)),
            tmp_path / "second",
        )
    )

    assert sha256_file(first_archive) == sha256_file(second_archive)
    assert first_manifest.read_bytes() == second_manifest.read_bytes()
    assert first_summary == second_summary
    assert first_summary["geo_source_expression_file_count"] == 2
    with tarfile.open(first_archive) as archive:
        members = archive.getmembers()
        assert [member.name for member in members] == [
            "GSM1.tsv.gz",
            "GSM2.tsv.gz",
        ]
        assert all(member.mtime == 0 for member in members)
        assert all(member.uid == member.gid == 0 for member in members)
        assert all(member.mode == 0o644 for member in members)


@pytest.mark.parametrize(
    "file_specs",
    [
        [],
        [
            {
                "name": "../unsafe.tsv.gz",
                "url": "https://example.test/unsafe.tsv.gz",
                "size": 1,
                "sha256": "a" * 64,
            }
        ],
        [
            {
                "name": "same.tsv.gz",
                "url": "https://example.test/one.tsv.gz",
                "size": 1,
                "sha256": "a" * 64,
            },
            {
                "name": "same.tsv.gz",
                "url": "https://example.test/two.tsv.gz",
                "size": 1,
                "sha256": "b" * 64,
            },
        ],
        [
            {
                "name": "bad.tsv.gz",
                "url": "https://example.test/bad.tsv.gz",
                "size": 0,
                "sha256": "not-a-sha256",
            }
        ],
    ],
)
def test_geo_individual_expression_files_reject_invalid_specs(
    tmp_path,
    file_specs,
):
    with pytest.raises(ValueError):
        materialize_geo_expression_files(file_specs, tmp_path)


def test_geo_download_reuses_a_local_cache(
    tmp_path,
    monkeypatch,
):
    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()
    cached = cache_dir / "GSE1_RAW.tar"
    cached.write_bytes(b"pinned-source")
    target = tmp_path / "source" / "GSE1_RAW.tar"
    monkeypatch.setenv("GEO_DOWNLOAD_CACHE_DIR", str(cache_dir))

    materialize_geo_download(
        "https://example.test/GSE1_RAW.tar",
        target,
        expected_size=len(b"pinned-source"),
    )

    assert target.read_bytes() == b"pinned-source"


def test_repository_download_retries_a_truncated_response(
    tmp_path,
    monkeypatch,
):
    class Response(io.BytesIO):
        def __init__(self, payload):
            super().__init__(payload)
            self.headers = {"Content-Length": "4"}

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            self.close()

    responses = iter([Response(b""), Response(b"data")])
    requests = []

    def urlopen(request, **_kwargs):
        requests.append(request)
        return next(responses)

    monkeypatch.setattr(
        "app.repository.adapters.cbioportal.urllib.request.urlopen",
        urlopen,
    )
    monkeypatch.setattr(
        "app.repository.adapters.cbioportal.time.sleep",
        lambda *_args: None,
    )
    target = tmp_path / "source.bin"
    download_repository_source("https://example.test/source.bin", target)
    assert target.read_bytes() == b"data"
    assert len(requests) == 2
    assert all(
        request.get_header("Accept-encoding") == "identity"
        for request in requests
    )


def test_repository_range_download_validates_each_interval(
    tmp_path,
    monkeypatch,
):
    source = b"0123456789abcdef"
    requests = []

    class Response(io.BytesIO):
        def __init__(self, payload, start, end):
            super().__init__(payload)
            self.status = 206
            self.headers = {
                "Content-Length": str(len(payload)),
                "Content-Range": (
                    f"bytes {start}-{end}/{len(source)}"
                ),
            }

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            self.close()

    def urlopen(request, **_kwargs):
        requests.append(request)
        value = request.get_header("Range")
        match = re.fullmatch(r"bytes=(\d+)-(\d+)", value)
        assert match is not None
        start, end = map(int, match.groups())
        return Response(source[start : end + 1], start, end)

    monkeypatch.setattr(
        "app.repository.adapters.cbioportal.urllib.request.urlopen",
        urlopen,
    )
    target = tmp_path / "source.bin"
    download_with_http_ranges(
        "https://example.test/source.bin",
        target,
        total_size=len(source),
        chunk_size=5,
    )

    assert target.read_bytes() == source
    assert len(requests) == 4
    assert all(
        request.get_header("Accept-encoding") == "identity"
        for request in requests
    )


def test_cbioportal_request_timeout_is_retryable():
    assert 408 in RETRYABLE_HTTP_CODES


def test_cgga_pinned_archive_requires_exact_size_and_sha256(tmp_path):
    archive = tmp_path / "clinical.zip"
    with zipfile.ZipFile(archive, "w") as handle:
        handle.writestr("clinical.tsv", "CGGA_ID\tOS\rP1\t100\r")
    observed_sha256 = sha256_file(archive)

    verify_pinned_archive(
        archive,
        expected_size=archive.stat().st_size,
        expected_sha256=observed_sha256,
        label="fixture",
    )
    with pytest.raises(ValueError, match="size differs"):
        verify_pinned_archive(
            archive,
            expected_size=archive.stat().st_size + 1,
            expected_sha256=observed_sha256,
            label="fixture",
        )
    with pytest.raises(ValueError, match="SHA-256 differs"):
        verify_pinned_archive(
            archive,
            expected_size=archive.stat().st_size,
            expected_sha256="0" * 64,
            label="fixture",
        )

    extracted = tmp_path / "clinical.tsv"
    extract_exact_zip_member(archive, "clinical.tsv", extracted)
    assert extracted.read_bytes() == b"CGGA_ID\tOS\rP1\t100\r"


def test_cgga_cohort_filter_is_disease_and_outcome_independent():
    rows = [
        {
            "CGGA_ID": "LGG_OK",
            "PRS_type": "Primary",
            "Histology": "A",
            "Grade": "WHO II",
            "Age": "18",
            "OS": "100",
            "Censor (alive=0; dead=1)": "0",
        },
        {
            "CGGA_ID": "GBM_OK",
            "PRS_type": "Primary",
            "Histology": "GBM",
            "Grade": "WHO IV",
            "Age": "40",
            "OS": "200",
            "Censor (alive=0; dead=1)": "1",
        },
        {
            "CGGA_ID": "RECURRENT",
            "PRS_type": "Recurrent",
            "Histology": "rA",
            "Grade": "WHO II",
            "Age": "35",
            "OS": "300",
            "Censor (alive=0; dead=1)": "1",
        },
        {
            "CGGA_ID": "PEDIATRIC",
            "PRS_type": "Primary",
            "Histology": "A",
            "Grade": "WHO III",
            "Age": "17",
            "OS": "400",
            "Censor (alive=0; dead=1)": "1",
        },
        {
            "CGGA_ID": "MISSING_OS",
            "PRS_type": "Primary",
            "Histology": "A",
            "Grade": "WHO III",
            "Age": "44",
            "OS": "NA",
            "Censor (alive=0; dead=1)": "0",
        },
    ]
    lgg = select_cgga_clinical_rows(
        rows,
        {
            "prs_type": "Primary",
            "grades": ["WHO II", "WHO III"],
            "minimum_age": 18,
        },
    )
    gbm = select_cgga_clinical_rows(
        rows,
        {
            "prs_type": "Primary",
            "grades": ["WHO IV"],
            "histologies": ["GBM"],
            "minimum_age": 18,
        },
    )
    assert [row["CGGA_ID"] for row in lgg] == ["LGG_OK"]
    assert [row["CGGA_ID"] for row in gbm] == ["GBM_OK"]
    validate_selected_cohort(
        lgg,
        expected_patients=1,
        expected_events=0,
        expected_censored=1,
    )
    validate_selected_cohort(
        gbm,
        expected_patients=1,
        expected_events=1,
        expected_censored=0,
    )


def test_cgga_expression_selects_columns_and_reports_case_duplicates(
    tmp_path,
):
    source = tmp_path / "expression.tsv"
    source.write_bytes(
        (
            "Gene_Name\tP1\tP2\tP3\r"
            "C15ORF37\t1\t2\t3\r"
            "C15orf37\t4\t5\t6\r"
            "TP53\t7\t8\t9\r"
        ).encode("utf-8")
    )
    output = tmp_path / "selected.tsv"

    summary = materialize_cgga_expression(
        source,
        output,
        ["P3", "P1"],
        expected_gene_rows=3,
    )

    with output.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle, delimiter="\t"))
    assert rows == [
        ["Hugo_Symbol", "P3", "P1"],
        ["C15ORF37", "3", "1"],
        ["C15orf37", "6", "4"],
        ["TP53", "9", "7"],
    ]
    assert summary == {
        "gene_rows": 3,
        "selected_samples": 2,
        "exact_unique_symbols": 3,
        "case_insensitive_unique_symbols": 2,
        "case_insensitive_ambiguous_symbols": 1,
    }


def test_cgga_access_evidence_ignores_dynamic_page_content():
    home = normalized_html_text(
        "<div>Counter 1042</div><p>Open access to read counts data for "
        "<a>mRNAseq_693</a> and <a>mRNAseq_325</a> datasets</p>"
    )
    download_page = normalized_html_text(
        "<div id='mRNAseq_693'>DataSet ID: mRNAseq_693 "
        "If you use this part of the data, please consider to cite: "
        "<a href='CGGA.mRNAseq_693_clinical.20200506.txt.zip'>"
        "Clinical Data</a>"
        "<a href='CGGA.mRNAseq_693.RSEM-genes.20200506.txt.zip'>"
        "Expression Data from STAR+RSEM (FPKM value)</a></div>"
    )
    validate_cgga_access_evidence(
        home,
        download_page,
        source_dataset_id="mRNAseq_693",
        clinical_member="CGGA.mRNAseq_693_clinical.20200506.txt",
        expression_member="CGGA.mRNAseq_693.RSEM-genes.20200506.txt",
    )
    with pytest.raises(ValueError, match="open-access statement"):
        validate_cgga_access_evidence(
            "mRNAseq_693 mRNAseq_325",
            download_page,
            source_dataset_id="mRNAseq_693",
            clinical_member="CGGA.mRNAseq_693_clinical.20200506.txt",
            expression_member="CGGA.mRNAseq_693.RSEM-genes.20200506.txt",
        )


class ExternalSample:
    def __init__(self, patient_id: str, sample_id: str, rank: int):
        self.patient_id = patient_id
        self.barcode = sample_id
        self.selection_rank = rank
        self.sample_type = "Tumor"
        self.stage = None
        self.grade = None
        self.gender = None
        self.race = None
        self.age_at_index = None
        self.os_time_days = None
        self.os_event = None


def write_tsv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


def test_pdc_expression_uses_pinned_gencode_mapping(tmp_path):
    expression = tmp_path / "expression.tsv"
    mapping = tmp_path / "mapping.tsv.gz"
    output = tmp_path / "mapped.tsv"
    sample_ids = [f"P{index:02d}" for index in range(1, 11)]
    gene_ids = [
        f"ENSG{index:011d}.1" for index in range(1, 10_003)
    ]

    with gzip.open(mapping, "wt", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["transcript", "gene", "gene_name"],
            delimiter="\t",
            lineterminator="\n",
        )
        writer.writeheader()
        for index, gene_id in enumerate(gene_ids):
            writer.writerow(
                {
                    "transcript": f"ENST{index:011d}.1",
                    "gene": gene_id,
                    "gene_name": (
                        "DUPLICATE"
                        if index < 2
                        else f"GENE{index:05d}"
                    ),
                }
            )
    with expression.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(
            handle, delimiter="\t", lineterminator="\n"
        )
        writer.writerow(["idx", *sample_ids])
        for index, gene_id in enumerate(gene_ids):
            writer.writerow(
                [gene_id, *[f"{index + sample / 10:.1f}" for sample in range(10)]]
            )

    summary = materialize_pdc_expression(
        expression, mapping, output
    )

    assert summary["sample_count"] == 10
    assert summary["gene_count"] == 10_002
    assert summary["unique_gene_symbols"] == 10_001
    assert summary["ambiguous_gene_symbols"] == 1
    assert summary["unambiguous_gene_symbols"] == 10_000
    rows = csv.reader(output.open(), delimiter="\t")
    assert next(rows) == ["gene_id", *sample_ids]
    assert next(rows)[0] == f"{gene_ids[0]}|DUPLICATE"


def test_pdc_clinical_requires_explicit_os_and_expression_linkage(
    tmp_path,
):
    clinical = tmp_path / "clinical.tsv"
    patients = tmp_path / "patients.tsv"
    samples = tmp_path / "samples.tsv"
    sample_ids = [f"P{index:02d}" for index in range(1, 11)]
    with clinical.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "idx",
                "Age",
                "Sex",
                "Histologic_Grade",
                "Stage",
                "OS_days",
                "OS_event",
                "PFS_days",
                "PFS_event",
            ],
            delimiter="\t",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerow(
            {
                "idx": "data_type",
                "Age": "CON",
                "Sex": "BIN",
                "OS_days": "CON",
                "OS_event": "BIN",
                "PFS_days": "CON",
                "PFS_event": "BIN",
            }
        )
        for index, sample_id in enumerate(sample_ids):
            writer.writerow(
                {
                    "idx": sample_id,
                    "Age": 50 + index,
                    "Sex": "Female",
                    "Histologic_Grade": "G3",
                    "Stage": "Stage III",
                    "OS_days": 100 + index,
                    "OS_event": index % 2,
                    "PFS_days": 80 + index,
                    "PFS_event": (index + 1) % 2,
                }
            )
        writer.writerow(
            {
                "idx": "CLINICAL_ONLY",
                "Age": 70,
                "Sex": "Female",
                "Histologic_Grade": "G3",
                "Stage": "Stage IV",
                "OS_days": "NA",
                "OS_event": "NA",
                "PFS_days": "NA",
                "PFS_event": "NA",
            }
        )

    summary = materialize_pdc_clinical(
        clinical, sample_ids, patients, samples
    )

    assert summary == {
        "clinical_patients": 11,
        "expression_clinical_patients": 10,
        "os_complete_patients": 10,
        "os_events": 5,
        "os_censored": 5,
        "pfs_complete_patients": 10,
        "pfs_events": 5,
        "pfs_censored": 5,
    }
    patient_rows = list(
        csv.DictReader(patients.open(), delimiter="\t")
    )
    sample_rows = list(
        csv.DictReader(samples.open(), delimiter="\t")
    )
    assert len(patient_rows) == 11
    assert len(sample_rows) == 10
    assert sample_rows[0] == {
        "SAMPLE_ID": "P01",
        "PATIENT_ID": "P01",
        "SAMPLE_TYPE": "Tumor",
    }
    assert patient_rows[0]["PFS_DAYS"] == "80"
    assert patient_rows[0]["PFS_STATUS"] == "1"


def test_figshare_counts_are_normalized_to_log2_cpm(tmp_path):
    source = tmp_path / "counts.csv"
    source.write_text(
        "Name,Description,S1,S2\n"
        "GENEA,ENSG1,10,30\n"
        "GENEB,ENSG2,30,10\n",
        encoding="utf-8",
    )
    target = tmp_path / "selected.tsv"

    summary = materialize_figshare_counts_matrix(
        source,
        target,
        ["S1", "S2"],
        delimiter=",",
        feature_column="Name",
        sample_start_column=3,
        minimum_gene_rows=2,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    expected_high = math.log2(750_000 + 1)
    expected_low = math.log2(250_000 + 1)
    assert rows[0] == ["Name", "S1", "S2"]
    assert float(rows[1][1]) == pytest.approx(expected_low)
    assert float(rows[1][2]) == pytest.approx(expected_high)
    assert float(rows[2][1]) == pytest.approx(expected_high)
    assert float(rows[2][2]) == pytest.approx(expected_low)
    assert summary["expression_normalization"] == "log2(CPM + 1)"


def test_figshare_patient_fields_must_match_across_samples():
    with pytest.raises(ValueError, match="conflicting"):
        ensure_patient_fields_are_consistent(
            [
                {"patient": "P1", "sample": "S1", "status": "0"},
                {"patient": "P1", "sample": "S2", "status": "1"},
            ],
            patient_id_column="patient",
            fields=["status"],
        )


def test_mendeley_selects_primary_ccrcc_with_exact_recurrence_time():
    base = {
        "Batch ID": "1A",
        "Primary Histology": "ccRCC",
        "Metastasectomy or [met] Biopsy": "No",
        "Date of resection or CR": "01/01/20",
        "Date of recurrence": "",
        "Date of last contact": "02/01/20",
    }
    selected, summary = select_validation_rows(
        [
            {
                **base,
                "Record ID": "VA001",
                "rna seq": "KR-3539-1",
                "Recurrence": "Yes",
                "Date of recurrence": "01/11/20",
            },
            {
                **base,
                "Record ID": "VA002",
                "rna seq": "KR-3539-2",
                "Recurrence": "No",
            },
            {
                **base,
                "Record ID": "VA003",
                "rna seq": "KR-3539-3",
                "Recurrence": "No",
                "Primary Histology": "papillary",
            },
            {
                **base,
                "Record ID": "VA004",
                "rna seq": "KR-3539-4",
                "Recurrence": "No",
            },
        ],
        {"1", "2", "3", "4"},
        excluded_sample_keys={"4"},
    )

    assert [
        (row["patient_id"], row["time_days"], row["event"])
        for row in selected
    ] == [
        ("VA001", 10, 1),
        ("VA002", 31, 0),
    ]
    assert summary["eligible_patients"] == 2
    assert summary["events"] == 1
    assert summary["censored"] == 1
    assert summary["exclusions"] == {
        "histology_mismatch": 1,
        "source_qc_outlier": 1,
    }


def test_mendeley_counts_use_all_rows_for_cpm_and_drop_duplicates(
    tmp_path,
):
    source = tmp_path / "counts.gct"
    source.write_text(
        "Name\tDescription\t3539-KR-1-A\t3539-KR-2-B\n"
        "ENSG1\tGENEA\t10\t30\n"
        "ENSG2\tGENEB\t30\t10\n"
        "ENSG3\tDUP\t10\t10\n"
        "ENSG4\tdup\t20\t20\n",
        encoding="utf-8",
    )
    target = tmp_path / "selected.tsv"

    summary = materialize_mendeley_counts_matrix(
        source,
        target,
        [("1", "VA001"), ("2", "VA002")],
        minimum_gene_rows=2,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    assert rows[0] == ["Hugo_Symbol", "VA001", "VA002"]
    assert [row[0] for row in rows[1:]] == ["GENEA", "GENEB"]
    assert float(rows[1][1]) == pytest.approx(
        math.log2((10 / 70) * 1_000_000 + 1)
    )
    assert float(rows[1][2]) == pytest.approx(
        math.log2((30 / 70) * 1_000_000 + 1)
    )
    assert summary["source_duplicated_symbol_labels"] == 1
    assert summary["source_duplicated_symbol_rows"] == 2
    assert summary["normalized_expression_gene_rows"] == 2


def test_mendeley_file_pin_requires_exact_sha256_and_size():
    pin = validate_mendeley_file_pin(
        {
            "file_id": "file-1",
            "name": "clinical.csv",
            "size": 12,
            "sha256": "a" * 64,
        },
        {
            "id": "file-1",
            "filename": "clinical.csv",
            "content_details": {
                "size": 12,
                "sha256_hash": "a" * 64,
                "download_url": "https://example.org/clinical.csv",
            },
        },
    )
    assert pin["file_id"] == "file-1"
    assert pin["sha256"] == "a" * 64

    with pytest.raises(ValueError, match="SHA-256"):
        validate_mendeley_file_pin(
            {
                "file_id": "file-1",
                "name": "clinical.csv",
                "size": 12,
                "sha256": "b" * 64,
            },
            {
                "id": "file-1",
                "filename": "clinical.csv",
                "content_details": {
                    "size": 12,
                    "sha256_hash": "a" * 64,
                    "download_url": "https://example.org/clinical.csv",
                },
            },
        )


def test_dryad_anubis_challenge_is_parsed_and_solved():
    html = (
        '<script id="anubis_challenge" type="application/json">'
        '{"rules":{"algorithm":"fast","difficulty":2},'
        '"challenge":{"id":"challenge-1","randomData":"abc",'
        '"difficulty":2}}</script>'
    )

    challenge = parse_anubis_challenge(html)
    digest, nonce = solve_anubis_pow(
        challenge["random_data"],
        challenge["difficulty"],
    )

    assert challenge == {
        "id": "challenge-1",
        "random_data": "abc",
        "difficulty": 2,
    }
    assert digest == hashlib.sha256(f"abc{nonce}".encode()).hexdigest()
    assert digest.startswith("00")


def test_dryad_counts_use_all_rows_for_cpm_and_drop_duplicates(
    tmp_path,
):
    source = tmp_path / "gene_counts.txt"
    source.write_text(
        "\tSample_1\tSample_2\n"
        "Chr~1~GENEA\t10\t30\n"
        "Chr~2~GENEB\t30\t10\n"
        "Chr~3~DUP\t10\t10\n"
        "Chr~X~dup\t20\t20\n",
        encoding="utf-8",
    )
    target = tmp_path / "selected.tsv"

    summary = materialize_dryad_counts_matrix(
        source,
        target,
        ["Sample_1", "Sample_2"],
        minimum_gene_rows=2,
        expected_gene_rows=4,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    assert rows[0] == ["Hugo_Symbol", "Sample_1", "Sample_2"]
    assert [row[0] for row in rows[1:]] == ["GENEA", "GENEB"]
    assert float(rows[1][1]) == pytest.approx(
        math.log2((10 / 70) * 1_000_000 + 1)
    )
    assert float(rows[1][2]) == pytest.approx(
        math.log2((30 / 70) * 1_000_000 + 1)
    )
    assert summary["source_duplicated_symbol_labels"] == 1
    assert summary["source_duplicated_symbol_rows"] == 2
    assert summary["normalized_expression_gene_rows"] == 2


def test_dryad_clinical_table_requires_exact_time_and_status(tmp_path):
    clinical = tmp_path / "recurrence.txt"
    clinical.write_text(
        "SampleID\tVSCode\tYears\n"
        "Sample_1\t0\t4.5\n"
        "Sample_2\t1\t2.25\n",
        encoding="utf-8",
    )

    rows = read_dryad_clinical(clinical)

    assert [
        (row["patient_id"], row["event"], row["raw_time"])
        for row in rows
    ] == [
        ("Sample_1", 0, "4.5"),
        ("Sample_2", 1, "2.25"),
    ]


def test_dryad_file_pin_requires_exact_checksums_and_size():
    remote = {
        "_links": {"self": {"href": "/api/v2/files/96050"}},
        "path": "gene_counts.txt",
        "size": 12,
        "digestType": "md5",
        "digest": "a" * 32,
    }
    pin = validate_dryad_file_pin(
        {
            "file_id": 96050,
            "name": "gene_counts.txt",
            "size": 12,
            "md5": "a" * 32,
            "sha256": "b" * 64,
        },
        remote,
    )

    assert pin["file_id"] == 96050
    assert pin["sha256"] == "b" * 64

    with pytest.raises(ValueError, match="MD5"):
        validate_dryad_file_pin(
            {
                "file_id": 96050,
                "name": "gene_counts.txt",
                "size": 12,
                "md5": "c" * 32,
                "sha256": "b" * 64,
            },
            remote,
        )


def test_zenodo_selects_deepest_library_per_patient_with_stable_ties():
    selected = select_deepest_library_per_patient(
        {
            "P1_T1": {
                "_PATIENT_ID": "P1",
                "_LIBRARY_SIZE": 10,
            },
            "P1_T2": {
                "_PATIENT_ID": "P1",
                "_LIBRARY_SIZE": 20,
            },
            "P2_T2": {
                "_PATIENT_ID": "P2",
                "_LIBRARY_SIZE": 15,
            },
            "P2_T1": {
                "_PATIENT_ID": "P2",
                "_LIBRARY_SIZE": 15,
            },
        }
    )

    assert list(selected) == ["P1_T2", "P2_T1"]


def test_zenodo_patient_level_counts_map_exact_samples_and_normalize(
    tmp_path,
):
    archive_path = tmp_path / "source.zip"
    member = "source/counts.tsv"
    source = (
        "gene_name\tmeta\tP1_T1\tP1_T2\tP2_T1\tP2_N\n"
        "G1\tx\t1\t5\t2\t9\n"
        "G2\tx\t1\t1\t2\t9\n"
    )
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr(member, source)
    target = tmp_path / "selected.tsv"

    result = materialize_patient_level_counts(
        archive_path,
        archive_member=member,
        target=target,
        mapped_samples={
            "P1_T1": {"_PATIENT_ID": "P1"},
            "P1_T2": {"_PATIENT_ID": "P1"},
            "P2_T1": {"_PATIENT_ID": "P2"},
        },
        delimiter="\t",
        feature_column="gene_name",
        sample_start_column=3,
        allowed_unmapped_sample_patterns=["_N$"],
        expected_source_sample_columns=4,
        expected_mapped_sample_columns=3,
        expected_selected_patients=2,
        minimum_gene_rows=2,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    assert rows[0] == ["gene_name", "P1_T2", "P2_T1"]
    assert float(rows[1][1]) == pytest.approx(
        math.log2((5 / 6) * 1_000_000 + 1)
    )
    assert result["selected_expression_patients"] == 2
    assert result["source_unmapped_allowed_columns"] == 1


def test_zenodo_direct_matrix_maps_one_profile_per_patient(tmp_path):
    source = tmp_path / "adjustedCounts.txt"
    source.write_text(
        "ID\tSYMBOL\tCOHORT.01\tCOHORT.02\n"
        "ENSG1\tGENE1\t0\t12.5\n"
        "ENSG2\tGENE2\t7\t3\n",
        encoding="utf-8",
    )
    target = tmp_path / "selected.tsv"

    result = materialize_direct_patient_matrix(
        source,
        target=target,
        clinical_by_patient={
            "COHORT-01": {"Patient ID": "COHORT-01"},
            "COHORT-02": {"Patient ID": "COHORT-02"},
        },
        delimiter="\t",
        feature_column="SYMBOL",
        sample_start_column=3,
        sample_to_patient_replacements=[(".", "-")],
        expected_source_sample_columns=2,
        expected_selected_sample_columns=2,
        expected_gene_rows=2,
        minimum_gene_rows=2,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    assert rows == [
        ["SYMBOL", "COHORT.01", "COHORT.02"],
        ["GENE1", "0", "12.5"],
        ["GENE2", "7", "3"],
    ]
    assert result["selected_expression_patients"] == 2
    assert result["selected_samples"]["COHORT.01"]["_PATIENT_ID"] == (
        "COHORT-01"
    )


def test_zenodo_direct_matrix_rejects_multiple_profiles_per_patient(
    tmp_path,
):
    source = tmp_path / "adjustedCounts.txt"
    source.write_text(
        "SYMBOL\tP1.A\tP1-A\nGENE1\t1\t2\n",
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="multiple samples"):
        materialize_direct_patient_matrix(
            source,
            target=tmp_path / "selected.tsv",
            clinical_by_patient={"P1-A": {"Patient ID": "P1-A"}},
            delimiter="\t",
            feature_column="SYMBOL",
            sample_start_column=2,
            sample_to_patient_replacements=[(".", "-")],
            expected_source_sample_columns=2,
            expected_gene_rows=1,
            minimum_gene_rows=1,
        )


def test_zenodo_direct_matrix_selects_eligible_negative_profiles(
    tmp_path,
):
    source = tmp_path / "normalized.csv"
    source.write_text(
        "ID,NCI-1,TCGA-A,CONTROL_T1\n"
        "A1BG,-1.25,3,4\n"
        "TP53:ENSG00000141510.18,2.5,5,6\n",
        encoding="utf-8",
    )
    target = tmp_path / "selected.tsv"

    result = materialize_direct_patient_matrix(
        source,
        target=target,
        clinical_by_patient={"NCI-1": {"Sample ID": "NCI-1"}},
        delimiter=",",
        feature_column="ID",
        sample_start_column=2,
        sample_to_patient_replacements=[],
        allowed_unmapped_sample_patterns=[
            r"^TCGA-",
            r"^CONTROL_T1$",
        ],
        allow_negative_values=True,
        expected_source_sample_columns=3,
        expected_selected_sample_columns=1,
        expected_gene_rows=2,
        minimum_gene_rows=2,
    )

    assert list(csv.reader(target.open(), delimiter="\t")) == [
        ["ID", "NCI-1"],
        ["A1BG", "-1.25"],
        ["TP53:ENSG00000141510.18", "2.5"],
    ]
    assert result["source_unmapped_allowed_columns"] == 2
    assert result["source_expression_allows_negative_values"] is True


def test_zenodo_external_clinical_archive_is_safely_pinned(
    tmp_path,
    monkeypatch,
):
    archive_buffer = io.BytesIO()
    member = "tables/clinical.xlsx"
    member_bytes = b"pinned clinical workbook"
    with zipfile.ZipFile(archive_buffer, "w") as archive:
        archive.writestr(member, member_bytes)
    archive_bytes = archive_buffer.getvalue()

    def fake_download(_url, target):
        target.write_bytes(archive_bytes)

    monkeypatch.setattr(
        "app.repository.adapters.zenodo.download",
        fake_download,
    )
    clinical_path, archive_path, snapshot = (
        materialize_direct_clinical_file(
            {},
            {
                "url": "https://example.test/clinical.zip",
                "name": "clinical.zip",
                "size": len(archive_bytes),
                "sha256": hashlib.sha256(
                    archive_bytes
                ).hexdigest(),
                "archive_member": member,
                "extracted_name": "clinical.xlsx",
                "member_size": len(member_bytes),
                "member_sha256": hashlib.sha256(
                    member_bytes
                ).hexdigest(),
            },
            tmp_path,
        )
    )

    assert clinical_path.read_bytes() == member_bytes
    assert archive_path == tmp_path / "clinical.zip"
    assert snapshot["member"]["path"] == member


def test_zenodo_record_file_and_crossref_license_are_pinned():
    md5 = "0123456789abcdef0123456789abcdef"
    assert (
        validate_zenodo_file_pin(
            {
                "size": 42,
                "checksum": f"md5:{md5}",
                "links": {"self": "https://example.test/file"},
            },
            {"size": 42, "md5": md5},
        )
        == md5
    )
    validate_crossref_license(
        {
            "message": {
                "DOI": "10.1000/example",
                "license": [
                    {
                        "URL": (
                            "http://creativecommons.org/licenses/"
                            "by-nc-nd/4.0/"
                        )
                    }
                ],
            }
        },
        expected_doi="10.1000/example",
        expected_license_url=(
            "https://creativecommons.org/licenses/by-nc-nd/4.0/"
        ),
    )
    with pytest.raises(ValueError, match="archive pin changed"):
        validate_zenodo_file_pin(
            {
                "size": 41,
                "checksum": f"md5:{md5}",
                "links": {"self": "https://example.test/file"},
            },
            {"size": 42, "md5": md5},
        )


def test_biostudies_uvm_clinical_table_has_explicit_os_semantics():
    text = "\n".join(
        [
            "1 20 10 T4 No Spindle cell 95 NED",
            "2 14 13.5 T3 No Epithelioid cell 73 NED",
            "3 15 12.9 T3 No p.K666M Mixed 91 NED",
            "4 20 6.8 T4 Yes Epithelioid cell 30 DFD",
            "5 11 8.2 T2 No Epithelioid cell 68 NED",
            "6 10 12.7 T3 No Spindle cell 60 DFO",
            "7 13 13.3 T3 No p.R625C Spindle cell 68 AWD",
            "8 17 16.7 T4 Yes Mixed 19 DFD",
            "9 24 13.2 T4 No Spindle cell 69 NED",
            "10 22 14.5 T4 Yes p.R625H Epithelioid cell 54 DFD",
        ]
    )

    records = parse_uvm_clinical_text(text)

    assert len(records) == 10
    assert records[2]["sf3b1_mutation"] == "p.K666M"
    assert records[5]["latest_status"] == "DFO"
    assert records[6]["latest_status"] == "AWD"


def test_biostudies_sdrf_links_count_filename_to_tumor_number(tmp_path):
    sdrf = tmp_path / "study.sdrf.txt"
    sdrf.write_text(
        "Source Name\tCharacteristics[genotype]\t"
        "Derived Array Data File\tDerived Array Data File\n"
        + "\n".join(
            (
                f"Sample {index}\twild type genotype\t"
                f"A{index}T{index}.junctions.bed\t"
                f"A{index}T{index}.hg19.ensembl.gene_id.counts"
            )
            for index in range(1, 11)
        )
        + "\n",
        encoding="utf-8",
    )

    records = parse_biostudies_sdrf(sdrf)

    assert len(records) == 10
    assert records[1] == {
        "source_name": "Sample 1",
        "count_file": "A1T1.hg19.ensembl.gene_id.counts",
        "genotype": "wild type genotype",
    }


def test_biostudies_featurecounts_archives_are_joined_and_normalized(
    tmp_path,
):
    first_archive = tmp_path / "counts-1.zip"
    second_archive = tmp_path / "counts-2.zip"
    samples = []
    for index in range(1, 11):
        filename = (
            f"A{index}T{index}.hg19.ensembl.gene_id.counts"
        )
        samples.append(
            {
                "sample_id": f"S{index}",
                "count_file": filename,
                "tumor_number": index,
            }
        )
        archive_path = first_archive if index <= 5 else second_archive
        mode = "a" if archive_path.exists() else "w"
        with zipfile.ZipFile(archive_path, mode) as archive:
            archive.writestr(
                filename,
                "# Program:featureCounts v1.4.6\n"
                "Geneid\tChr\tStart\tEnd\tStrand\tLength\tSample\n"
                f"ENSG000001\tchr1\t1\t2\t+\t2\t{index}\n"
                f"ENSG000002\tchr1\t3\t4\t+\t2\t{index * 3}\n",
            )
    target = tmp_path / "expression.tsv"

    summary = materialize_featurecounts_archives(
        [first_archive, second_archive],
        target,
        samples,
        {
            "ENSG000001": "GENEA",
            "ENSG000002": "GENEB",
        },
        minimum_gene_rows=2,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    expected_low = math.log2(250_000 + 1)
    expected_high = math.log2(750_000 + 1)
    assert rows[0] == [
        "Ensembl_Gene_Id|Hugo_Symbol",
        *[f"S{index}" for index in range(1, 11)],
    ]
    assert float(rows[1][1]) == pytest.approx(expected_low)
    assert float(rows[2][1]) == pytest.approx(expected_high)
    assert summary["source_count_samples"] == 10
    assert summary["mapped_expression_gene_rows"] == 2
    assert summary["expression_normalization"] == "log2(CPM + 1)"


def test_biostudies_uromol_sdrf_preserves_progression_semantics(
    tmp_path,
):
    sdrf = tmp_path / "uromol.sdrf.txt"
    header = [
        "Source Name",
        "Characteristics[disease]",
        "Characteristics[tumor grading]",
        "Characteristics[disease staging]",
        "Characteristics[sex]",
        "Characteristics[age]",
        "Unit[time unit]",
        "Characteristics[clinical center]",
        "Characteristics[tumor size]",
        "Characteristics[growth pattern]",
        "Characteristics[BCG treatment]",
        "Characteristics[CIS in disease course]",
        "Characteristics[cystectomy]",
        "Characteristics[progression to T2+]",
        "Characteristics[progression free survival]",
        "Unit[time unit]",
    ]
    with sdrf.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(
            handle, delimiter="\t", lineterminator="\n"
        )
        writer.writerow(header)
        for index in range(10):
            writer.writerow(
                [
                    f"U{index:04d}",
                    "bladder tumor",
                    "high grade" if index % 2 else "low grade",
                    "T1" if index % 2 else "Ta",
                    "male" if index % 2 else "female",
                    str(50 + index),
                    "year",
                    "AAR",
                    "< 3",
                    "papillary",
                    "no",
                    "no",
                    "no",
                    "yes" if index < 5 else "no",
                    "0" if index == 9 else str(index + 1),
                    "month",
                ]
            )

    records = parse_uromol_sdrf(sdrf)

    assert len(records) == 10
    assert sum(int(row["progression_event"]) for row in records) == 5
    assert records[0]["progression_free_survival_months"] == "1"
    assert records[-1]["progression_free_survival_months"] == "0"
    assert records[1]["stage"] == "T1"


def test_biostudies_processed_fpkm_matrix_is_pinned_to_sdrf_order(
    tmp_path,
):
    source = tmp_path / "source-fpkm.tsv"
    target = tmp_path / "selected-fpkm.tsv"
    samples = [
        {
            "source_name": f"U{index:04d}",
            "sample_id": f"E-MTAB-4321-U{index:04d}-RNA",
        }
        for index in range(10)
    ]
    with source.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(
            handle, delimiter="\t", lineterminator="\n"
        )
        writer.writerow(
            [
                "tracking_id",
                "gene.type",
                "gene.status",
                "gene.name",
                *[row["source_name"] for row in samples],
            ]
        )
        writer.writerow(
            [
                "ENSG000001.4",
                "protein_coding",
                "KNOWN",
                "GeneA",
                *[str(index) for index in range(10)],
            ]
        )
        writer.writerow(
            [
                "ENSG000002.1",
                "protein_coding",
                "KNOWN",
                "TRUNCATED",
                *["1"] * 8,
            ]
        )
        writer.writerow(
            [
                "ENSG000003.2",
                "protein_coding",
                "KNOWN",
                "GeneB",
                *["0.5"] * 10,
            ]
        )
    malformed = [
        {
            "source_line": 3,
            "feature_id": "ENSG000002.1",
            "gene_symbol": "TRUNCATED",
            "observed_columns": 12,
            "expected_columns": 14,
        }
    ]

    summary = materialize_processed_fpkm_matrix(
        source,
        target,
        samples,
        expected_gene_rows=3,
        expected_malformed_rows=malformed,
        minimum_gene_rows=2,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    assert rows[0] == [
        "Ensembl_Gene_Id|Hugo_Symbol",
        *[row["sample_id"] for row in samples],
    ]
    assert rows[1][0] == "ENSG000001|GENEA"
    assert rows[2][0] == "ENSG000003|GENEB"
    assert summary["source_fpkm_samples"] == 10
    assert summary["source_fpkm_gene_rows"] == 3
    assert summary["source_fpkm_complete_gene_rows"] == 2
    assert summary["source_fpkm_malformed_rows"] == malformed


def build_minimal_bundle(root: Path) -> Path:
    derived = root / "derived"
    source = root / "source"
    source.mkdir(parents=True)
    (source / "LICENSE").write_text("Fixture license\n", encoding="utf-8")
    patient_ids = [f"P{index:02d}" for index in range(10)]
    sample_ids = [f"S{index:02d}" for index in range(10)]
    write_tsv(
        derived / "patients.tsv",
        ["patient_id"],
        [{"patient_id": patient_id} for patient_id in patient_ids],
    )
    write_tsv(
        derived / "samples.tsv",
        ["sample_id", "patient_id"],
        [
            {"sample_id": sample_id, "patient_id": patient_id}
            for sample_id, patient_id in zip(
                sample_ids, patient_ids, strict=True
            )
        ],
    )
    write_tsv(
        derived / "endpoint_os.tsv",
        ["patient_id", "time_days", "event"],
        [
            {
                "patient_id": patient_id,
                "time_days": 100 + index,
                "event": int(index < 5),
            }
            for index, patient_id in enumerate(patient_ids)
        ],
    )
    write_tsv(
        derived / "genes.tsv",
        [
            "gene_symbol",
            "original_gene_id",
            "row_number",
            "mapping_source",
        ],
        [
            {
                "gene_symbol": "GENEA",
                "original_gene_id": "1",
                "row_number": 0,
                "mapping_source": "fixture",
            },
            {
                "gene_symbol": "GENEB",
                "original_gene_id": "2",
                "row_number": 1,
                "mapping_source": "fixture",
            },
        ],
    )
    write_float32le_matrix(
        derived / "expression.float32le.bin",
        [
            [float(index) for index in range(10)],
            [float(index + 10) for index in range(10)],
        ],
    )
    metadata = {
        "dtype": "float32_le",
        "layout": "row_major_gene_by_sample",
        "gene_count": 2,
        "sample_count": 10,
        "sample_ids": sample_ids,
    }
    (derived / "expression.metadata.json").write_text(
        json.dumps(metadata),
        encoding="utf-8",
    )
    relative_files = [
        "derived/patients.tsv",
        "derived/samples.tsv",
        "derived/endpoint_os.tsv",
        "derived/genes.tsv",
        "derived/expression.float32le.bin",
        "derived/expression.metadata.json",
    ]
    manifest = {
        "schema_version": importer.BUNDLE_SCHEMA_VERSION,
        "dataset": {
            "id": "fixture-study",
            "cancer_code": "SKCM",
            "name": "Fixture study",
            "source_provider": "fixture",
            "source_accession": "FIXTURE",
            "source_url": "https://example.test/fixture",
            "assay": "bulk_rna_seq",
            "independence_status": "verified_external",
            "license_id": "LicenseRef-Fixture",
            "license_url": "https://example.test/license",
        },
        "release": {
            "id": "fixture-release-v1",
            "version": "v1",
            "source_snapshot": "fixture-sha",
        },
        "files": {
            "patients": "derived/patients.tsv",
            "samples": "derived/samples.tsv",
        },
        "endpoints": [
            {
                "endpoint_id": "OS",
                "label": "Overall survival",
                "time_origin": "Diagnosis",
                "event_definition": "Death from any cause",
                "source_time_column": "OS_DAYS",
                "source_event_column": "OS_STATUS",
                "source_time_unit": "days",
                "values_file": "derived/endpoint_os.tsv",
            }
        ],
        "expression_layers": [
            {
                "layer_id": "log2_rpkm",
                "label": "log2(RPKM + 1)",
                "source_unit": "RPKM",
                "analysis_unit": "log2(RPKM + 1)",
                "transform": "log2p",
                "is_default": True,
                "downloadable": True,
                "genes_file": "derived/genes.tsv",
                "matrix_file": "derived/expression.float32le.bin",
                "metadata_file": "derived/expression.metadata.json",
            }
        ],
        "source_files": {
            "license": "source/LICENSE",
        },
        "checksums": {
            relative: sha256_file(root / relative)
            for relative in [*relative_files, "source/LICENSE"]
        },
    }
    (root / "manifest.json").write_text(
        json.dumps(manifest),
        encoding="utf-8",
    )
    return root


def stub_survival_registry_disposition(monkeypatch) -> None:
    def capability(available: bool) -> SimpleNamespace:
        return SimpleNamespace(available=available)

    disposition = SimpleNamespace(
        universe_id="fixture-study",
        source_kind=SourceKind.EXTERNAL,
        cancer_code="SKCM",
        active=True,
        registry_category=StudyUniverseCategory.CATALOG_ONLY,
        endpoint_class="OS",
        endpoint_id="OS",
        capabilities=SimpleNamespace(
            catalog=capability(True),
            expression_comparison=capability(True),
            gsea=capability(True),
            survival=capability(True),
            hierarchical_pancancer=capability(False),
        ),
    )
    monkeypatch.setattr(
        importer,
        "require_study_registry_disposition",
        lambda registry_root, dataset_id: disposition,
    )
    monkeypatch.setattr(
        importer,
        "load_dataset_candidate_registry",
        lambda path: {"candidates": []},
    )


def test_repository_matrix_round_trip_and_safe_paths(tmp_path):
    matrix = tmp_path / "matrix.bin"
    write_float32le_matrix(matrix, [[1.25, 2.5], [3.75, 4.0]])

    assert read_float32le_row(
        matrix,
        row_number=1,
        sample_ids=["S1", "S2"],
    ) == pytest.approx({"S1": 3.75, "S2": 4.0})
    assert safe_bundle_path(tmp_path, "derived/matrix.bin").parent == (
        tmp_path / "derived"
    )
    with pytest.raises(ValueError, match="escapes"):
        safe_bundle_path(tmp_path, "../outside.bin")


def test_bundle_validation_enforces_checksums_and_qc(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(importer, "MIN_GENES", 2)
    bundle = build_minimal_bundle(tmp_path)

    validation = importer.validate_bundle(bundle)
    assert validation["qc"]["status"] == "passed"
    assert validation["qc"]["endpoints"]["OS"] == {
        "patients": 10,
        "events": 5,
        "censored": 5,
        "available": True,
    }

    matrix = bundle / "derived/expression.float32le.bin"
    matrix.write_bytes(matrix.read_bytes()[:-4])
    invalid = importer.validate_bundle(bundle)
    assert invalid["qc"]["status"] == "failed"
    assert any(
        "Checksum mismatch" in error for error in invalid["qc"]["errors"]
    )
    assert any(
        "matrix size" in error for error in invalid["qc"]["errors"]
    )


def test_repository_catalog_includes_endpoint_decision_summary(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(importer, "MIN_GENES", 2)
    stub_survival_registry_disposition(monkeypatch)
    bundle = build_minimal_bundle(tmp_path / "bundle")
    repository_root = tmp_path / "repository"
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    with factory() as db:
        db.add(
            CancerType(
                code="SKCM",
                tcga_cohort="TCGA-SKCM",
                name="Skin Cutaneous Melanoma",
                sort_order=0,
                coverage_status="available",
            )
        )
        db.commit()
        importer.promote_bundle(
            db,
            bundle,
            repository_root,
            tmp_path / "registry",
        )

        [dataset] = list_repository_datasets(db)

    assert dataset["id"] == "fixture-study"
    assert dataset["endpoints"] == [
        {
            "value": "OS",
            "label": "Overall survival",
            "standard_code": None,
            "patient_count": 10,
            "event_count": 5,
            "time_origin": "Diagnosis",
            "event_definition": "Death from any cause",
            "source_time_unit": "days",
        }
    ]


def test_bundle_validation_keeps_molecular_release_with_event_only_follow_up(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(importer, "MIN_GENES", 2)
    bundle = build_minimal_bundle(tmp_path)
    endpoint_path = bundle / "derived/endpoint_os.tsv"
    rows = importer.read_tsv(endpoint_path)
    write_tsv(
        endpoint_path,
        ["patient_id", "time_days", "event"],
        [{**row, "event": 1} for row in rows],
    )
    manifest_path = bundle / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["checksums"]["derived/endpoint_os.tsv"] = sha256_file(endpoint_path)
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    validation = importer.validate_bundle(bundle)

    assert validation["qc"]["status"] == "passed"
    assert validation["qc"]["endpoints"]["OS"] == {
        "patients": 10,
        "events": 10,
        "censored": 0,
        "available": False,
    }
    assert validation["qc"]["capabilities"]["expression"][
        "available"
    ] is True
    assert validation["qc"]["capabilities"]["gsea"][
        "available"
    ] is True
    assert validation["qc"]["capabilities"]["survival"][
        "available"
    ] is False


def test_bundle_validation_rejects_expression_fields_that_exceed_db_limits(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(importer, "MIN_GENES", 2)
    bundle = build_minimal_bundle(tmp_path)
    manifest_path = bundle / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["expression_layers"][0]["source_unit"] = "x" * 65
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    validation = importer.validate_bundle(bundle)

    assert validation["qc"]["status"] == "failed"
    assert any(
        "source_unit exceeds 64 characters" in error
        for error in validation["qc"]["errors"]
    )


def test_bundle_validation_rejects_dataset_fields_that_exceed_db_limits(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(importer, "MIN_GENES", 2)
    bundle = build_minimal_bundle(tmp_path)
    manifest_path = bundle / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["dataset"]["publication_id"] = "x" * 129
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    validation = importer.validate_bundle(bundle)

    assert validation["qc"]["status"] == "failed"
    assert any(
        "Dataset field publication_id exceeds 128 characters" in error
        for error in validation["qc"]["errors"]
    )


def test_active_release_revalidation_disables_only_ineligible_endpoint(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(importer, "MIN_GENES", 2)
    stub_survival_registry_disposition(monkeypatch)
    bundle = build_minimal_bundle(tmp_path / "bundle")
    repository_root = tmp_path / "repository"
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        db.add(
            CancerType(
                code="SKCM",
                tcga_cohort="TCGA-SKCM",
                name="Skin Cutaneous Melanoma",
                sort_order=0,
                coverage_status="available",
            )
        )
        db.commit()
        importer.promote_bundle(
            db,
            bundle,
            repository_root,
            tmp_path / "registry",
        )
        assert db.get(RepositoryDataset, "fixture-study").status == "available"

        monkeypatch.setattr(importer, "MIN_CENSORED", 6)
        result = importer.revalidate_active_releases(db, repository_root)

        dataset = db.get(RepositoryDataset, "fixture-study")
        endpoint = db.scalars(
            select(RepositoryEndpointDefinition)
        ).one()
        release = db.get(RepositoryRelease, dataset.active_release_id)
        assert result["available"] == 1
        assert result["qc_failed"] == 0
        assert dataset.status == "available"
        assert release.qc_status == "passed"
        assert endpoint.available is False
        assert endpoint.reason == (
            "Requires at least 6 censored observations."
        )
        assert len(list_repository_datasets(db)) == 1
        assert list_repository_datasets(
            db, analysis_type="survival"
        ) == []
        [expression_dataset] = list_repository_datasets(
            db, analysis_type="expression"
        )
        assert expression_dataset["id"] == "fixture-study"


def test_active_release_revalidation_fails_closed_on_invalid_expression(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(importer, "MIN_GENES", 2)
    stub_survival_registry_disposition(monkeypatch)
    bundle = build_minimal_bundle(tmp_path / "bundle")
    repository_root = tmp_path / "repository"
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        db.add(
            CancerType(
                code="SKCM",
                tcga_cohort="TCGA-SKCM",
                name="Skin Cutaneous Melanoma",
                sort_order=0,
                coverage_status="available",
            )
        )
        db.commit()
        importer.promote_bundle(
            db,
            bundle,
            repository_root,
            tmp_path / "registry",
        )
        release = db.get(RepositoryRelease, "fixture-release-v1")
        release_path = Path(release.repository_path)
        matrix_path = release_path / "derived/expression.float32le.bin"
        write_float32le_matrix(
            matrix_path,
            [
                [
                    float("nan"),
                    *[float(index) for index in range(1, 10)],
                ],
                [float(index + 10) for index in range(10)],
            ],
        )
        manifest_path = release_path / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["checksums"]["derived/expression.float32le.bin"] = (
            sha256_file(matrix_path)
        )
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

        result = importer.revalidate_active_releases(db, repository_root)

        dataset = db.get(RepositoryDataset, "fixture-study")
        assert result["available"] == 0
        assert result["qc_failed"] == 1
        assert dataset.status == "qc_failed"
        assert release.qc_status == "failed"
        assert release.qc_json["layers"]["log2_rpkm"][
            "nonfinite_values"
        ] == 1
        assert all(
            not details["available"]
            for details in release.qc_json["capabilities"].values()
        )
        assert list_repository_datasets(db) == []


def test_active_release_revalidation_rebases_legacy_manifest_path(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(importer, "MIN_GENES", 2)
    stub_survival_registry_disposition(monkeypatch)
    bundle = build_minimal_bundle(tmp_path / "bundle")
    repository_root = tmp_path / "repository"
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        db.add(
            CancerType(
                code="SKCM",
                tcga_cohort="TCGA-SKCM",
                name="Skin Cutaneous Melanoma",
                sort_order=0,
                coverage_status="available",
            )
        )
        db.commit()
        importer.promote_bundle(
            db,
            bundle,
            repository_root,
            tmp_path / "registry",
        )
        release = db.get(RepositoryRelease, "fixture-release-v1")
        layer = db.scalars(select(RepositoryExpressionLayer)).one()
        legacy_release = Path(
            "/repo/studies/fixture-study/releases/fixture-release-v1"
        )
        release.repository_path = str(legacy_release / "manifest.json")
        release.manifest_path = str(legacy_release / "manifest.json")
        layer.matrix_path = str(legacy_release / "derived/expression.float32le.bin")
        layer.metadata_path = str(
            legacy_release / "derived/expression.metadata.json"
        )
        db.commit()

        result = importer.revalidate_active_releases(db, repository_root)

        expected = (
            repository_root
            / "studies"
            / "fixture-study"
            / "releases"
            / "fixture-release-v1"
        )
        assert result["available"] == 1
        assert result["datasets"][0]["storage_path_rebased"] is True
        assert release.repository_path == str(expected)
        assert release.manifest_path == str(expected / "manifest.json")
        assert layer.matrix_path == str(
            expected / "derived/expression.float32le.bin"
        )
        assert layer.metadata_path == str(
            expected / "derived/expression.metadata.json"
        )


def test_repromotion_rebases_release_storage_paths(tmp_path, monkeypatch):
    monkeypatch.setattr(importer, "MIN_GENES", 2)
    stub_survival_registry_disposition(monkeypatch)
    bundle = build_minimal_bundle(tmp_path / "bundle")
    first_root = tmp_path / "first-repository"
    second_root = tmp_path / "second-repository"
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    with factory() as db:
        db.add(
            CancerType(
                code="SKCM",
                tcga_cohort="TCGA-SKCM",
                name="Skin Cutaneous Melanoma",
                sort_order=0,
                coverage_status="available",
            )
        )
        db.commit()
        importer.promote_bundle(
            db,
            bundle,
            first_root,
            tmp_path / "registry",
        )
        importer.promote_bundle(
            db,
            bundle,
            second_root,
            tmp_path / "registry",
        )

        release = db.get(RepositoryRelease, "fixture-release-v1")
        layer = db.scalars(select(RepositoryExpressionLayer)).one()
        source = db.get(DataSource, "external:fixture-study")
        data_manifest = db.scalars(select(DataManifest)).one()
        expected_release_path = (
            second_root
            / "studies"
            / "fixture-study"
            / "releases"
            / "fixture-release-v1"
        )

        assert release.repository_path == str(expected_release_path)
        assert release.manifest_path == str(
            expected_release_path / "manifest.json"
        )
        assert layer.matrix_path == str(
            expected_release_path / "derived/expression.float32le.bin"
        )
        assert layer.metadata_path == str(
            expected_release_path / "derived/expression.metadata.json"
        )
        assert Path(layer.matrix_path).is_file()
        assert source.source_path == str(expected_release_path)
        assert data_manifest.manifest_path == str(
            expected_release_path / "manifest.json"
        )


def test_external_analysis_contract_requires_pinned_release():
    with pytest.raises(
        ValidationError,
        match="dataset_id and dataset_release_id",
    ):
        AnalysisRequest(
            cohort="TCGA-SKCM",
            dataset_id="fixture-study",
            gene_symbol="TMEM176B",
        )

    request = AnalysisRequest(
        cohort="TCGA-SKCM",
        dataset_id="fixture-study",
        dataset_release_id="fixture-release-v1",
        expression_layer_id="log2_rpkm",
        gene_symbol="TMEM176B",
    )
    assert request.dataset_release_id == "fixture-release-v1"


def test_multiverse_propagates_external_release_to_every_cell():
    request = MultiverseAnalysisRequest(
        cohort="TCGA-SKCM",
        dataset_id="fixture-study",
        dataset_release_id="fixture-release-v1",
        expression_layer_id="log2_rpkm",
        genes=[SignatureGene(gene_symbol="TMEM176B")],
        endpoints=["OS"],
        scoring_methods=["single"],
        cutpoint_methods=["median", "upper_quartile"],
    )
    specifications = expand_multiverse_request(request)
    assert {
        row["analysis_request"].dataset_release_id
        for row in specifications
    } == {"fixture-release-v1"}

    result = summarize_multiverse(
        session_id="mv_external_fixture",
        request=request,
        specifications=specifications,
        completed_items=[],
        pipeline_version="test",
        data_version={
            "external_repository": {
                "dataset_id": "fixture-study",
                "release_id": "fixture-release-v1",
            }
        },
    )
    validated = MultiverseAnalysisOut(**result)
    assert validated.dataset_id == "fixture-study"
    assert validated.dataset_release_id == "fixture-release-v1"
    assert result["audit"]["data_version"]["external_repository"][
        "release_id"
    ] == "fixture-release-v1"


def test_external_sample_selection_uses_release_rank():
    preferred = ExternalSample("P1", "S-Z", 0)
    fallback = ExternalSample("P1", "S-A", 1)
    second_patient = ExternalSample("P2", "S-2", 0)
    candidates, warnings, summary = filter_sample_candidates(
        [fallback, preferred, second_patient],
        EmptyFilters(),
        endpoint_by_patient={
            "P1": ClinicalOutcome(
                endpoint="OS",
                time_days=100,
                event=1,
                source="external:fixture",
            ),
            "P2": ClinicalOutcome(
                endpoint="OS",
                time_days=200,
                event=0,
                source="external:fixture",
            ),
        },
        selection_rule="external_curated",
        endpoint_source="external:fixture",
    )
    retained, warnings, summary = select_expression_complete_samples(
        candidates,
        {preferred.barcode, fallback.barcode, second_patient.barcode},
        warnings=warnings,
        summary=summary,
    )

    assert {row.patient_id: row.barcode for row in retained} == {
        "P1": "S-Z",
        "P2": "S-2",
    }
    assert summary["selection_rule_kind"] == "external_curated"
    assert summary["priority_order"] == ["selection_rank", "sample_id"]


def test_expression_transform_is_explicit_and_does_not_double_log():
    assert apply_expression_transform([0.0, 3.0], "log2p") == [0.0, 2.0]
    provided = [1.25, -0.5]
    assert apply_expression_transform(provided, "identity") is provided
    with pytest.raises(ValueError, match="Unsupported"):
        apply_expression_transform([1.0], "inferred")


def test_source_feature_match_only_normalizes_hugo_symbol_case():
    assert source_feature_ids_match("C2orf15", "C2ORF15", "hugo_symbol")
    assert not source_feature_ids_match("PARK2", "PRKN", "hugo_symbol")
    assert not source_feature_ids_match("101", "0101", "entrez_gene_id")


def test_importer_treats_cbioportal_missing_tokens_as_null():
    assert importer._optional_float("[Not Available]") is None
    assert importer._optional_float("[Not Applicable]") is None
    assert importer._optional_float("62.5") == 62.5


def test_build_recipe_changes_with_spec_source_or_adapter():
    baseline = build_recipe_metadata(
        {"dataset": {"id": "example"}},
        "source-a",
        "adapter-a",
    )
    assert baseline == build_recipe_metadata(
        {"dataset": {"id": "example"}},
        "source-a",
        "adapter-a",
    )
    assert baseline["build_recipe_sha256"] != build_recipe_metadata(
        {"dataset": {"id": "changed"}},
        "source-a",
        "adapter-a",
    )["build_recipe_sha256"]
    assert baseline["build_recipe_sha256"] != build_recipe_metadata(
        {"dataset": {"id": "example"}},
        "source-b",
        "adapter-a",
    )["build_recipe_sha256"]
    assert baseline["build_recipe_sha256"] != build_recipe_metadata(
        {"dataset": {"id": "example"}},
        "source-a",
        "adapter-b",
    )["build_recipe_sha256"]
    assert baseline["build_recipe_sha256"] != build_recipe_metadata(
        {"dataset": {"id": "example"}},
        "source-a",
        "adapter-a",
        adapter_version="cbioportal_api_v1",
    )["build_recipe_sha256"]


def test_gdc_survival_merge_restores_censored_follow_up():
    patients = {
        "P-DEAD": {"OS_STATUS": "1:DECEASED"},
        "P-ALIVE": {"OS_STATUS": "0:LIVING"},
        "P-UPDATED": {"OS_STATUS": "0:LIVING"},
        "P-MISSING": {"OS_STATUS": "0:LIVING"},
    }
    payload = {
        "results": [
            {
                "donors": [
                    {
                        "submitter_id": "P-DEAD",
                        "project_id": "CPTAC-3",
                        "id": "case-dead",
                        "time": 100,
                        "censored": False,
                    },
                    {
                        "submitter_id": "P-ALIVE",
                        "project_id": "CPTAC-3",
                        "id": "case-alive",
                        "time": 250.5,
                        "censored": True,
                    },
                    {
                        "submitter_id": "P-UPDATED",
                        "project_id": "CPTAC-3",
                        "id": "case-updated",
                        "time": 300,
                        "censored": False,
                    },
                ]
            }
        ]
    }

    summary = merge_gdc_survival_into_patients(
        patients,
        set(patients),
        payload,
        project_id="CPTAC-3",
    )

    assert summary["matched_patients"] == 3
    assert summary["events"] == 2
    assert summary["censored"] == 1
    assert summary["missing_patients"] == 1
    assert summary["source_status_updates"] == 1
    assert patients["P-DEAD"]["GDC_OS_STATUS"] == "1:DECEASED"
    assert patients["P-ALIVE"]["GDC_OS_STATUS"] == "0:CENSORED"
    assert patients["P-ALIVE"]["GDC_OS_DAYS"] == "250.5"


def test_gdc_file_selection_is_patient_level_and_deterministic():
    def hit(
        patient: str,
        sample: str,
        aliquot: str,
        file_id: str,
        sample_type: str = "Primary Tumor",
    ):
        return {
            "file_id": file_id,
            "file_name": f"{file_id}.tsv",
            "file_size": 100,
            "md5sum": "a" * 32,
            "cases": [
                {
                    "case_id": f"case-{patient}",
                    "submitter_id": patient,
                    "samples": [
                        {
                            "sample_id": f"uuid-{sample}",
                            "submitter_id": sample,
                            "sample_type": sample_type,
                            "tissue_type": "Tumor",
                            "portions": [
                                {
                                    "analytes": [
                                        {
                                            "aliquots": [
                                                {
                                                    "aliquot_id": (
                                                        f"uuid-{aliquot}"
                                                    ),
                                                    "submitter_id": aliquot,
                                                }
                                            ]
                                        }
                                    ]
                                }
                            ],
                        }
                    ],
                }
            ],
        }

    selected, summary = select_gdc_expression_files(
        [
            hit("P1", "P1-01B", "P1-01B-01R", "file-c"),
            hit("P1", "P1-01A", "P1-01A-02R", "file-b"),
            hit("P1", "P1-01A", "P1-01A-01R", "file-a"),
            hit("P2", "P2-06A", "P2-06A-01R", "file-d", "Metastatic"),
            hit("P3", "P3-10A", "P3-10A-01R", "file-e", "Blood Derived Normal"),
        ],
        {"P1", "P2", "P3"},
        sample_type_priority=["Primary Tumor", "Metastatic"],
    )

    assert [row["file_id"] for row in selected] == ["file-a", "file-d"]
    assert summary["discarded_duplicate_file_candidates"] == 2
    assert summary["excluded_sample_types"] == {"Blood Derived Normal": 1}


def test_gdc_file_link_uses_associated_biospecimen_entity():
    hit = {
        "file_id": "file-1",
        "file_name": "file-1.tsv",
        "file_size": 100,
        "md5sum": "a" * 32,
        "associated_entities": [
            {
                "entity_id": "aliquot-b",
                "entity_type": "aliquot",
            }
        ],
        "cases": [
            {
                "case_id": "case-1",
                "submitter_id": "P1",
                "samples": [
                    {
                        "sample_id": "sample-a",
                        "submitter_id": "P1-01A",
                        "sample_type": "Primary Tumor",
                        "portions": [
                            {
                                "analytes": [
                                    {
                                        "aliquots": [
                                            {
                                                "aliquot_id": "aliquot-a",
                                                "submitter_id": "A",
                                            }
                                        ]
                                    }
                                ]
                            }
                        ],
                    },
                    {
                        "sample_id": "sample-b",
                        "submitter_id": "P1-01B",
                        "sample_type": "Primary Tumor",
                        "portions": [
                            {
                                "analytes": [
                                    {
                                        "aliquots": [
                                            {
                                                "aliquot_id": "aliquot-b",
                                                "submitter_id": "B",
                                            }
                                        ]
                                    }
                                ]
                            }
                        ],
                    },
                ],
            }
        ],
    }

    selected, _ = select_gdc_expression_files(
        [hit],
        {"P1"},
        sample_type_priority=["Primary Tumor"],
    )

    assert selected[0]["sample_id"] == "P1-01B"
    assert selected[0]["aliquot_id"] == "aliquot-b"
    assert selected[0]["linked_sample_candidates"] == 1


def test_gdc_case_eligibility_uses_prespecified_clinical_fields():
    case = {
        "primary_site": "Brain",
        "diagnoses": [
            {
                "diagnosis_id": "secondary",
                "diagnosis_is_primary_disease": False,
                "primary_diagnosis": "Astrocytoma, NOS",
            },
            {
                "diagnosis_id": "primary",
                "diagnosis_is_primary_disease": True,
                "primary_diagnosis": "Glioblastoma",
            },
        ],
        "demographic": {"sex_at_birth": "female"},
    }
    rules = [
        {
            "field": "case.primary_site",
            "include": ["Brain"],
            "required": True,
        },
        {
            "field": "diagnosis.primary_diagnosis",
            "include": ["Glioblastoma"],
            "required": True,
        },
    ]

    assert gdc_case_passes_eligibility(case, rules)
    assert not gdc_case_passes_eligibility(
        case,
        [
            *rules[:1],
            {
                "field": "diagnosis.primary_diagnosis",
                "include": ["Oligodendroglioma, NOS"],
            },
        ],
    )


def test_gdc_star_counts_are_materialized_in_shared_gene_order(tmp_path):
    def star(path: Path, sample_offset: int) -> None:
        path.write_text(
            "# gene-model: GENCODE v36\n"
            "gene_id\tgene_name\tgene_type\tunstranded\tstranded_first\t"
            "stranded_second\ttpm_unstranded\tfpkm_unstranded\t"
            "fpkm_uq_unstranded\n"
            "N_unmapped\t\t\t1\t1\t1\t\t\t\n"
            f"ENSG000001.1\tGENEA\tprotein_coding\t1\t1\t1\t"
            f"{1 + sample_offset}.0\t0\t0\n"
            f"ENSG000002.2\tGENEB\tprotein_coding\t1\t1\t1\t"
            f"{2 + sample_offset}.0\t0\t0\n",
            encoding="utf-8",
        )

    first = tmp_path / "first.tsv"
    second = tmp_path / "second.tsv"
    star(first, 0)
    star(second, 10)
    records = [
        {"file_id": "F1", "sample_id": "S1"},
        {"file_id": "F2", "sample_id": "S2"},
    ]
    output = tmp_path / "matrix.tsv"

    # Lower the production gene floor only by exercising the shard primitive
    # through a padded fixture.
    iterator, model = gdc_star_tpm_iterator(first)
    assert model == "GENCODE v36"
    assert list(iterator) == [
        ("ENSG000001.1", "GENEA", "1"),
        ("ENSG000002.2", "GENEB", "2"),
    ]

    def large_star(path: Path, sample_offset: int) -> None:
        rows = [
            "# gene-model: GENCODE v36",
            (
                "gene_id\tgene_name\tgene_type\tunstranded\t"
                "stranded_first\tstranded_second\ttpm_unstranded\t"
                "fpkm_unstranded\tfpkm_uq_unstranded"
            ),
        ]
        rows.extend(
            f"ENSG{index:011d}.1\tGENE{index}\tprotein_coding\t1\t1\t1\t"
            f"{index + sample_offset}.0\t0\t0"
            for index in range(1, 10_001)
        )
        path.write_text("\n".join(rows) + "\n", encoding="utf-8")

    large_star(first, 0)
    large_star(second, 10)
    summary = materialize_gdc_star_tpm_matrix(
        records,
        {"F1": first, "F2": second},
        output,
        shard_size=2,
    )
    with output.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == [
            "Ensembl_Gene_Id|Hugo_Symbol",
            "S1",
            "S2",
        ]
        assert next(reader) == [
            "ENSG00000000001.1|GENE1",
            "1",
            "11",
        ]
    assert summary["materialized_genes"] == 10_000
    assert summary["gene_models"] == ["GENCODE v36"]


def test_gdc_embedded_ensembl_hugo_mapping_is_offline_and_explicit():
    mapping = fetch_provided_ensembl_hugo_mapping(
        [
            "ENSG00000141510.18|TP53",
            "ENSG00000146648.22|EGFR",
            "N_unmapped|",
        ]
    )
    assert mapping == {
        "ENSG00000141510.18|TP53": "TP53",
        "ENSG00000146648.22|EGFR": "EGFR",
    }


def test_geo_embedded_hugo_entrez_mapping_is_offline_and_explicit():
    mapping = fetch_provided_hugo_entrez_mapping(
        [
            "A1BG-AS1|503538",
            "TP53|7157",
            "missing_entrez|",
            "missing_separator",
        ]
    )
    assert mapping == {
        "A1BG-AS1|503538": "A1BG-AS1",
        "TP53|7157": "TP53",
    }
    assert source_feature_ids_match(
        "Tp53|7157",
        "TP53|7157",
        "provided_hugo_entrez",
    )


def test_mixed_hugo_optional_ensembl_mapping_is_offline_and_explicit():
    mapping = fetch_provided_hugo_optional_ensembl_mapping(
        [
            "A1BG",
            "TP53:ENSG00000141510.18",
            "ASMT:ENSG00000196433.13_PAR_Y",
            "invalid:not-an-ensembl-id",
        ]
    )
    assert mapping == {
        "A1BG": "A1BG",
        "TP53:ENSG00000141510.18": "TP53",
        "ASMT:ENSG00000196433.13_PAR_Y": "ASMT",
    }
    assert source_feature_ids_match(
        "Tp53:ENSG00000141510.18",
        "TP53:ENSG00000141510.18",
        "provided_hugo_optional_ensembl",
    )


def test_geo_soft_and_compressed_csv_are_parsed_without_expansion(
    tmp_path,
):
    soft = tmp_path / "family.soft.gz"
    with gzip.open(soft, "wt", encoding="utf-8") as handle:
        handle.write(
            "^SERIES = GSE1\n"
            "!Series_title = Fixture series\n"
            "!Series_last_update_date = Jan 02 2024\n"
            "^SAMPLE = GSM1\n"
            "!Sample_title = F1\n"
            "!Sample_source_name_ch1 = Sigmoid colon\n"
            "!Sample_geo_accession = GSM1\n"
            "!Sample_description = Library name: Sample 1\n"
            "!Sample_supplementary_file_1 = "
            "ftp://example.org/GSM1_F1.counts.txt.gz\n"
            "!Sample_characteristics_ch1 = overall survival days: 100\n"
            "!Sample_characteristics_ch1 = overall survival event: 0\n"
            "^SAMPLE = GSM2\n"
            "!Sample_title = F1repl\n"
            "!Sample_geo_accession = GSM2\n"
            "!Sample_characteristics_ch1 = overall survival days: 100\n"
            "!Sample_characteristics_ch1 = overall survival event: 0\n"
        )
    series, samples = parse_geo_family_soft(soft)
    assert series["title"] == "Fixture series"
    assert [sample["title"] for sample in samples] == ["F1", "F1repl"]
    assert samples[0]["source_name"] == "Sigmoid colon"
    assert geo_sample_metadata_values(
        samples[0],
        "source_name",
    ) == ["Sigmoid colon"]
    assert samples[0]["descriptions"] == ["Library name: Sample 1"]
    assert samples[0]["supplementary_files"] == [
        "ftp://example.org/GSM1_F1.counts.txt.gz"
    ]
    assert geo_sample_metadata_values(
        samples[0],
        "supplementary_file",
    ) == ["ftp://example.org/GSM1_F1.counts.txt.gz"]
    assert samples[0]["characteristics"]["overall survival days"] == "100"

    expression = tmp_path / "expression.csv.gz"
    with gzip.open(
        expression, "wt", newline="", encoding="utf-8"
    ) as handle:
        writer = csv.writer(handle)
        writer.writerow(["", "F1", "F1repl"])
        writer.writerow(["TP53", 1.5, 1.6])
    assert read_expression_header(expression) == ["", "F1", "F1repl"]
    assert duplicated_source_feature_symbols(
        ["LRRc37A", "LRRC37A", "TP53"],
        "provided_hugo_symbol",
    ) == {"LRRC37A"}
    assert parse_geo_date(series["last_update_date"]) == (
        "2024-01-02T00:00:00Z"
    )
    assert "retrieved_at" not in series


def test_geo_plain_text_soft_is_supported(tmp_path):
    soft = tmp_path / "family.soft"
    soft.write_text(
        "^SERIES = GSE1\n"
        "!Series_title = Plain fixture\n"
        "!Series_last_update_date = Jan 02 2024\n"
        "^SAMPLE = GSM1\n"
        "!Sample_title = Tumor 1\n"
        "!Sample_geo_accession = GSM1\n",
        encoding="utf-8",
    )

    series, samples = parse_geo_family_soft(soft)

    assert series["title"] == "Plain fixture"
    assert [sample["geo_accession"] for sample in samples] == ["GSM1"]


def test_geo_soft_conflicts_require_an_explicit_policy(tmp_path):
    soft = tmp_path / "conflicting-family.soft.gz"
    with gzip.open(soft, "wt", encoding="utf-8") as handle:
        handle.write(
            "^SERIES = GSE3\n"
            "!Series_title = Conflicting fixture\n"
            "^SAMPLE = GSM30\n"
            "!Sample_title = Tumor 30\n"
            "!Sample_geo_accession = GSM30\n"
            "!Sample_characteristics_ch1 = tissue: baseline biopsy\n"
            "!Sample_characteristics_ch1 = tissue: breast\n"
        )

    with pytest.raises(
        ValueError,
        match="conflicting characteristic 'tissue'",
    ):
        parse_geo_family_soft(soft)

    _, first_samples = parse_geo_family_soft(
        soft,
        characteristic_conflict_policy={"Tissue": "first"},
    )
    assert first_samples[0]["characteristics"]["tissue"] == (
        "baseline biopsy"
    )
    assert first_samples[0]["characteristic_conflicts"]["tissue"] == [
        "baseline biopsy",
        "breast",
    ]

    _, last_samples = parse_geo_family_soft(
        soft,
        characteristic_conflict_policy={"tissue": "last"},
    )
    assert last_samples[0]["characteristics"]["tissue"] == "breast"
    assert last_samples[0]["characteristic_conflicts"]["tissue"] == [
        "baseline biopsy",
        "breast",
    ]


def test_geo_additional_family_soft_is_pinned_and_parsed(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "historical.soft.gz"
    with gzip.open(source, "wt", encoding="utf-8") as handle:
        handle.write(
            "^SERIES = GSE2\n"
            "!Series_title = Historical series\n"
            "!Series_last_update_date = Jan 03 2024\n"
            "^SAMPLE = GSM20\n"
            "!Sample_title = Historical sample\n"
            "!Sample_geo_accession = GSM20\n"
        )
    payload = source.read_bytes()

    def fake_download(_url, target):
        target.write_bytes(payload)

    monkeypatch.setattr(
        "app.repository.adapters.geo.download",
        fake_download,
    )
    materialized = materialize_geo_additional_soft_sources(
        [
            {
                "accession": "GSE2",
                "url": "https://example.test/GSE2_family.soft.gz",
                "size": len(payload),
                "sha256": hashlib.sha256(payload).hexdigest(),
            }
        ],
        tmp_path / "materialized",
    )

    assert materialized[0]["accession"] == "GSE2"
    assert materialized[0]["series"]["accession"] == "GSE2"
    assert materialized[0]["samples"][0]["geo_accession"] == "GSM20"


def test_geo_characteristic_filters_and_date_interval_are_explicit():
    characteristics = {
        "histology": "2",
        "stage tnm": "5",
        "mapping rate": "65.5",
        "surgery date": "2006-05-09",
        "vital date": "2008-02-03",
    }
    assert sample_passes_characteristic_filters(
        characteristics,
        [
            {
                "key": "histology",
                "include": ["2"],
                "required": True,
            },
            {
                "key": "stage tnm",
                "exclude": ["7"],
                "required": True,
            },
            {
                "key": "mapping rate",
                "minimum_numeric": 30,
                "maximum_numeric": 90,
                "required": True,
            },
        ],
    )
    assert not sample_passes_characteristic_filters(
        characteristics,
        [{"key": "histology", "include": ["1"], "required": True}],
    )
    assert derive_date_interval_days("2006-05-09", "2008-02-03") == "635"
    assert derive_date_interval_days("n/a", "2008-02-03") is None


def test_geo_patient_sample_selection_is_outcome_independent():
    samples = [
        {
            "title": "library-low",
            "geo_accession": "GSM1",
            "characteristics": {
                "patient id": "P1",
                "unique mapping": "44.0",
                "total mapping": "75.0",
                "os": "20",
                "event": "1",
            },
        },
        {
            "title": "library-high",
            "geo_accession": "GSM2",
            "characteristics": {
                "patient id": "P1",
                "unique mapping": "72.0",
                "total mapping": "80.0",
                "os": "20",
                "event": "1",
            },
        },
        {
            "title": "library-other",
            "geo_accession": "GSM3",
            "characteristics": {
                "patient id": "P2",
                "unique mapping": "61.0",
                "total mapping": "82.0",
                "os": "50",
                "event": "0",
            },
        },
    ]
    selected, summary = select_geo_patient_samples(
        samples,
        patient_id_characteristic="patient id",
        selection_spec={
            "rank_by": [
                {
                    "key": "unique mapping",
                    "type": "numeric",
                    "order": "desc",
                },
                {
                    "key": "total mapping",
                    "type": "numeric",
                    "order": "desc",
                },
            ],
            "require_equal_characteristics": ["os", "event"],
            "expected_selected_samples": 2,
            "expected_duplicate_patient_groups": 1,
            "expected_removed_samples": 1,
        },
    )
    assert [sample["geo_accession"] for sample in selected] == [
        "GSM2",
        "GSM3",
    ]
    assert summary["geo_duplicate_patient_groups"] == 1
    assert summary["geo_duplicate_samples_removed"] == 1

    samples[1]["characteristics"]["event"] = "0"
    with pytest.raises(ValueError, match="conflicting"):
        select_geo_patient_samples(
            samples,
            patient_id_characteristic="patient id",
            selection_spec={
                "rank_by": [
                    {
                        "key": "unique mapping",
                        "type": "numeric",
                        "order": "desc",
                    }
                ],
                "require_equal_characteristics": ["event"],
            },
        )


def test_geo_event_and_sample_identifiers_are_explicit():
    event_map = normalized_event_value_map({"yes": 1, "no": 0})
    assert map_geo_event("YES", event_map) == 1
    assert map_geo_event(" no ", event_map) == 0
    assert map_geo_event("unknown", event_map) is None
    assert map_geo_event("1", event_map) == 1
    with pytest.raises(ValueError, match="0/1"):
        normalized_event_value_map({"yes": 2})
    assert normalized_geo_categorical_value_map(
        {"0": "Female", "1": "Male"},
        field="sex",
    ) == {"0": "Female", "1": "Male"}
    with pytest.raises(ValueError, match="non-empty"):
        normalized_geo_categorical_value_map(
            {"0": ""},
            field="sex",
        )

    characteristics = {"alternative sample_id": "RIB0001"}
    assert (
        geo_sample_identifier(
            "Prostate_tumor_RIB0001",
            characteristics,
            "alternative sample_id",
            label="expression sample",
        )
        == "RIB0001"
    )
    assert (
        geo_sample_identifier(
            "Prostate_tumor_RIB0001",
            characteristics,
            "",
            label="patient",
        )
        == "Prostate_tumor_RIB0001"
    )
    assert (
        geo_sample_identifier(
            "Prostate_tumor_RIB0001",
            characteristics,
            "geo_accession",
            label="expression sample",
            geo_accession="GSM1234567",
        )
        == "GSM1234567"
    )
    with pytest.raises(ValueError, match="lacks"):
        geo_sample_identifier(
            "Prostate_tumor_RIB0001",
            characteristics,
            "missing",
            label="patient",
        )
    assert (
        geo_sample_metadata_identifier(
            {
                "title": "Discovery cohort 1",
                "descriptions": ["JMT1_84"],
            },
            "description",
            label="expression sample",
        )
        == "JMT1_84"
    )
    with pytest.raises(ValueError, match="exactly one"):
        geo_sample_metadata_identifier(
            {
                "title": "Discovery cohort 1",
                "descriptions": ["JMT1_84", "unexpected"],
            },
            "description",
            label="expression sample",
        )
    assert (
        normalize_geo_identifier(
            "T_06_01_A033a_T_SL9583",
            r"(?:_T)?_SL[0-9]+$",
            "",
            label="expression sample",
        )
        == "T_06_01_A033a"
    )
    with pytest.raises(ValueError, match="did not match"):
        normalize_geo_identifier(
            "T_06_01_A033a",
            r"_SL[0-9]+$",
            "",
            label="expression sample",
        )


def test_geo_curated_clinical_records_match_sample_descriptions():
    samples = [
        {
            "title": "230830-RN01",
            "geo_accession": "GSM1",
            "descriptions": ["Library name: Sample 1"],
            "characteristics": {"tissue": "Papillary Thyroid Carcinoma"},
        },
        {
            "title": "230831-RN01",
            "geo_accession": "GSM2",
            "descriptions": ["Library name: Sample 2"],
            "characteristics": {"tissue": "Papillary Thyroid Carcinoma"},
        },
        {
            "title": "230832-RN01",
            "geo_accession": "GSM3",
            "descriptions": ["Library name: Sample 3"],
            "characteristics": {"tissue": "Papillary Thyroid Carcinoma"},
        },
    ]
    records = [
        {
            "Sample ID": "Sample 1",
            "PFS": 1,
            "PFS duration (years)": 0.48,
        },
        {
            "Sample ID": "Sample 2",
            "PFS": 0,
            "PFS duration (years)": 19.73,
        },
    ]
    summary = attach_geo_curated_clinical_records(
        samples,
        {
            "records": records,
            "record_id_column": "Sample ID",
            "sample_metadata_field": "description",
            "sample_id_pattern": r"Library name: (Sample [0-9]+)",
            "field_name_map": {
                "PFS": "pfs event",
                "PFS duration (years)": "pfs years",
            },
            "expected_records": 2,
            "evidence_url": "https://example.test/supplement.xlsx",
            "evidence_size": 123,
            "evidence_sha256": "a" * 64,
        },
    )
    assert summary["geo_curated_clinical_records_matched"] == 2
    assert samples[0]["characteristics"]["sample id"] == "Sample 1"
    assert samples[0]["characteristics"]["pfs event"] == 1
    assert samples[1]["characteristics"]["pfs years"] == 19.73
    assert summary["geo_curated_clinical_field_name_map"] == {
        "PFS": "pfs event",
        "PFS duration (years)": "pfs years",
    }
    assert "curated_clinical_record" not in samples[2]

    with pytest.raises(ValueError, match="absent"):
        attach_geo_curated_clinical_records(
            samples,
            {
                "records": [
                    {
                        "Sample ID": "Sample 99",
                        "PFS": 1,
                        "PFS duration (years)": 1,
                    }
                ],
                "record_id_column": "Sample ID",
                "sample_metadata_field": "description",
                "sample_id_pattern": r"Library name: (Sample [0-9]+)",
                "evidence_url": "https://example.test/supplement.xlsx",
                "evidence_size": 123,
                "evidence_sha256": "a" * 64,
            },
        )


def test_geo_curated_clinical_records_normalize_linkage_keys():
    samples = [
        {
            "title": "RNA-Seq PT4_1",
            "geo_accession": "GSM4",
            "descriptions": [],
            "characteristics": {},
        },
        {
            "title": "RNA-Seq PT10_1",
            "geo_accession": "GSM10",
            "descriptions": [],
            "characteristics": {},
        },
    ]
    records = [
        {"Sample ID": "Pt 4", "DFS days": 264, "DFS event": 1},
        {"Sample ID": "Pt 10", "DFS days": 42, "DFS event": 1},
        {"Sample ID": "Pt 11", "DFS days": 500, "DFS event": 0},
    ]

    summary = attach_geo_curated_clinical_records(
        samples,
        {
            "records": records,
            "record_id_column": "Sample ID",
            "sample_metadata_field": "title",
            "sample_id_pattern": (
                r"(?i)(?:Pt\s*|RNA-Seq\s+PT)([0-9]+)(?:_1)?"
            ),
            "sample_id_replacement": r"PT\1",
            "allow_unmatched_records": True,
            "expected_records": 3,
            "expected_matched_records": 2,
            "expected_unmatched_records": 1,
            "evidence_url": "https://example.test/supplement.xlsx",
            "evidence_size": 123,
            "evidence_sha256": "a" * 64,
        },
    )

    assert samples[0]["characteristics"]["sample id"] == "Pt 4"
    assert samples[1]["characteristics"]["dfs days"] == 42
    assert summary["geo_curated_clinical_records_matched"] == 2
    assert summary["geo_curated_clinical_records_unmatched"] == 1
    assert (
        summary["geo_curated_clinical_sample_id_replacement"]
        == r"PT\1"
    )


def test_geo_curated_clinical_fields_use_event_time_or_censor_interval():
    records = [
        {
            "Case ID": "R01-001",
            "Survival Status": "Dead",
            "Time to Death (days)": "28",
            "CT Date": "01/01/2020",
            "Days between CT and surgery": "5",
            "Date of Last Known Alive": "12/31/2020",
        },
        {
            "Case ID": "R01-002",
            "Survival Status": "Alive",
            "Time to Death (days)": "",
            "CT Date": "01/01/2020",
            "Days between CT and surgery": "5",
            "Date of Last Known Alive": "02/01/2020",
        },
        {
            "Case ID": "R01-003",
            "Survival Status": "Unknown",
            "Time to Death (days)": "",
            "CT Date": "",
            "Days between CT and surgery": "",
            "Date of Last Known Alive": "",
        },
    ]

    derived, summary = derive_geo_curated_clinical_fields(
        records,
        [
            {
                "operation": "event_time_or_censor_interval",
                "target": "TCGA_TRACE_OS_DAYS",
                "status_field": "Survival Status",
                "event_values": ["Dead"],
                "censor_values": ["Alive"],
                "event_time_field": "Time to Death (days)",
                "censor_start_date_field": "CT Date",
                "censor_start_offset_days_field": (
                    "Days between CT and surgery"
                ),
                "censor_end_date_field": "Date of Last Known Alive",
                "date_format": "%m/%d/%Y",
            }
        ],
    )

    assert derived[0]["TCGA_TRACE_OS_DAYS"] == "28"
    assert derived[1]["TCGA_TRACE_OS_DAYS"] == "26"
    assert derived[2]["TCGA_TRACE_OS_DAYS"] == ""
    assert summary["TCGA_TRACE_OS_DAYS"] == {
        "operation": "event_time_or_censor_interval",
        "records": 3,
        "direct_event_records": 1,
        "date_derived_censor_records": 1,
        "unresolved_records": 1,
        "date_format": "%m/%d/%Y",
        "censor_start_offset_days_field": (
            "Days between CT and surgery"
        ),
    }


def test_geo_curated_clinical_fields_scale_numeric_difference():
    records = [
        {
            "Patient": "P01",
            "Age at diagnosis": "50",
            "Age at follow-up": "52.5",
        },
        {
            "Patient": "P02",
            "Age at diagnosis": "60",
            "Age at follow-up": "59",
        },
        {
            "Patient": "P03",
            "Age at diagnosis": "",
            "Age at follow-up": "63",
        },
    ]

    derived, summary = derive_geo_curated_clinical_fields(
        records,
        [
            {
                "operation": "numeric_difference_scaled",
                "target": "OS days",
                "start_field": "Age at diagnosis",
                "end_field": "Age at follow-up",
                "scale": 365.25,
                "require_positive": True,
            }
        ],
    )

    assert float(derived[0]["OS days"]) == pytest.approx(913.125)
    assert derived[1]["OS days"] == ""
    assert derived[2]["OS days"] == ""
    assert summary["OS days"] == {
        "operation": "numeric_difference_scaled",
        "records": 3,
        "derived_records": 1,
        "unresolved_records": 2,
        "start_field": "Age at diagnosis",
        "end_field": "Age at follow-up",
        "scale": 365.25,
        "require_positive": True,
    }


def test_geo_curated_clinical_fields_use_event_or_censor_time():
    records = [
        {
            "Patient": "AML01",
            "Time to death": "120",
            "Time to follow-up": "120",
        },
        {
            "Patient": "AML02",
            "Time to death": "",
            "Time to follow-up": "365",
        },
        {
            "Patient": "AML03",
            "Time to death": "",
            "Time to follow-up": "",
        },
    ]

    derived, summary = derive_geo_curated_clinical_fields(
        records,
        [
            {
                "operation": "event_time_or_censor_time",
                "target": "OS days",
                "status_target": "OS status",
                "event_time_field": "Time to death",
                "censor_time_field": "Time to follow-up",
                "event_value": "Dead",
                "censor_value": "Alive",
            }
        ],
    )

    assert derived[0]["OS days"] == "120"
    assert derived[0]["OS status"] == "Dead"
    assert derived[1]["OS days"] == "365"
    assert derived[1]["OS status"] == "Alive"
    assert derived[2]["OS days"] == ""
    assert derived[2]["OS status"] == ""
    assert summary["OS days"] == {
        "operation": "event_time_or_censor_time",
        "records": 3,
        "event_records": 1,
        "censor_records": 1,
        "unresolved_records": 1,
        "status_target": "OS status",
        "event_time_field": "Time to death",
        "censor_time_field": "Time to follow-up",
        "event_value": "Dead",
        "censor_value": "Alive",
    }


@pytest.mark.parametrize(
    "rule, message",
    [
        (
            {
                "operation": "numeric_difference_scaled",
                "target": "OS days",
                "start_field": "Missing",
                "end_field": "End",
                "scale": 1,
            },
            "existing start/end fields",
        ),
        (
            {
                "operation": "numeric_difference_scaled",
                "target": "OS days",
                "start_field": "Start",
                "end_field": "End",
                "scale": 0,
            },
            "positive finite scale",
        ),
    ],
)
def test_geo_numeric_difference_rejects_invalid_rules(rule, message):
    with pytest.raises(ValueError, match=message):
        derive_geo_curated_clinical_fields(
            [{"Start": "1", "End": "2"}],
            [rule],
        )


def test_geo_curated_clinical_xlsx_is_pinned_and_filtered(
    tmp_path,
    monkeypatch,
):
    from openpyxl import Workbook

    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clinical"
    sheet.append(["Patient ID", "DFS Event", "DFS Time"])
    sheet.append(["Patient_001", 1, 12])
    sheet.append(["Patient_002", 0, 24])
    sheet.append(["DFS means disease-free survival", "", ""])
    workbook.save(source)
    payload = source.read_bytes()

    def fake_download(_url, target):
        target.write_bytes(payload)

    monkeypatch.setattr(
        "app.repository.adapters.geo.download",
        fake_download,
    )
    resolved, evidence_path = (
        materialize_geo_curated_clinical_evidence(
            {
                "evidence_url": "https://example.test/source.xlsx",
                "evidence_size": len(payload),
                "evidence_sha256": hashlib.sha256(payload).hexdigest(),
                "evidence_format": "xlsx",
                "sheet_name": "Clinical",
                "header_row": 1,
                "record_id_column": "Patient ID",
                "sample_id_pattern": r"(Patient_[0-9]+)",
            },
            tmp_path / "materialized",
        )
    )

    assert evidence_path.read_bytes() == payload
    assert resolved["records"] == [
        {
            "Patient ID": "Patient_001",
            "DFS Event": "1",
            "DFS Time": "12",
        },
        {
            "Patient ID": "Patient_002",
            "DFS Event": "0",
            "DFS Time": "24",
        },
    ]


def test_geo_curated_clinical_soft_is_pinned_and_flattened(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "clinical.soft.gz"
    with gzip.open(source, "wt", encoding="utf-8") as handle:
        handle.write("^SERIES = GSE2\n")
        handle.write("!Series_title = Endpoint cohort\n")
        for index in (1, 2):
            handle.write(f"^SAMPLE = GSM{index}\n")
            handle.write(f"!Sample_title = CASE{index:03d} [2]\n")
            handle.write(
                f"!Sample_geo_accession = GSM{index}\n"
            )
            handle.write(
                f"!Sample_characteristics_ch1 = os day: {index * 10}\n"
            )
            handle.write(
                "!Sample_characteristics_ch1 = os bin: 0\n"
            )
    payload = source.read_bytes()

    def fake_download(_url, target):
        target.write_bytes(payload)

    monkeypatch.setattr(
        "app.repository.adapters.geo.download",
        fake_download,
    )
    resolved, evidence_path = materialize_geo_curated_clinical_evidence(
        {
            "evidence_url": "https://example.test/clinical.soft.gz",
            "evidence_size": len(payload),
            "evidence_sha256": hashlib.sha256(payload).hexdigest(),
            "evidence_format": "geo_soft",
            "record_id_column": "sample_title",
            "sample_id_pattern": r"(CASE[0-9]{3})(?: \[2\])?",
        },
        tmp_path / "materialized",
    )

    assert evidence_path.read_bytes() == payload
    assert resolved["records"] == [
        {
            "sample_title": "CASE001 [2]",
            "geo_accession": "GSM1",
            "os day": "10",
            "os bin": "0",
        },
        {
            "sample_title": "CASE002 [2]",
            "geo_accession": "GSM2",
            "os day": "20",
            "os bin": "0",
        },
    ]


def test_geo_curated_clinical_xlsx_supports_pmc_pow_download(
    tmp_path,
    monkeypatch,
):
    from openpyxl import Workbook

    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clinical"
    sheet.append(["Patient", "RFS Event", "RFS Months"])
    sheet.append(["Patient2", 1, 12])
    workbook.save(source)
    payload = source.read_bytes()
    calls = []

    def fake_download(_url, target, *, mode):
        calls.append(mode)
        target.write_bytes(payload)

    monkeypatch.setattr(
        "app.repository.adapters.pmc.download_publication_source",
        fake_download,
    )
    resolved, evidence_path = (
        materialize_geo_curated_clinical_evidence(
            {
                "evidence_url": "https://example.test/source.xlsx",
                "evidence_size": len(payload),
                "evidence_sha256": hashlib.sha256(payload).hexdigest(),
                "evidence_format": "xlsx",
                "download_mode": "pmc_pow",
                "sheet_name": "Clinical",
                "record_id_column": "Patient",
                "sample_id_pattern": r"(Patient[0-9]+)",
            },
            tmp_path / "materialized",
        )
    )

    assert calls == ["pmc_pow"]
    assert evidence_path.read_bytes() == payload
    assert resolved["download_mode"] == "pmc_pow"
    assert resolved["records"] == [
        {
            "Patient": "Patient2",
            "RFS Event": "1",
            "RFS Months": "12",
        }
    ]


def test_geo_curated_clinical_zip_xlsx_supports_blank_header(
    tmp_path,
    monkeypatch,
):
    from openpyxl import Workbook
    import zipfile

    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Figure 4F-1"
    sheet.append(["Figure 4F-1", "", ""])
    sheet.append(["", "", ""])
    sheet.append(["", "RFS time", "RFS event"])
    sheet.append(["Sample-01", 12, 1])
    sheet.append(["Sample-02", 24, 0])
    workbook.save(source)
    archive_path = tmp_path / "source-data.zip"
    archive_member = "source-data/clinical/source.xlsx"
    with zipfile.ZipFile(archive_path, mode="w") as archive:
        archive.write(source, archive_member)
    payload = archive_path.read_bytes()

    def fake_download(_url, target):
        target.write_bytes(payload)

    monkeypatch.setattr(
        "app.repository.adapters.geo.download",
        fake_download,
    )
    resolved, evidence_path = (
        materialize_geo_curated_clinical_evidence(
            {
                "evidence_url": "https://example.test/source-data.zip",
                "evidence_size": len(payload),
                "evidence_sha256": hashlib.sha256(payload).hexdigest(),
                "evidence_format": "zip_xlsx",
                "archive_member": archive_member,
                "sheet_name": "Figure 4F-1",
                "header_row": 3,
                "header_name_overrides": {"1": "Sample"},
                "record_id_column": "Sample",
                "sample_id_pattern": r"(Sample-[0-9]+)",
            },
            tmp_path / "materialized",
        )
    )

    assert evidence_path.read_bytes() == payload
    assert resolved["records"] == [
        {
            "Sample": "Sample-01",
            "RFS time": "12",
            "RFS event": "1",
        },
        {
            "Sample": "Sample-02",
            "RFS time": "24",
            "RFS event": "0",
        },
    ]


def test_geo_curated_clinical_zip_can_pin_the_consumed_member(
    tmp_path,
    monkeypatch,
):
    from openpyxl import Workbook

    source = tmp_path / "clinical.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clinical"
    sheet.append(["Patient", "OS"])
    sheet.append(["P01", 12])
    workbook.save(source)
    member_payload = source.read_bytes()
    archive_path = tmp_path / "dynamic-container.zip"
    archive_member = "clinical.xlsx"
    with zipfile.ZipFile(archive_path, mode="w") as archive:
        archive.writestr(archive_member, member_payload)
    archive_payload = archive_path.read_bytes()

    def fake_download(_url, target):
        target.write_bytes(archive_payload)

    monkeypatch.setattr(
        "app.repository.adapters.geo.download",
        fake_download,
    )
    resolved, evidence_path = (
        materialize_geo_curated_clinical_evidence(
            {
                "evidence_url": "https://example.test/dynamic.zip",
                "evidence_size": len(member_payload),
                "evidence_sha256": hashlib.sha256(
                    member_payload
                ).hexdigest(),
                "evidence_container_size": len(archive_payload),
                "evidence_pin_scope": "archive_member",
                "evidence_format": "zip_xlsx",
                "archive_member": archive_member,
                "sheet_name": "Clinical",
                "record_id_column": "Patient",
                "sample_id_pattern": r"(P[0-9]+)",
            },
            tmp_path / "materialized",
        )
    )

    assert evidence_path.read_bytes() == member_payload
    assert resolved["evidence_pin_scope"] == "archive_member"
    assert resolved["evidence_container_size"] == len(archive_payload)
    assert resolved["records"] == [{"Patient": "P01", "OS": "12"}]


def test_geo_curated_clinical_gzip_csv_is_pinned_and_filtered(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "clinical.csv.gz"
    with gzip.open(
        source,
        mode="wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "subject",
                "rnaseq_name",
                "Time.to.os",
                "OS.code",
                "Histology ",
            ]
        )
        writer.writerow(
            [
                "patient_2236",
                "rnaseq_2236",
                "74",
                "1",
                "Adenocarcinoma",
            ]
        )
        writer.writerow(
            [
                "patient_2976",
                "rnaseq_2976",
                "73",
                "0",
                "Squamous cell carcinoma",
            ]
        )
        writer.writerow(["footer", "not-a-library", "", "", ""])
    payload = source.read_bytes()

    def fake_download(_url, target):
        target.write_bytes(payload)

    monkeypatch.setattr(
        "app.repository.adapters.geo.download",
        fake_download,
    )
    resolved, evidence_path = (
        materialize_geo_curated_clinical_evidence(
            {
                "evidence_url": "https://example.test/clinical.csv.gz",
                "evidence_size": len(payload),
                "evidence_sha256": hashlib.sha256(payload).hexdigest(),
                "evidence_format": "csv.gz",
                "record_id_column": "rnaseq_name",
                "sample_id_pattern": r"(rnaseq_[0-9]+)",
            },
            tmp_path / "materialized",
        )
    )

    assert evidence_path.read_bytes() == payload
    assert resolved["records"] == [
        {
            "subject": "patient_2236",
            "rnaseq_name": "rnaseq_2236",
            "Time.to.os": "74",
            "OS.code": "1",
            "Histology": "Adenocarcinoma",
        },
        {
            "subject": "patient_2976",
            "rnaseq_name": "rnaseq_2976",
            "Time.to.os": "73",
            "OS.code": "0",
            "Histology": "Squamous cell carcinoma",
        },
    ]


def test_geo_delimited_clinical_evidence_rejects_duplicate_headers(
    tmp_path,
):
    source = tmp_path / "clinical.csv"
    source.write_text(
        "sample_id,sample_id,time\nS1,S1,12\n",
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="duplicate columns"):
        read_geo_delimited_clinical_evidence(
            source,
            delimiter=",",
            compressed=False,
            header_row=1,
        )


def test_geo_ensembl_matrix_uses_pinned_gtf_and_preserves_alias_columns(
    tmp_path,
):
    gtf = tmp_path / "Homo_sapiens.GRCh37.73.gtf.gz"
    with gzip.open(gtf, "wt", encoding="utf-8") as handle:
        for index in range(10_000):
            handle.write(
                "1\tensembl\texon\t1\t2\t.\t+\t.\t"
                f'gene_id "ENSG{index:011d}"; '
                f'gene_name "GENE{index}";\n'
            )
    source = tmp_path / "expression.tsv.gz"
    sample_ids = [f"L{index}T" for index in range(10)]
    sample_ids[-1] = "L9T_1"
    with gzip.open(
        source, "wt", newline="", encoding="utf-8"
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["Ensembl_gene_id", *sample_ids])
        for index in range(10_000):
            writer.writerow(
                [
                    f"LEGACY{index}_ENSG{index:011d}",
                    *([index] * len(sample_ids)),
                ]
            )
    target = tmp_path / "mapped.tsv.gz"
    summary = materialize_ensembl_expression(
        source,
        target,
        gtf,
        expression_sample_aliases={"L9T": "L9T_1"},
        feature_id_pattern=r".*_(ENSG[0-9]+)$",
    )
    with gzip.open(target, "rt", newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == [
            "Ensembl_Gene_Id|Hugo_Symbol",
            *sample_ids,
        ]
        assert next(reader)[0] == "ENSG00000000000|GENE0"
    assert summary["geo_mapped_expression_gene_rows"] == 10_000
    assert summary["geo_unambiguous_mapped_symbols"] == 10_000
    assert summary["geo_expression_feature_id_pattern"] == (
        r".*_(ENSG[0-9]+)$"
    )


def test_geo_entrez_matrix_uses_pinned_ncbi_gene_info(tmp_path):
    gene_info = tmp_path / "Homo_sapiens.gene_info.gz"
    with gzip.open(
        gene_info,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["#tax_id", "GeneID", "Symbol"])
        writer.writerow(["10090", "999999", "MouseGene"])
        for index in range(10_000):
            writer.writerow(["9606", index + 1, f"Gene{index}"])

    parsed = parse_ncbi_gene_info_map(gene_info)
    assert parsed["1"] == "GENE0"
    assert "999999" not in parsed

    source = tmp_path / "expression.csv.gz"
    sample_ids = [f"S{index:03d}T1R1" for index in range(10)]
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle)
        writer.writerow(["Entrez_Gene_Id", *sample_ids])
        for index in range(10_001):
            writer.writerow(
                [
                    index + 1,
                    *([index] * len(sample_ids)),
                ]
            )

    target = tmp_path / "mapped.tsv.gz"
    summary = materialize_entrez_expression(
        source,
        target,
        gene_info,
        expression_sample_aliases={"S001": "S001T1R1"},
    )
    with gzip.open(
        target,
        "rt",
        newline="",
        encoding="utf-8",
    ) as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == [
            "Hugo_Symbol|Entrez_Gene_Id",
            *sample_ids,
        ]
        assert next(reader)[0] == "GENE0|1"
    assert summary["geo_source_expression_gene_rows"] == 10_001
    assert summary["geo_mapped_expression_gene_rows"] == 10_000
    assert summary["geo_unmapped_expression_gene_rows"] == 1
    assert summary["geo_unambiguous_mapped_symbols"] == 10_000


def test_geo_count_matrix_is_normalized_to_log2_cpm(tmp_path):
    source = tmp_path / "counts.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["gene", "S1", "S2"])
        for index in range(10_000):
            writer.writerow([f"GENE{index}", 1, 3])

    target = tmp_path / "normalized.tsv"
    summary = materialize_geo_count_matrix(
        source,
        target,
        ["S2", "S1"],
        feature_column="gene",
        sample_start_column=2,
    )
    assert summary["source_expression_gene_rows"] == 10_000
    assert summary["expression_normalization"] == "log2(CPM + 1)"
    assert summary["expression_library_sizes"] == {
        "S2": "30000",
        "S1": "10000",
    }
    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["gene", "S2", "S1"]
        row = next(reader)
    assert row[0] == "GENE0"
    expected = math.log2(100 + 1)
    assert float(row[1]) == pytest.approx(expected)
    assert float(row[2]) == pytest.approx(expected)


def test_geo_count_matrix_drops_only_all_missing_features(tmp_path):
    source = tmp_path / "counts_with_annotations.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["gene", "S1", "S2", "", ""])
        writer.writerow(["MISSING", "NA", "NA", "annotation", "value"])
        for index in range(10_000):
            writer.writerow(
                [f"GENE{index}", 2, 4, "annotation", "value"]
            )

    target = tmp_path / "normalized.tsv"
    summary = materialize_geo_count_matrix(
        source,
        target,
        ["S1", "S2"],
        feature_column="gene",
        sample_start_column=2,
        drop_all_missing_features=True,
        missing_value_tokens=["NA"],
    )

    assert summary["source_expression_gene_rows"] == 10_001
    assert summary["normalized_expression_gene_rows"] == 10_000
    assert summary["source_expression_dropped_all_missing_features"] == 1
    assert summary["source_expression_missing_value_tokens"] == ["NA"]
    with target.open(encoding="utf-8") as handle:
        rows = list(csv.reader(handle, delimiter="\t"))
    assert len(rows) == 10_001
    assert rows[0] == ["gene", "S1", "S2"]
    assert rows[1][0] == "GENE0"


def test_geo_count_matrix_rejects_partially_missing_features(tmp_path):
    source = tmp_path / "partially_missing_counts.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["gene", "S1", "S2"])
        writer.writerow(["GENE0", "NA", 1])

    with pytest.raises(ValueError, match="partially missing"):
        materialize_geo_count_matrix(
            source,
            tmp_path / "normalized.tsv",
            ["S1", "S2"],
            feature_column="gene",
            sample_start_column=2,
            drop_all_missing_features=True,
            missing_value_tokens=["NA"],
        )


def test_geo_transcript_counts_are_aggregated_before_cpm(tmp_path):
    gene_reference = tmp_path / "Homo_sapiens.GRCh38.100.gtf.gz"
    with gzip.open(gene_reference, "wt", encoding="utf-8") as handle:
        handle.write("##gtf-version 3\n")
        handle.write(
            '1\ttest\ttranscript\t1\t10\t.\t+\t.\t'
            'gene_id "ENSG00000000001.1"; '
            'transcript_id "ENST00000000001.1"; '
            'gene_name "GENEA";\n'
        )
        handle.write(
            '1\ttest\ttranscript\t11\t20\t.\t+\t.\t'
            'gene_id "ENSG00000000001.1"; '
            'transcript_id "ENST00000000002.1"; '
            'gene_name "GENEA";\n'
        )
        for index in range(2, 10_001):
            handle.write(
                f'1\ttest\ttranscript\t{index * 10 + 1}\t'
                f'{index * 10 + 10}\t.\t+\t.\t'
                f'gene_id "ENSG{index:011d}.1"; '
                f'transcript_id "ENST{index + 1:011d}.1"; '
                f'gene_name "GENE{index}";\n'
            )

    source = tmp_path / "transcript_counts.tsv.gz"
    with gzip.open(source, "wt", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["transcript_id", "S1_count", "S2_count"])
        writer.writerow(["ENST00000000001", 1, 3])
        writer.writerow(["ENST00000000002", 2, 1])
        for index in range(2, 10_001):
            writer.writerow([f"ENST{index + 1:011d}", 0, 0])

    target = tmp_path / "gene_log2_cpm.tsv"
    summary = materialize_geo_transcript_count_matrix(
        source,
        target,
        ["S1_count", "S2_count"],
        gene_reference=gene_reference,
        feature_column="transcript_id",
        sample_start_column=2,
        count_value_type="integer",
        require_complete_reference=True,
    )

    with target.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle, delimiter="\t"))
    first_gene = rows[1]
    expected = math.log2(1_000_000 + 1)
    assert first_gene[0] == "ENSG00000000001|GENEA"
    assert float(first_gene[1]) == pytest.approx(expected)
    assert float(first_gene[2]) == pytest.approx(expected)
    assert summary["source_expression_transcript_rows"] == 10_001
    assert summary["geo_mapped_expression_transcript_rows"] == 10_001
    assert summary["geo_unmapped_expression_transcript_rows"] == 0
    assert summary["normalized_expression_gene_rows"] == 10_000
    assert summary["geo_transcript_reference_complete"] is True


def test_geo_recount3_matrix_is_auc_scaled_and_sample_linked(tmp_path):
    gene_sums = tmp_path / "sra.gene_sums.SRP1.G026.gz"
    with gzip.open(
        gene_sums,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        handle.write("##annotation=G026\n")
        handle.write("##date.generated=2020-05-12 01:01:40.180617\n")
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(
            ["gene_id", *[f"SRR{index}" for index in range(1, 11)]]
        )
        for index in range(10_000):
            writer.writerow(
                [f"ENSG{index:011d}.1", *([4_000] * 10)]
            )

    sample_metadata = tmp_path / "sra.sra.SRP1.MD.gz"
    with gzip.open(
        sample_metadata,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["external_id", "sample_title"])
        for index in range(1, 11):
            writer.writerow([f"SRR{index}", f"R01-{index:03d}"])

    qc_metadata = tmp_path / "sra.recount_qc.SRP1.MD.gz"
    with gzip.open(
        qc_metadata,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["external_id", "bc_auc.all_reads_all_bases"])
        for index in range(1, 11):
            writer.writerow(
                [
                    f"SRR{index}",
                    "20000000" if index == 2 else "40000000",
                ]
            )

    target = tmp_path / "normalized.tsv"
    summary = materialize_geo_recount3_matrix(
        gene_sums,
        target,
        [
            "R01-002",
            "R01-001",
            *[f"R01-{index:03d}" for index in range(3, 11)],
        ],
        sample_metadata=sample_metadata,
        qc_metadata=qc_metadata,
        run_id_column="external_id",
        sample_id_column="sample_title",
        auc_column="bc_auc.all_reads_all_bases",
        target_size=40_000_000,
        expected_annotation="G026",
        expected_source_runs=10,
    )

    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == [
            "Ensembl_Gene_Id",
            "R01-002",
            "R01-001",
            *[f"R01-{index:03d}" for index in range(3, 11)],
        ]
        row = next(reader)
    assert row[0] == "ENSG00000000000.1"
    assert float(row[1]) == pytest.approx(math.log2(8_001))
    assert float(row[2]) == pytest.approx(math.log2(4_001))
    assert summary["source_expression_columns"] == 10
    assert summary["selected_expression_columns"] == 10
    assert summary["source_expression_gene_rows"] == 10_000
    assert summary["recount3_scale_factor_min"] == pytest.approx(1)
    assert summary["recount3_scale_factor_max"] == pytest.approx(2)
    assert summary["recount3_annotation"] == "G026"


def test_geo_count_matrix_accepts_explicit_fractional_expected_counts(
    tmp_path,
):
    source = tmp_path / "rsem_expected_counts.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["gene", "S1", "S2"])
        for index in range(10_000):
            writer.writerow([f"GENE{index}", 1.25, 3.75])

    with pytest.raises(ValueError, match="integer"):
        materialize_geo_count_matrix(
            source,
            tmp_path / "rejected.tsv",
            ["S1", "S2"],
            feature_column="gene",
            sample_start_column=2,
        )

    target = tmp_path / "normalized.tsv"
    summary = materialize_geo_count_matrix(
        source,
        target,
        ["S1", "S2"],
        feature_column="gene",
        sample_start_column=2,
        count_value_type="nonnegative_numeric",
    )
    assert (
        summary["source_expression_count_value_type"]
        == "nonnegative_numeric"
    )
    assert summary["expression_library_sizes"] == {
        "S1": "12500",
        "S2": "37500",
    }


def test_geo_count_matrix_extracts_documented_feature_parts(tmp_path):
    source = tmp_path / "counts.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["GeneID", "GSM1", "GSM2"])
        for index in range(10_000):
            writer.writerow(
                [
                    (
                        f"ENSG{index:011d}|1|10|20|+|11|10|"
                        f"GENE{index}|protein_coding"
                    ),
                    2,
                    4,
                ]
            )

    target = tmp_path / "normalized.tsv"
    summary = materialize_geo_count_matrix(
        source,
        target,
        ["GSM1", "GSM2"],
        feature_column="GeneID",
        feature_column_output="Ensembl_Gene_Id|Hugo_Symbol",
        sample_start_column=2,
        feature_split_delimiter="|",
        feature_part_indexes=[1, 8],
        expected_feature_parts=9,
    )

    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == [
            "Ensembl_Gene_Id|Hugo_Symbol",
            "GSM1",
            "GSM2",
        ]
        assert next(reader)[0] == "ENSG00000000000|GENE0"
    assert summary["source_expression_feature_part_indexes"] == [1, 8]
    assert summary["source_expression_expected_feature_parts"] == 9


def test_geo_wide_normalized_matrix_preserves_source_scale_and_order(
    tmp_path,
):
    source = tmp_path / "normalized.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["gene_id", "gene_type", "S1", "S2", "S3"])
        for index in range(10_000):
            writer.writerow(
                [
                    f"ENSG{index:011d}",
                    "protein_coding",
                    0.5,
                    3.0,
                    15.0,
                ]
            )

    target = tmp_path / "selected.tsv"
    summary = materialize_geo_wide_normalized_matrix(
        source,
        target,
        ["S3", "S1"],
        feature_column="gene_id",
        feature_column_output="Ensembl_Gene_Id",
        sample_start_column=3,
        value_transform="log2p",
    )

    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["Ensembl_Gene_Id", "S3", "S1"]
        row = next(reader)
    assert row[0] == "ENSG00000000000"
    assert float(row[1]) == pytest.approx(4.0)
    assert float(row[2]) == pytest.approx(math.log2(1.5))
    assert summary["source_expression_columns"] == 3
    assert summary["selected_expression_columns"] == 2
    assert summary["source_expression_value_transform"] == "log2p"


def test_geo_wide_normalized_matrix_normalizes_source_sample_ids(
    tmp_path,
):
    source = tmp_path / "normalized.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(
            ["gene_id", "CASE001_lane_A", "CASE002_lane_B"]
        )
        for index in range(10_000):
            writer.writerow([f"GENE{index}", index, index + 1])

    target = tmp_path / "selected.tsv"
    summary = materialize_geo_wide_normalized_matrix(
        source,
        target,
        ["CASE002", "CASE001"],
        feature_column="gene_id",
        feature_column_output="Hugo_Symbol",
        sample_start_column=2,
        value_transform="identity",
        source_sample_id_pattern=r"^(CASE[0-9]{3})_.+$",
        source_sample_id_replacement=r"\1",
    )

    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["Hugo_Symbol", "CASE002", "CASE001"]
        assert next(reader) == ["GENE0", "1", "0"]
    assert summary["source_expression_sample_id_normalizations"] == 2


def test_geo_wide_normalized_matrix_excludes_ambiguous_features(
    tmp_path,
):
    source = tmp_path / "normalized.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["gene_id", "gene_name", "S1", "S2"])
        writer.writerow(["ENSG1", "TP53", 1, 3])
        writer.writerow(["ENSG2", "GeneA", 2, 4])
        writer.writerow(["ENSG3", "GENEA", 5, 6])
        for index in range(9_999):
            writer.writerow(
                [
                    f"ENSG{index + 4}",
                    f"GENE{index + 4}",
                    0,
                    1,
                ]
            )

    target = tmp_path / "selected.tsv"
    summary = materialize_geo_wide_normalized_matrix(
        source,
        target,
        ["S2", "S1"],
        feature_column="gene_name",
        feature_column_output="Hugo_Symbol",
        sample_start_column=3,
        value_transform="identity",
        duplicate_feature_policy="exclude_ambiguous",
        duplicate_feature_case_insensitive=True,
    )

    with target.open(encoding="utf-8") as handle:
        rows = list(csv.reader(handle, delimiter="\t"))
    assert rows[0] == ["Hugo_Symbol", "S2", "S1"]
    assert rows[1] == ["TP53", "3", "1"]
    assert all(row[0].casefold() != "genea" for row in rows[1:])
    assert summary["source_expression_gene_rows"] == 10_002
    assert summary["normalized_expression_gene_rows"] == 10_000
    assert summary["source_expression_duplicate_feature_symbols"] == 1
    assert summary["source_expression_duplicate_feature_rows"] == 2
    assert (
        summary["source_expression_excluded_duplicate_feature_rows"]
        == 2
    )


def test_geo_wide_normalized_matrix_drops_incomplete_selected_features(
    tmp_path,
):
    source = tmp_path / "normalized.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["gene_name", "S1", "S2", "S3"])
        writer.writerow(["DROP_S1", "NA", 1, 2])
        writer.writerow(["DROP_S2", 1, "NA", 2])
        writer.writerow(["KEEP_UNSELECTED_NA", 1, 2, "NA"])
        for index in range(10_000):
            writer.writerow([f"GENE{index}", 1, 2, 3])

    target = tmp_path / "selected.tsv"
    summary = materialize_geo_wide_normalized_matrix(
        source,
        target,
        ["S1", "S2"],
        feature_column="gene_name",
        feature_column_output="Hugo_Symbol",
        sample_start_column=2,
        value_transform="log2p",
        drop_incomplete_features=True,
        missing_value_tokens=["NA"],
    )

    with target.open(encoding="utf-8") as handle:
        rows = list(csv.reader(handle, delimiter="\t"))
    assert rows[0] == ["Hugo_Symbol", "S1", "S2"]
    assert rows[1][0] == "KEEP_UNSELECTED_NA"
    assert all(row[0] not in {"DROP_S1", "DROP_S2"} for row in rows)
    assert summary["source_expression_gene_rows"] == 10_003
    assert summary["normalized_expression_gene_rows"] == 10_001
    assert summary["source_expression_dropped_incomplete_features"] == 2
    assert summary["source_expression_missing_value_tokens"] == ["NA"]


def test_geo_wide_normalized_matrix_supports_missing_feature_header(
    tmp_path,
):
    source = tmp_path / "normalized.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["S1", "S2"])
        for index in range(10_000):
            writer.writerow([f"ENSG{index:011d}", 0.5, 3.0])

    target = tmp_path / "selected.tsv"
    summary = materialize_geo_wide_normalized_matrix(
        source,
        target,
        ["S2", "S1"],
        feature_column="",
        feature_column_index=1,
        feature_column_output="Ensembl_Gene_Id",
        sample_start_column=2,
        header_missing_feature_column=True,
        value_transform="log2p",
    )

    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["Ensembl_Gene_Id", "S2", "S1"]
        row = next(reader)
    assert row[0] == "ENSG00000000000"
    assert float(row[1]) == pytest.approx(2.0)
    assert float(row[2]) == pytest.approx(math.log2(1.5))
    assert summary["source_expression_columns"] == 2
    assert summary["source_expression_header_missing_feature_column"]


def test_geo_wide_normalized_matrix_supports_missing_leading_header(
    tmp_path,
):
    source = tmp_path / "normalized.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["gene_name", "accession_number", "S1", "S2"])
        for index in range(10_000):
            writer.writerow(
                [
                    index + 1,
                    f"GENE{index}",
                    f"NM_{index:06d}",
                    0.5,
                    3.0,
                ]
            )

    target = tmp_path / "selected.tsv"
    summary = materialize_geo_wide_normalized_matrix(
        source,
        target,
        ["S2", "S1"],
        feature_column="gene_name",
        feature_column_output="Hugo_Symbol",
        sample_start_column=4,
        header_missing_leading_column=True,
        value_transform="log2p",
    )

    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["Hugo_Symbol", "S2", "S1"]
        row = next(reader)
    assert row[0] == "GENE0"
    assert float(row[1]) == pytest.approx(2.0)
    assert float(row[2]) == pytest.approx(math.log2(1.5))
    assert summary["source_expression_columns"] == 2
    assert summary["source_expression_header_missing_leading_column"]


def test_geo_wide_normalized_matrix_allows_declared_negative_identity_values(
    tmp_path,
):
    source = tmp_path / "normalized.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["S1", "S2"])
        for index in range(10_000):
            writer.writerow([f"GENE{index}", -0.5, 3.0])

    target = tmp_path / "selected.tsv"
    summary = materialize_geo_wide_normalized_matrix(
        source,
        target,
        ["S2", "S1"],
        feature_column="",
        feature_column_index=1,
        feature_column_output="Hugo_Symbol",
        sample_start_column=2,
        header_missing_feature_column=True,
        value_transform="identity",
        allow_negative_values=True,
    )

    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["Hugo_Symbol", "S2", "S1"]
        assert next(reader) == ["GENE0", "3", "-0.5"]
    assert summary["source_expression_allow_negative_values"] is True

    with pytest.raises(
        ValueError,
        match="Negative source values may be allowed only",
    ):
        materialize_geo_wide_normalized_matrix(
            source,
            target,
            ["S1", "S2"],
            feature_column="",
            feature_column_index=1,
            feature_column_output="Hugo_Symbol",
            sample_start_column=2,
            header_missing_feature_column=True,
            value_transform="log2p",
            allow_negative_values=True,
        )


def test_geo_count_matrix_supports_explicit_unnamed_feature_column(
    tmp_path,
):
    source = tmp_path / "counts.tsv.gz"
    with gzip.open(
        source,
        "wt",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["", "S1", "S2"])
        for index in range(10_000):
            writer.writerow([f"ENSG{index:011d}", 2, 4])

    target = tmp_path / "normalized.tsv"
    summary = materialize_geo_count_matrix(
        source,
        target,
        ["S1", "S2"],
        feature_column="",
        feature_column_index=1,
        feature_column_output="Ensembl_Gene_Id",
        sample_start_column=2,
    )

    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["Ensembl_Gene_Id", "S1", "S2"]
        assert next(reader)[0] == "ENSG00000000000"
    assert summary["source_expression_feature_column"] == ""
    assert summary["source_expression_feature_column_index"] == 1
    assert (
        summary["normalized_expression_feature_column"]
        == "Ensembl_Gene_Id"
    )


def test_geo_featurecounts_tar_is_linked_and_normalized(tmp_path):
    source = tmp_path / "GSE1_RAW.tar"
    samples = []
    with tarfile.open(source, "w") as archive:
        for index in range(10):
            sample_id = f"Sample_{index:02d}"
            member_name = f"GSM{index:04d}_{sample_id}.counts.txt.gz"
            samples.append(
                {
                    "sample_id": sample_id,
                    "geo_accession": f"GSM{index:04d}",
                    "url": f"ftp://example.org/{member_name}",
                    "member_name": member_name,
                }
            )
            comment = (
                '"# Program:featureCounts; Command:""featureCounts"" '
                '""-o"" "\t\t\t\t\t\t\n'
                if index == 0
                else ""
            )
            text = comment + (
                "Geneid\tChr\tStart\tEnd\tStrand\tLength\t"
                f"Library_{index:02d}\n"
                f"ENSG000001\t1\t1\t2\t+\t2\t{index + 1}\n"
                f"ENSG000002\t1\t3\t4\t+\t2\t{3 * (index + 1)}\n"
            )
            payload = gzip.compress(text.encode("utf-8"), mtime=0)
            member = tarfile.TarInfo(member_name)
            member.size = len(payload)
            archive.addfile(member, io.BytesIO(payload))
    target = tmp_path / "expression.tsv"

    summary = materialize_geo_featurecounts_tar(
        source,
        target,
        samples,
        feature_column="Geneid",
        feature_column_output="Ensembl_Gene_Id",
        count_column_index=7,
        supplementary_file_pattern=re.compile(
            r".+\.counts\.txt\.gz"
        ),
        minimum_gene_rows=2,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    assert rows[0] == [
        "Ensembl_Gene_Id",
        *[f"Sample_{index:02d}" for index in range(10)],
    ]
    assert rows[1][0] == "ENSG000001"
    assert rows[2][0] == "ENSG000002"
    assert float(rows[1][1]) == pytest.approx(
        math.log2(250_000 + 1)
    )
    assert float(rows[2][1]) == pytest.approx(
        math.log2(750_000 + 1)
    )
    assert summary["source_expression_columns"] == 10
    assert summary["selected_expression_columns"] == 10
    assert summary["source_expression_gene_rows"] == 2
    assert summary["expression_normalization"] == "log2(CPM + 1)"
    assert summary["source_expression_count_columns"]["Sample_00"] == (
        "Library_00"
    )


def test_geo_featurecounts_tar_rejects_unsafe_members(tmp_path):
    source = tmp_path / "unsafe.tar"
    samples = []
    with tarfile.open(source, "w") as archive:
        for index in range(10):
            member_name = f"GSM{index:04d}.counts.txt.gz"
            samples.append(
                {
                    "sample_id": f"S{index}",
                    "geo_accession": f"GSM{index:04d}",
                    "url": f"ftp://example.org/{member_name}",
                    "member_name": member_name,
                }
            )
            payload = gzip.compress(
                (
                    "Geneid\tChr\tStart\tEnd\tStrand\tLength\tSample\n"
                    "ENSG000001\t1\t1\t2\t+\t2\t1\n"
                ).encode("utf-8"),
                mtime=0,
            )
            member = tarfile.TarInfo(member_name)
            member.size = len(payload)
            archive.addfile(member, io.BytesIO(payload))
        unsafe = tarfile.TarInfo("../escape.counts.txt.gz")
        unsafe_payload = gzip.compress(b"unsafe", mtime=0)
        unsafe.size = len(unsafe_payload)
        archive.addfile(unsafe, io.BytesIO(unsafe_payload))

    with pytest.raises(ValueError, match="unsafe member path"):
        materialize_geo_featurecounts_tar(
            source,
            tmp_path / "expression.tsv",
            samples,
            feature_column="Geneid",
            feature_column_output="Ensembl_Gene_Id",
            count_column_index=7,
            supplementary_file_pattern=re.compile(
                r".+\.counts\.txt\.gz"
            ),
            minimum_gene_rows=1,
        )


def test_geo_transcript_count_files_tar_is_linked(tmp_path):
    source = tmp_path / "GSE1_RAW.tar"
    samples = []
    with tarfile.open(source, "w") as archive:
        for index in range(10):
            sample_id = f"Sample_{index:02d}"
            member_name = f"GSM{index:04d}_{sample_id}_quant.sf.gz"
            samples.append(
                {
                    "sample_id": sample_id,
                    "geo_accession": f"GSM{index:04d}",
                    "url": f"ftp://example.org/{member_name}",
                    "member_name": member_name,
                }
            )
            text = (
                "Name\tLength\tEffectiveLength\tTPM\tNumReads\n"
                f"ENST000001.1\t100\t90\t1\t{index + 0.5}\n"
                f"ENST000002.2\t200\t180\t2\t{2 * index + 1.5}\n"
            )
            payload = gzip.compress(text.encode("utf-8"), mtime=0)
            member = tarfile.TarInfo(member_name)
            member.size = len(payload)
            archive.addfile(member, io.BytesIO(payload))
    target = tmp_path / "transcript_counts.tsv"

    summary = materialize_geo_transcript_count_files_tar(
        source,
        target,
        samples,
        feature_column="Name",
        count_column="NumReads",
        supplementary_file_pattern=re.compile(r".+_quant\.sf\.gz"),
        minimum_transcript_rows=2,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    assert rows[0] == [
        "Name",
        *[f"Sample_{index:02d}" for index in range(10)],
    ]
    assert rows[1][0] == "ENST000001.1"
    assert float(rows[1][1]) == pytest.approx(0.5)
    assert float(rows[2][-1]) == pytest.approx(19.5)
    assert summary["source_expression_columns"] == 10
    assert summary["selected_expression_columns"] == 10
    assert summary["source_expression_transcript_rows"] == 2
    assert summary["source_expression_count_column"] == "NumReads"


def test_geo_transcript_count_files_tar_rejects_feature_drift(tmp_path):
    source = tmp_path / "GSE1_RAW.tar"
    samples = []
    with tarfile.open(source, "w") as archive:
        for index in range(10):
            member_name = f"GSM{index:04d}_quant.sf.gz"
            samples.append(
                {
                    "sample_id": f"S{index}",
                    "geo_accession": f"GSM{index:04d}",
                    "url": f"ftp://example.org/{member_name}",
                    "member_name": member_name,
                }
            )
            second_feature = (
                "ENST000003.1" if index == 9 else "ENST000002.1"
            )
            text = (
                "Name\tLength\tEffectiveLength\tTPM\tNumReads\n"
                "ENST000001.1\t100\t90\t1\t1.5\n"
                f"{second_feature}\t200\t180\t2\t2.5\n"
            )
            payload = gzip.compress(text.encode("utf-8"), mtime=0)
            member = tarfile.TarInfo(member_name)
            member.size = len(payload)
            archive.addfile(member, io.BytesIO(payload))

    with pytest.raises(ValueError, match="feature order differs"):
        materialize_geo_transcript_count_files_tar(
            source,
            tmp_path / "transcript_counts.tsv",
            samples,
            feature_column="Name",
            count_column="NumReads",
            supplementary_file_pattern=re.compile(
                r".+_quant\.sf\.gz"
            ),
            minimum_transcript_rows=2,
        )


def test_geo_normalized_files_tar_is_linked_and_preserves_scale(tmp_path):
    source = tmp_path / "GSE1_RAW.tar"
    samples = []
    with tarfile.open(source, "w") as archive:
        for index in range(10):
            sample_id = f"Sample_{index:02d}"
            member_name = (
                f"GSM{index:04d}_{sample_id}.normalized.tsv.gz"
            )
            samples.append(
                {
                    "sample_id": sample_id,
                    "geo_accession": f"GSM{index:04d}",
                    "url": f"ftp://example.org/{member_name}",
                    "member_name": member_name,
                }
            )
            text = (
                f"GeneID\t{sample_id}\n"
                f"GENE1\t{index + 0.25}\n"
                f"GENE2\t{index + 1.5}\n"
            )
            payload = gzip.compress(text.encode("utf-8"), mtime=0)
            member = tarfile.TarInfo(member_name)
            member.size = len(payload)
            archive.addfile(member, io.BytesIO(payload))
    target = tmp_path / "expression.tsv"

    summary = materialize_geo_normalized_files_tar(
        source,
        target,
        samples,
        feature_column="GeneID",
        feature_column_output="Hugo_Symbol",
        value_column_index=2,
        supplementary_file_pattern=re.compile(
            r".+\.normalized\.tsv\.gz"
        ),
        value_transform="identity",
        require_nonnegative=True,
        require_value_header_matches_sample_id=True,
        minimum_gene_rows=2,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    assert rows[0] == [
        "Hugo_Symbol",
        *[f"Sample_{index:02d}" for index in range(10)],
    ]
    assert rows[1][0] == "GENE1"
    assert float(rows[1][1]) == pytest.approx(0.25)
    assert float(rows[2][-1]) == pytest.approx(10.5)
    assert summary["source_expression_columns"] == 10
    assert summary["selected_expression_columns"] == 10
    assert summary["source_expression_gene_rows"] == 2
    assert summary["normalized_expression_gene_rows"] == 2
    assert summary["expression_normalization"] == (
        "source normalized value"
    )
    assert summary["source_expression_value_columns"]["Sample_00"] == (
        "Sample_00"
    )


def test_geo_normalized_files_tar_can_align_feature_order(tmp_path):
    source = tmp_path / "GSE2_RAW.tar"
    samples = []
    with tarfile.open(source, "w") as archive:
        for index in range(10):
            sample_id = f"Sample_{index:02d}"
            member_name = f"GSM{index:04d}_{sample_id}.tsv.gz"
            samples.append(
                {
                    "sample_id": sample_id,
                    "geo_accession": f"GSM{index:04d}",
                    "url": f"https://example.test/{member_name}",
                    "member_name": member_name,
                }
            )
            rows = [
                ("GENE1", index + 1),
                ("GENE2", index + 101),
            ]
            if index % 2:
                rows.reverse()
            text = "GeneID\tFPKM\n" + "".join(
                f"{gene}\t{value}\n" for gene, value in rows
            )
            payload = gzip.compress(text.encode("utf-8"), mtime=0)
            member = tarfile.TarInfo(member_name)
            member.size = len(payload)
            archive.addfile(member, io.BytesIO(payload))
    target = tmp_path / "expression.tsv"

    summary = materialize_geo_normalized_files_tar(
        source,
        target,
        samples,
        feature_column="GeneID",
        feature_column_output="Hugo_Symbol",
        value_column_index=2,
        supplementary_file_pattern=re.compile(r".+\.tsv\.gz"),
        feature_order_policy="align",
        minimum_gene_rows=2,
    )

    rows = list(csv.reader(target.open(), delimiter="\t"))
    assert rows[1] == [
        "GENE1",
        *[str(index + 1) for index in range(10)],
    ]
    assert rows[2] == [
        "GENE2",
        *[str(index + 101) for index in range(10)],
    ]
    assert summary["source_expression_feature_order_policy"] == "align"


def test_icgc_listing_endpoint_and_analysis_selection_are_deterministic():
    listing = ET.fromstring(
        """
        <ListBucketResult xmlns="http://s3.amazonaws.com/doc/2006-03-01/">
          <Contents>
            <Key>release_28/data/LIRI-JP/DO2/exp_seq/file.tsv.gz</Key>
            <LastModified>2019-11-26T00:00:00.000Z</LastModified>
            <ETag>"bbbbbbbbbbbbbbbbbbbbbbbbbbbbbbbb"</ETag>
            <Size>20</Size>
          </Contents>
          <Contents>
            <Key>release_28/data/LIRI-JP/DO1/exp_seq/file.tsv.gz</Key>
            <LastModified>2019-11-25T00:00:00.000Z</LastModified>
            <ETag>"aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa"</ETag>
            <Size>10</Size>
          </Contents>
          <IsTruncated>true</IsTruncated>
          <NextContinuationToken>page-2</NextContinuationToken>
        </ListBucketResult>
        """
    )
    records, truncated, token = parse_object_listing(listing)
    assert records == [
        {
            "key": "release_28/data/LIRI-JP/DO2/exp_seq/file.tsv.gz",
            "etag": "bbbbbbbbbbbbbbbbbbbbbbbbbbbbbbbb",
            "size": 20,
            "last_modified": "2019-11-26T00:00:00.000Z",
        },
        {
            "key": "release_28/data/LIRI-JP/DO1/exp_seq/file.tsv.gz",
            "etag": "aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa",
            "size": 10,
            "last_modified": "2019-11-25T00:00:00.000Z",
        },
    ]
    assert truncated is True
    assert token == "page-2"
    assert all("retrieved_at" not in record for record in records)

    assert resolve_icgc_os(
        {
            "donor_vital_status": "deceased",
            "donor_survival_time": "730.5",
            "donor_interval_of_last_followup": "700",
        }
    ) == {
        "time_days": 730.5,
        "event": 1,
        "time_source": "donor_survival_time",
    }
    assert resolve_icgc_os(
        {
            "donor_vital_status": "alive",
            "donor_survival_time": "",
            "donor_interval_of_last_followup": "900",
        }
    ) == {
        "time_days": 900.0,
        "event": 0,
        "time_source": "donor_interval_of_last_followup",
    }

    primary = {
        "icgc_sample_id": "SA-PRIMARY",
        "icgc_specimen_id": "SP-PRIMARY",
        "analysis_id": "AN-Z",
        "gene_model": "RefSeq",
        "normalization_algorithm": "FKPM",
        "gene_rows": 13_405,
    }
    later_analysis_id = {
        **primary,
        "analysis_id": "AN-A",
    }
    metastatic = {
        **primary,
        "icgc_sample_id": "SA-MET",
        "icgc_specimen_id": "SP-MET",
        "analysis_id": "AN-0",
    }
    selected = select_expression_analysis(
        [primary, metastatic, later_analysis_id],
        {
            "SA-PRIMARY": {
                "icgc_specimen_id": "SP-PRIMARY",
                "submitted_sample_id": "SUB-PRIMARY",
            },
            "SA-MET": {
                "icgc_specimen_id": "SP-MET",
                "submitted_sample_id": "SUB-MET",
            },
        },
        {
            "SP-PRIMARY": {
                "specimen_type": "Primary tumour - solid tissue",
                "tumour_histological_type": (
                    "Pancreatic Ductal Adenocarcinoma"
                ),
            },
            "SP-MET": {
                "specimen_type": (
                    "Metastatic tumour - metastasis to distant location"
                ),
                "tumour_histological_type": (
                    "Pancreatic Ductal Adenocarcinoma"
                ),
            },
        },
        specimen_priority=(
            "Primary tumour - solid tissue",
            "Metastatic tumour - metastasis to distant location",
        ),
        required_gene_model="RefSeq",
        required_normalization_algorithm="FKPM",
        required_tumour_histological_types=(
            "Pancreatic Ductal Adenocarcinoma",
        ),
    )
    assert selected is later_analysis_id

    assert (
        select_expression_analysis(
            [primary],
            {
                "SA-PRIMARY": {
                    "icgc_specimen_id": "SP-PRIMARY",
                    "submitted_sample_id": "SUB-PRIMARY",
                },
            },
            {
                "SP-PRIMARY": {
                    "specimen_type": "Primary tumour - solid tissue",
                    "tumour_histological_type": "Acinar Cell Carcinoma",
                },
            },
            specimen_priority=("Primary tumour - solid tissue",),
            required_gene_model="RefSeq",
            required_normalization_algorithm="FKPM",
            required_tumour_histological_types=(
                "Pancreatic Ductal Adenocarcinoma",
            ),
        )
        is None
    )


def test_icgc_recurrence_interval_requires_exact_event_or_censor_data():
    censor_statuses = ("complete remission",)

    assert resolve_icgc_recurrence_interval(
        {
            "disease_status_last_followup": "relapse",
            "donor_relapse_type": "local recurrence",
            "donor_relapse_interval": "169",
            "donor_interval_of_last_followup": "2703",
        },
        censor_disease_statuses=censor_statuses,
    ) == {
        "time_days": 169.0,
        "event": 1,
        "time_source": "donor_relapse_interval",
    }
    assert resolve_icgc_recurrence_interval(
        {
            "disease_status_last_followup": "complete remission",
            "donor_relapse_type": "",
            "donor_relapse_interval": "",
            "donor_interval_of_last_followup": "1542",
        },
        censor_disease_statuses=censor_statuses,
    ) == {
        "time_days": 1542.0,
        "event": 0,
        "time_source": "donor_interval_of_last_followup",
    }
    assert (
        resolve_icgc_recurrence_interval(
            {
                "disease_status_last_followup": "complete remission",
                "donor_relapse_type": "local recurrence",
                "donor_relapse_interval": "",
                "donor_interval_of_last_followup": "2000",
            },
            censor_disease_statuses=censor_statuses,
        )
        is None
    )
    assert (
        resolve_icgc_recurrence_interval(
            {
                "disease_status_last_followup": "progression",
                "donor_relapse_type": "",
                "donor_relapse_interval": "",
                "donor_interval_of_last_followup": "926",
            },
            censor_disease_statuses=censor_statuses,
        )
        is None
    )
    assert (
        resolve_icgc_recurrence_interval(
            {
                "disease_status_last_followup": "",
                "donor_relapse_type": "",
                "donor_relapse_interval": "",
                "donor_interval_of_last_followup": "900",
            },
            censor_disease_statuses=censor_statuses,
        )
        is None
    )


def test_icgc_duplicate_source_genes_require_an_explicit_policy(tmp_path):
    expression_path = tmp_path / "expression.tsv.gz"
    header = [
        "icgc_sample_id",
        "analysis_id",
        "gene_id",
        "normalized_read_count",
    ]
    with gzip.open(
        expression_path, "wt", encoding="utf-8", newline=""
    ) as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        for index in range(10_000):
            writer.writerow(
                ["SA-1", "AN-1", f"GENE{index:05d}", str(index)]
            )
        writer.writerow(["SA-1", "AN-1", "DUPLICATED", "1"])
        writer.writerow(["SA-1", "AN-1", "DUPLICATED", "2"])

    selected = {
        "analysis": {
            "icgc_sample_id": "SA-1",
            "analysis_id": "AN-1",
        },
        "expression_paths": [expression_path],
    }
    with pytest.raises(ValueError, match="contains duplicate gene"):
        selected_expression_values(selected, header)

    duplicate_counts = []
    values = selected_expression_values(
        selected,
        header,
        duplicate_gene_policy="drop",
        duplicate_gene_counts=duplicate_counts,
    )
    assert len(values) == 10_000
    assert "DUPLICATED" not in values
    assert duplicate_counts == [1]


def test_icgc_gene_reference_requires_pinned_content(
    tmp_path, monkeypatch
):
    payload = b"pinned-gene-reference"

    def fake_download(url, path):
        assert url == "https://example.test/reference.gtf.gz"
        path.write_bytes(payload)

    monkeypatch.setattr(
        "app.repository.adapters.icgc.download", fake_download
    )
    reference = {
        "url": "https://example.test/reference.gtf.gz",
        "name": "reference.gtf.gz",
        "size": len(payload),
        "sha256": hashlib.sha256(payload).hexdigest(),
    }
    path = materialize_gene_reference(reference, tmp_path)
    assert path == tmp_path / "reference.gtf.gz"
    assert path.read_bytes() == payload

    with pytest.raises(ValueError, match="changed SHA-256"):
        materialize_gene_reference(
            {**reference, "sha256": "0" * 64}, tmp_path
        )


def test_publication_adapter_reads_named_xlsx_sheet_and_header_row(
    tmp_path,
):
    workbook = tmp_path / "clinical.xlsx"
    spreadsheet_ns = (
        "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    )
    relationship_ns = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    package_relationship_ns = (
        "http://schemas.openxmlformats.org/package/2006/relationships"
    )
    with zipfile.ZipFile(workbook, "w") as archive:
        archive.writestr(
            "xl/sharedStrings.xml",
            (
                f'<sst xmlns="{spreadsheet_ns}">'
                "<si><t>Sample ID</t></si>"
                "<si><t>OS days</t></si>"
                "<si><t>S1</t></si>"
                "</sst>"
            ),
        )
        archive.writestr(
            "xl/workbook.xml",
            (
                f'<workbook xmlns="{spreadsheet_ns}" '
                f'xmlns:r="{relationship_ns}"><sheets>'
                '<sheet name="Summary" sheetId="1" r:id="rId1"/>'
                '<sheet name="Patient data" sheetId="2" r:id="rId2"/>'
                "</sheets></workbook>"
            ),
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            (
                f'<Relationships xmlns="{package_relationship_ns}">'
                '<Relationship Id="rId1" Target="worksheets/sheet1.xml"/>'
                '<Relationship Id="rId2" Target="worksheets/sheet2.xml"/>'
                "</Relationships>"
            ),
        )
        archive.writestr(
            "xl/worksheets/sheet2.xml",
            (
                f'<worksheet xmlns="{spreadsheet_ns}"><sheetData>'
                '<row r="3"><c r="A3" t="s"><v>0</v></c>'
                '<c r="B3" t="s"><v>1</v></c></row>'
                '<row r="4"><c r="A4" t="s"><v>2</v></c>'
                '<c r="B4"><v>365</v></c></row>'
                "</sheetData></worksheet>"
            ),
        )

    assert read_xlsx_sheet(
        workbook,
        sheet_name="Patient data",
        header_row=3,
    ) == [{"Sample ID": "S1", "OS days": "365"}]


def test_publication_adapter_reads_pinned_csv_with_unnamed_id_column(
    tmp_path,
):
    clinical = tmp_path / "clinical.csv"
    clinical.write_text(
        ",status,OS\nS1,1,12.5\nS2,0,18.0\n",
        encoding="utf-8",
    )

    assert read_delimited_records(
        clinical,
        delimiter=",",
        header_name_overrides={"1": "sample_id"},
    ) == [
        {"sample_id": "S1", "status": "1", "OS": "12.5"},
        {"sample_id": "S2", "status": "0", "OS": "18.0"},
    ]


def test_publication_adapter_preserves_curated_records_and_license(
    tmp_path,
):
    records = normalize_curated_records(
        [
            {"ID": "1T", "Death": "Yes", "OS": 14.3},
            {"ID": "2T", "Death": "No", "OS": None},
        ]
    )
    assert records == [
        {"ID": "1T", "Death": "Yes", "OS": "14.3"},
        {"ID": "2T", "Death": "No", "OS": ""},
    ]

    full_text = tmp_path / "article.xml"
    full_text.write_text(
        '<article xmlns:xlink="http://www.w3.org/1999/xlink">'
        '<license xlink:href="https://creativecommons.org/licenses/'
        'by-nc/4.0/">CC BY-NC 4.0</license></article>',
        encoding="utf-8",
    )
    assert (
        detect_creative_commons_license(full_text)
        == "CC-BY-NC-4.0"
    )
    full_text.write_text(
        '<article xmlns:xlink="http://www.w3.org/1999/xlink">'
        '<license xlink:href="https://creativecommons.org/licenses/'
        'by-nc-nd/4.0/">CC BY-NC-ND 4.0</license></article>',
        encoding="utf-8",
    )
    assert (
        detect_creative_commons_license(full_text)
        == "CC-BY-NC-ND-4.0"
    )


def test_publication_expression_layouts_are_explicit_and_deterministic(
    tmp_path,
):
    gencode = tmp_path / "gencode.gtf.gz"
    with gzip.open(gencode, "wt", encoding="utf-8") as handle:
        handle.write(
            "chr1\tHAVANA\tgene\t1\t2\t.\t+\t.\t"
            'gene_id "ENSG000001.4"; gene_name "GENEA";\n'
            "chr1\tHAVANA\ttranscript\t1\t2\t.\t+\t.\t"
            'gene_id "ENSG000001.4"; gene_name "GENEA";\n'
        )
    assert parse_gencode_gene_map(gencode) == {
        "ENSG000001": "GENEA",
        "ENSG000001.4": "GENEA",
    }
    assert selected_sample_indexes(
        ["feature", "az-1", "AZ-2"],
        ["AZ-1", "az-2"],
        case_insensitive=True,
    ) == [1, 2]

    source = tmp_path / "rpkm.txt.gz"
    with gzip.open(
        source, "wt", newline="", encoding="utf-8"
    ) as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(
            [
                "No",
                "Chr",
                "Ensembl",
                "Symbol",
                "Start",
                "Stop",
                "Length",
                "source-a",
                "source-b",
            ]
        )
        writer.writerow(["", "", "", "", "", "", "", "1T", "2T"])
        for index in range(10_000):
            writer.writerow(
                [
                    index,
                    "chr1",
                    f"ENSG{index:011d}",
                    f"GENE{index}",
                    1,
                    2,
                    2,
                    1.0,
                    2.0,
                ]
            )
    assert publication_expression_sample_ids(
        source,
        {
            "layout": "geo_rpkm_two_row_header",
            "delimiter": "\t",
            "metadata_columns": 7,
        },
    ) == ["1T", "2T"]
    output = tmp_path / "selected.tsv"
    summary = materialize_geo_rpkm_matrix(
        source,
        output,
        ["2T", "1T"],
        delimiter="\t",
        metadata_columns=7,
        ensembl_column=2,
        symbol_column=3,
    )
    assert summary["mapped_expression_gene_rows"] == 10_000
    with output.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == [
            "Ensembl_Gene_Id|Hugo_Symbol",
            "2T",
            "1T",
        ]
        assert next(reader) == [
            "ENSG00000000000|GENE0",
            "2.0",
            "1.0",
        ]


def test_publication_xlsx_expression_layout_repairs_declared_features(
    tmp_path,
):
    from openpyxl import Workbook

    source = tmp_path / "expression.xlsx"
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("processed")
    worksheet.append(["notes"])
    worksheet.append([])
    worksheet.append(["genes", "S1", "S2"])
    worksheet.append([datetime(2020, 3, 1), 1.0, 2.0])
    for index in range(9_999):
        worksheet.append([f"GENE{index}", index, index + 1])
    workbook.save(source)

    expression_spec = {
        "layout": "xlsx_matrix",
        "sheet_name": "processed",
        "header_row": 3,
        "feature_column": "genes",
        "sample_start_column": 2,
        "feature_replacements": {
            "2020-03-01T00:00:00": "MARCH1",
        },
    }
    assert publication_expression_sample_ids(
        source,
        expression_spec,
    ) == ["S1", "S2"]

    target = tmp_path / "selected.tsv"
    summary = materialize_xlsx_expression_matrix(
        source,
        target,
        ["S2", "S1"],
        sheet_name="processed",
        header_row=3,
        feature_column="genes",
        sample_start_column=2,
        feature_replacements={
            "2020-03-01T00:00:00": "MARCH1",
        },
        case_insensitive_samples=False,
    )
    assert summary["source_expression_gene_rows"] == 10_000
    assert summary["source_expression_feature_replacements"] == 1
    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["genes", "S2", "S1"]
        assert next(reader) == ["MARCH1", "2", "1"]


def test_geo_xlsx_gene_counts_selects_normalizes_and_tracks_source(
    tmp_path,
):
    from openpyxl import Workbook

    source = tmp_path / "counts.xlsx"
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("counts")
    worksheet.append(["Gene", "Unused", "S1", "S2"])
    worksheet.append(["DUPLICATE", 3, 1, 1])
    worksheet.append(["DUPLICATE", 4, 2, 2])
    for index in range(10_000):
        worksheet.append([f"GENE{index}", index, index + 1, index + 2])
    workbook.save(source)

    target = tmp_path / "log2_cpm.tsv"
    summary = materialize_geo_xlsx_count_matrix(
        source,
        target,
        ["S2", "S1"],
        sheet_name="counts",
        header_row=1,
        feature_column="Gene",
        feature_column_output="Hugo_Symbol",
        sample_start_column=2,
        feature_replacements={},
        case_insensitive_samples=False,
        duplicate_feature_policy="exclude_ambiguous",
        duplicate_feature_case_insensitive=True,
        count_value_type="integer",
    )

    assert summary["source_expression_layout"] == "xlsx_gene_counts"
    assert summary["source_expression_columns"] == 3
    assert summary["selected_expression_columns"] == 2
    assert summary["source_expression_gene_rows"] == 10_000
    assert summary["normalized_expression_gene_rows"] == 10_000
    assert summary["source_expression_duplicate_feature_symbols"] == 1
    assert summary["source_expression_excluded_duplicate_feature_rows"] == 2
    assert summary["expression_normalization"] == "log2(CPM + 1)"
    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["Hugo_Symbol", "S2", "S1"]
        first = next(reader)
        assert first[0] == "GENE0"
        assert float(first[1]) > float(first[2])


def test_publication_xlsx_expression_can_drop_date_formatted_features(
    tmp_path,
):
    from openpyxl import Workbook

    source = tmp_path / "expression.xlsx"
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("processed")
    worksheet.append(["genes", "S1", "S2"])
    worksheet.append([datetime(2020, 3, 1), 1.0, 2.0])
    for index in range(10_000):
        worksheet.append([f"GENE{index}", index, index + 1])
    workbook.save(source)

    target = tmp_path / "selected.tsv"
    summary = materialize_xlsx_expression_matrix(
        source,
        target,
        ["S1", "S2"],
        sheet_name="processed",
        header_row=1,
        feature_column="genes",
        sample_start_column=2,
        feature_replacements={},
        case_insensitive_samples=False,
        drop_date_formatted_features=True,
    )

    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["genes", "S1", "S2"]
        assert next(reader) == ["GENE0", "0", "1"]
    assert summary["source_expression_gene_rows"] == 10_000
    assert (
        summary["source_expression_dropped_date_formatted_features"]
        == 1
    )


def test_publication_xlsb_tables_drop_ambiguous_numeric_features(
    tmp_path,
    monkeypatch,
):
    class FakeCell:
        def __init__(self, value):
            self.v = value

    class FakeSheet:
        def __init__(self, rows):
            self._rows = rows

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def rows(self):
            for row in self._rows:
                yield [FakeCell(value) for value in row]

    class FakeWorkbook:
        def __init__(self, sheets):
            self._sheets = sheets
            self.sheets = list(sheets)

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def get_sheet(self, name):
            return FakeSheet(self._sheets[name])

    clinical_rows = [
        ["Clinical table", None, None],
        ["ID", "PFS", "CNSR"],
        ["P1", 4.5, 0.0],
        ["P2", 8.0, 1.0],
    ]
    expression_rows = [
        ["Expression table", None, None],
        ["HUGO", "P1", "P2"],
        [44166.0, 1.0, 2.0],
        *[
            [f"GENE{index}", float(index), float(index + 1)]
            for index in range(10_000)
        ],
    ]
    sheets = {
        "clinical": clinical_rows,
        "expression": expression_rows,
    }
    monkeypatch.setattr(
        "app.repository.adapters.pmc.open_workbook",
        lambda _path: FakeWorkbook(sheets),
    )

    source = tmp_path / "source.xlsb"
    source.write_bytes(b"fake xlsb handled by the patched reader")
    assert read_xlsb_sheet(
        source,
        sheet_name="clinical",
        header_row=2,
    ) == [
        {"ID": "P1", "PFS": "4.5", "CNSR": "0"},
        {"ID": "P2", "PFS": "8", "CNSR": "1"},
    ]

    target = tmp_path / "selected.tsv"
    summary = materialize_xlsb_expression_matrix(
        source,
        target,
        ["P2", "P1"],
        sheet_name="expression",
        header_row=2,
        feature_column="HUGO",
        sample_start_column=2,
        feature_replacements={},
        case_insensitive_samples=False,
        drop_non_string_features=True,
        duplicate_feature_policy="error",
        duplicate_feature_case_insensitive=True,
    )

    assert summary["source_expression_gene_rows"] == 10_000
    assert (
        summary["source_expression_dropped_non_string_features"]
        == 1
    )
    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["HUGO", "P2", "P1"]
        assert next(reader) == ["GENE0", "1", "0"]


def test_publication_xlsx_expression_drops_only_fully_missing_features(
    tmp_path,
):
    from openpyxl import Workbook

    source = tmp_path / "expression.xlsx"
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("processed")
    worksheet.append(["genes", "S1", "S2"])
    worksheet.append(["ALL_NA", "NA", "NA"])
    for index in range(10_000):
        worksheet.append([f"GENE{index}", index, index + 1])
    workbook.save(source)

    target = tmp_path / "selected.tsv"
    summary = materialize_xlsx_expression_matrix(
        source,
        target,
        ["S1", "S2"],
        sheet_name="processed",
        header_row=1,
        feature_column="genes",
        sample_start_column=2,
        feature_replacements={},
        case_insensitive_samples=False,
        drop_all_missing_features=True,
        missing_value_tokens=["NA"],
    )
    assert summary["source_expression_gene_rows"] == 10_000
    assert summary[
        "source_expression_dropped_all_missing_features"
    ] == 1
    assert summary["source_expression_missing_value_tokens"] == ["na"]


def test_publication_xlsx_expression_rejects_partially_missing_features(
    tmp_path,
):
    from openpyxl import Workbook

    source = tmp_path / "expression.xlsx"
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("processed")
    worksheet.append(["genes", "S1", "S2"])
    worksheet.append(["PARTIAL_NA", "NA", 1.0])
    workbook.save(source)

    with pytest.raises(ValueError, match="missing value"):
        materialize_xlsx_expression_matrix(
            source,
            tmp_path / "selected.tsv",
            ["S1", "S2"],
            sheet_name="processed",
            header_row=1,
            feature_column="genes",
            sample_start_column=2,
            feature_replacements={},
            case_insensitive_samples=False,
            drop_all_missing_features=True,
            missing_value_tokens=["NA"],
        )


def test_publication_xlsx_expression_can_drop_any_incomplete_feature(
    tmp_path,
):
    from openpyxl import Workbook

    source = tmp_path / "expression.xlsx"
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("processed")
    worksheet.append(["genes", "S1", "S2"])
    worksheet.append(["ALL_NA", "NA", "NA"])
    worksheet.append(["PARTIAL_NA", "NA", 1.0])
    for index in range(10_000):
        worksheet.append([f"GENE{index}", index, index + 1])
    workbook.save(source)

    target = tmp_path / "selected.tsv"
    summary = materialize_xlsx_expression_matrix(
        source,
        target,
        ["S1", "S2"],
        sheet_name="processed",
        header_row=1,
        feature_column="genes",
        sample_start_column=2,
        feature_replacements={},
        case_insensitive_samples=False,
        drop_any_missing_features=True,
        missing_value_tokens=["NA"],
    )
    assert summary["source_expression_gene_rows"] == 10_000
    assert summary[
        "source_expression_dropped_all_missing_features"
    ] == 1
    assert summary[
        "source_expression_dropped_any_missing_features"
    ] == 2
    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["genes", "S1", "S2"]
        assert next(reader) == ["GENE0", "0", "1"]


def test_publication_xlsx_expression_excludes_ambiguous_features(
    tmp_path,
):
    from openpyxl import Workbook

    source = tmp_path / "expression.xlsx"
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("processed")
    worksheet.append(["genes", "S1", "S2"])
    worksheet.append(["DUPLICATE", 1.0, 2.0])
    worksheet.append(["duplicate", 3.0, 4.0])
    for index in range(10_000):
        worksheet.append([f"GENE{index}", index, index + 1])
    workbook.save(source)

    target = tmp_path / "selected.tsv"
    summary = materialize_xlsx_expression_matrix(
        source,
        target,
        ["S1", "S2"],
        sheet_name="processed",
        header_row=1,
        feature_column="genes",
        sample_start_column=2,
        feature_replacements={},
        case_insensitive_samples=False,
        duplicate_feature_policy="exclude_ambiguous",
        duplicate_feature_case_insensitive=True,
    )
    assert summary["source_expression_gene_rows"] == 10_000
    assert summary[
        "source_expression_duplicate_feature_symbols"
    ] == 1
    assert summary["source_expression_duplicate_feature_rows"] == 2
    assert summary[
        "source_expression_excluded_duplicate_feature_rows"
    ] == 2
    with target.open(encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        assert next(reader) == ["genes", "S1", "S2"]
        assert next(reader) == ["GENE0", "0", "1"]


def test_europe_pmc_supplement_pin_requires_exact_size_and_sha256(
    tmp_path,
):
    source = tmp_path / "supplement.xlsx"
    source.write_bytes(b"pinned supplement")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()

    verify_pinned_supplement(
        source,
        expected_size=source.stat().st_size,
        expected_sha256=digest,
        label="test supplement",
    )
    with pytest.raises(ValueError, match="changed size"):
        verify_pinned_supplement(
            source,
            expected_size=source.stat().st_size + 1,
            expected_sha256=digest,
            label="test supplement",
        )
    with pytest.raises(ValueError, match="changed SHA-256"):
        verify_pinned_supplement(
            source,
            expected_size=source.stat().st_size,
            expected_sha256="0" * 64,
            label="test supplement",
        )


def test_pmc_proof_of_work_challenge_is_parsed_and_solved():
    challenge_html = """
    <script type="module">
      const POW_CHALLENGE = "fixed:test-challenge"
      const POW_DIFFICULTY = "2"
      const POW_COOKIE_NAME = "cloudpmc-viewer-pow"
    </script>
    """
    challenge, difficulty, cookie_name = parse_pmc_pow_challenge(
        challenge_html
    )
    nonce = solve_pmc_pow_nonce(challenge, difficulty)

    assert challenge == "fixed:test-challenge"
    assert difficulty == 2
    assert cookie_name == "cloudpmc-viewer-pow"
    assert hashlib.sha256(
        f"{challenge}{nonce}".encode("utf-8")
    ).hexdigest().startswith("00")


def test_cbioportal_source_and_sample_rules_are_declarative():
    assert resolve_source_file_spec(
        "public/example",
        "data_mrna_seq.txt",
    ) == ("public/example/data_mrna_seq.txt", "data_mrna_seq.txt")
    assert resolve_source_file_spec(
        "public/example",
        {
            "path": "README.md",
            "scope": "repository",
            "target_name": "DATAHUB_LICENSE.md",
        },
    ) == ("README.md", "DATAHUB_LICENSE.md")
    with pytest.raises(ValueError, match="Invalid source"):
        resolve_source_file_spec("public/example", "../LICENSE")

    patient = {"DISEASE": "AML"}
    initial = {
        "STAGE": "Initial Diagnosis",
        "SPECIMEN": "Bone Marrow",
    }
    relapse = {
        "STAGE": "Relapse",
        "SPECIMEN": "Peripheral Blood",
    }
    rules = [
        {
            "field": "patient.DISEASE",
            "include": ["AML"],
            "required": True,
        },
        {
            "field": "sample.STAGE",
            "include": ["Initial Diagnosis"],
        },
    ]
    assert sample_passes_eligibility(patient, initial, rules)
    assert not sample_passes_eligibility(patient, relapse, rules)
    sample_spec = {
        "selection_rank_rules": [
            {
                "field": "sample.SPECIMEN",
                "order": ["Bone Marrow", "Peripheral Blood"],
            }
        ]
    }
    assert sample_selection_rank(patient, initial, sample_spec) == 0
    assert sample_selection_rank(patient, relapse, sample_spec) == 1


def test_cbioportal_endpoint_supports_explicit_status_sets():
    endpoint = {
        "time_column": "OS_DAYS",
        "event_column": "STATUS",
        "time_unit": "days",
        "event_values": [
            "DOD",
            "DOC",
            "Deceased - Of Unknown Cause",
        ],
        "censor_values": ["NED", "AWD"],
    }
    assert parse_endpoint(
        {"OS_DAYS": "173", "STATUS": "DOD"}, endpoint
    ) == {
        "time_days": "173",
        "event": "1",
        "raw_time": "173",
        "raw_event": "DOD",
    }
    assert parse_endpoint(
        {"OS_DAYS": "1307", "STATUS": "awd"}, endpoint
    ) == {
        "time_days": "1307",
        "event": "0",
        "raw_time": "1307",
        "raw_event": "awd",
    }
    assert (
        parse_endpoint(
            {"OS_DAYS": "500", "STATUS": "Unknown"}, endpoint
        )
        is None
    )


def test_cbioportal_expression_feature_id_joins_hugo_and_entrez():
    assert expression_feature_id(
        ["KRAS", "3845", "7.2", "8.1"],
        feature_id_type="provided_hugo_entrez",
        feature_column_count=2,
    ) == "KRAS|3845"
    assert expression_feature_id(
        ["KRAS", "7.2", "8.1"],
        feature_id_type="provided_hugo",
        feature_column_count=1,
    ) == "KRAS"


def test_cbioportal_endpoint_rejects_ambiguous_status_sets():
    endpoint = {
        "time_column": "OS_DAYS",
        "event_column": "STATUS",
        "time_unit": "days",
        "event_values": ["DOD"],
        "censor_values": ["DOD"],
    }
    with pytest.raises(ValueError, match="overlap"):
        parse_endpoint(
            {"OS_DAYS": "173", "STATUS": "DOD"}, endpoint
        )


def test_cbioportal_discovery_keeps_screening_separate_from_promotion(
    tmp_path,
):
    (tmp_path / "studies").mkdir()
    (tmp_path / "cancer_types.json").write_text(
        json.dumps(
            {
                "cancer_types": [
                    {
                        "code": "BLCA",
                        "tcga_cohort": "TCGA-BLCA",
                        "name": "Bladder cancer",
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    (tmp_path / "coverage.json").write_text(
        json.dumps(
            {
                "cancers": [
                    {
                        "code": "BLCA",
                        "status": "available",
                        "candidates": [
                            {
                                "accession": "blca_external",
                                "review_status": "promoted",
                            }
                        ],
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    (tmp_path / "studies" / "blca.json").write_text(
        json.dumps(
            {
                "dataset": {
                    "id": "fixture-blca",
                    "source_accession": "blca_external",
                }
            }
        ),
        encoding="utf-8",
    )

    def requester(url, **_kwargs):
        if "/studies?" in url:
            return [
                {
                    "studyId": "blca_external",
                    "cancerTypeId": "blca",
                    "name": "Independent bladder cohort",
                    "description": "Bulk RNA cohort",
                    "allSampleCount": 20,
                    "mrnaRnaSeqSampleCount": 20,
                },
                {
                    "studyId": "blca_tcga",
                    "cancerTypeId": "blca",
                    "name": "TCGA",
                    "description": "TCGA cohort",
                    "allSampleCount": 400,
                    "mrnaRnaSeqSampleCount": 400,
                },
            ]
        if "/molecular-profiles?" in url:
            return [
                {
                    "studyId": "blca_external",
                    "molecularProfileId": "blca_external_rna_seq_mrna",
                    "molecularAlterationType": "MRNA_EXPRESSION",
                    "datatype": "CONTINUOUS",
                    "name": "mRNA expression (RPKM)",
                }
            ]
        if url.endswith("/blca_external/clinical-attributes"):
            return [
                {
                    "clinicalAttributeId": "OS_STATUS",
                    "patientAttribute": True,
                },
                {
                    "clinicalAttributeId": "OS_MONTHS",
                    "patientAttribute": True,
                },
            ]
        raise AssertionError(url)

    result = discover_cbioportal_candidates(
        tmp_path,
        requester=requester,
    )
    assert result["summary"]["candidate_studies"] == 1
    candidate = result["cancers"][0]["candidates"][0]
    assert candidate["study_id"] == "blca_external"
    assert candidate["review_status"] == "promoted"
    assert candidate["endpoint_pairs"][0]["endpoint"] == "OS"
    assert result["schema_version"] == "tcga-trace-cbioportal-discovery-v2"
    assert result["cancers"][0]["rejection_reasons"] == {
        "disallowed_accession": 1
    }
    assert result["cancers"][0]["rejections"][0]["study_id"] == "blca_tcga"


def test_cbioportal_rna_sample_lists_recover_missing_summary_count():
    assert _rna_seq_sample_count(
        [
            {
                "category": "all_cases_with_mrna_rnaseq_data",
                "sampleIds": ["S1", "S2"],
            },
            {
                "category": "all_cases_with_mrna_rnaseq_data",
                "sampleIds": ["S2", "S3"],
            },
            {
                "category": "all_cases_with_mrna_array_data",
                "sampleIds": ["A1"],
            },
        ]
    ) == 3


def test_external_manifests_do_not_replace_tcga_summary_manifest():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    with factory() as db:
        db.add_all(
            [
                DataManifest(
                    source_id="tcga_cdr",
                    status="ready",
                    manifest_hash="tcga-hash",
                    created_at=now,
                ),
                DataManifest(
                    source_id="external:fixture",
                    status="ready",
                    manifest_hash="external-hash",
                    created_at=now + timedelta(seconds=1),
                ),
            ]
        )
        db.commit()
        assert latest_data_manifest(db).manifest_hash == "tcga-hash"


def test_select_metaprism_samples_requires_unique_expression_linkage():
    cohort_rows = [
        {
            "sampleId": f"source-{index:02d}",
            "patientId": f"patient-{index:02d}",
            "value": "THCA - Not_TCGA",
        }
        for index in range(10)
    ]
    cohort_rows.append(
        {
            "sampleId": "other-source",
            "patientId": "other-patient",
            "value": "LUAD",
        }
    )
    rna_rows = [
        {
            "sampleId": f"source-{index:02d}",
            "value": f"rna-{index:02d}",
        }
        for index in range(10)
    ]

    selected = select_metaprism_samples(
        cohort_rows,
        rna_rows,
        cohort_values={"THCA - Not_TCGA"},
        expression_sample_ids={
            f"rna-{index:02d}" for index in range(10)
        },
    )

    assert len(selected) == 10
    assert selected[0] == {
        "source_sample_id": "source-00",
        "expression_sample_id": "rna-00",
        "patient_id": "patient-00",
    }
    with pytest.raises(ValueError, match="Fewer than"):
        select_metaprism_samples(
            cohort_rows[:9],
            rna_rows[:9],
            cohort_values={"THCA - Not_TCGA"},
            expression_sample_ids={
                f"rna-{index:02d}" for index in range(9)
            },
        )


def test_materialize_metaprism_expression_selects_and_maps_samples(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "expression.tsv.gz"
    with gzip.open(source, "wt", newline="", encoding="utf-8") as handle:
        writer = csv.writer(
            handle, delimiter="\t", lineterminator="\n"
        )
        writer.writerow(["ensembl_gene_id", "rna-1", "rna-2", "rna-3"])
        writer.writerow(["ENSG000001.1", "1", "2", "3"])
        writer.writerow(["ENSG000002.2", "4", "5", "6"])
        writer.writerow(["ENSG000003.3", "7", "8", "9"])
    reference = tmp_path / "reference.gtf.gz"
    reference.write_bytes(b"fixture")
    monkeypatch.setattr(
        metaprism_adapter,
        "parse_ensembl_gene_map",
        lambda _path: {
            "ENSG000001.1": "GENE1",
            "ENSG000002.2": "GENE2",
        },
    )
    target = tmp_path / "selected.tsv.gz"

    metadata = materialize_metaprism_expression(
        source,
        target,
        reference,
        selected_sample_ids=["rna-3", "rna-1"],
        minimum_mapped_genes=2,
    )

    with gzip.open(target, "rt", newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle, delimiter="\t"))
    assert rows == [
        [
            "Ensembl_Gene_Id|Hugo_Symbol",
            "rna-3",
            "rna-1",
        ],
        ["ENSG000001.1|GENE1", "3", "1"],
        ["ENSG000002.2|GENE2", "6", "4"],
    ]
    assert metadata == {
        "source_expression_gene_rows": 3,
        "mapped_expression_gene_rows": 2,
        "unmapped_expression_gene_rows": 1,
        "selected_expression_sample_ids": ["rna-3", "rna-1"],
    }
