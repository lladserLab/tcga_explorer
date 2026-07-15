#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


BLUE = RGBColor(31, 119, 180)
RED = RGBColor(214, 39, 40)
INK = RGBColor(31, 36, 42)
MUTED = RGBColor(91, 103, 112)
LIGHT = RGBColor(238, 241, 244)
RULE = RGBColor(195, 204, 212)
WHITE = RGBColor(255, 255, 255)


SIGNATURE_IDS = [
    "Trm",
    "Tcirc",
    "HLA-DR",
    "HLA-DP",
    "HLA-DQ",
    "CXCL10_CXCR3",
    "CCL3_CCL4_CCR5",
    "CX3CL1_CX3CR1",
]

GROUPS = [
    ("Firmas principales", "signatures", SIGNATURE_IDS),
    ("Genes CD4 memoria / reactividad / citotoxicidad", "individual_genes", ["CD4", "CD69", "CXCR6", "ITGA1", "CXCL13", "GZMB", "CX3CR1"]),
    ("Genes HLA-DR", "individual_genes", ["HLA-DRA", "HLA-DRB1", "HLA-DRB5", "HLA-DRB6", "HLA-DRB9"]),
    ("Genes HLA-DP", "individual_genes", ["HLA-DPA1", "HLA-DPA2", "HLA-DPA3", "HLA-DPB1", "HLA-DPB2"]),
    ("Genes HLA-DQ", "individual_genes", ["HLA-DQA1", "HLA-DQA2", "HLA-DQB1", "HLA-DQB2"]),
    ("Genes ejes quimioquina / receptor", "individual_genes", ["CXCL10", "CXCR3", "CCL3", "CCL4", "CCR5", "CX3CL1"]),
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build scientific PPTX from TCGA-KIRC survival results.")
    parser.add_argument("--results-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def add_textbox(slide, text: str, x: float, y: float, w: float, h: float, size: int = 14, bold: bool = False, color: RGBColor = INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    add_textbox(slide, title, 0.45, 0.18, 12.45, 0.38, size=22, bold=True)
    if subtitle:
        add_textbox(slide, subtitle, 0.47, 0.58, 12.3, 0.26, size=9, color=MUTED)
    line = slide.shapes.add_shape(1, Inches(0.45), Inches(0.88), Inches(12.45), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = RULE
    line.line.color.rgb = RULE


def add_footer(slide, index: int, total: int) -> None:
    add_textbox(slide, "TCGA-KIRC OS | log2(TPM + 1) | High rojo / Low azul", 0.45, 7.18, 8.8, 0.18, size=7, color=MUTED)
    add_textbox(slide, f"{index}/{total}", 12.05, 7.18, 0.85, 0.18, size=7, color=MUTED, align=PP_ALIGN.RIGHT)


def add_chip(slide, text: str, x: float, y: float, w: float, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(0.24))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    shape.text = text
    frame = shape.text_frame
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.name = "Arial"
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = WHITE


def add_image(slide, path: Path, x: float, y: float, w: float | None = None, h: float | None = None) -> None:
    if w is not None and h is not None:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    elif w is not None:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    elif h is not None:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(h))
    else:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y))


def add_stat_line(slide, row: dict[str, str], x: float, y: float, w: float) -> None:
    text = (
        f"log-rank {pvalue(row.get('logrank_p_value'))}  |  "
        f"Cox stage+grade HR {fmt(row.get('stage_grade_adjusted_hr'))}, {pvalue(row.get('stage_grade_adjusted_p_value'))}"
    )
    add_textbox(slide, text, x, y, w, 0.22, size=8, color=MUTED)


def pvalue(value: Any) -> str:
    number = to_float(value)
    if number is None:
        return "p = NA"
    if number < 0.001:
        return "p < 0.001"
    return f"p = {number:.3f}"


def fmt(value: Any, digits: int = 2) -> str:
    number = to_float(value)
    if number is None:
        return "NA"
    return f"{number:.{digits}f}"


def to_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def image_paths(results_dir: Path, row: dict[str, str]) -> tuple[Path, Path]:
    root = results_dir / ("signatures" if row["analysis_type"] == "signature" else "individual_genes") / row["analysis_id"]
    return root / f"{row['analysis_id']}.km.png", root / f"{row['analysis_id']}.cox_forest.png"


def add_analysis_row(slide, results_dir: Path, row: dict[str, str], y: float) -> None:
    km_path, cox_path = image_paths(results_dir, row)
    title = row["label"]
    genes = row.get("genes", "").replace(";", ", ")
    add_textbox(slide, title, 0.45, y, 3.0, 0.26, size=12, bold=True)
    add_textbox(slide, genes, 3.0, y + 0.02, 6.7, 0.2, size=7, color=MUTED)
    add_stat_line(slide, row, 9.2, y + 0.02, 3.65)
    add_image(slide, km_path, 0.47, y + 0.36, w=2.55)
    add_image(slide, cox_path, 3.22, y + 0.47, w=6.25)
    add_chip(slide, "Low", 9.86, y + 0.72, 0.56, BLUE)
    add_chip(slide, "High", 10.52, y + 0.72, 0.62, RED)
    add_textbox(slide, "KM: maxstat\nCox: High vs Low\nModelos: uni, +estadio, +grado, +ambos", 9.85, y + 1.08, 2.75, 0.74, size=8, color=MUTED)


def add_title_slide(prs: Presentation, sample_n: str, events: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_textbox(slide, "TCGA-KIRC: Kaplan-Meier y Cox", 0.75, 1.2, 11.8, 0.52, size=28, bold=True)
    add_textbox(slide, "Firmas y genes del documento Tesis_MF", 0.78, 1.78, 11.2, 0.35, size=15, color=MUTED)
    add_textbox(slide, f"OS | n = {sample_n} pacientes | eventos = {events} | corte KM maxstat", 0.78, 2.35, 10.8, 0.3, size=13)
    add_chip(slide, "Low", 0.82, 3.02, 0.75, BLUE)
    add_chip(slide, "High", 1.72, 3.02, 0.85, RED)
    add_textbox(slide, "Cox reportado como High vs Low: univariado, ajustado por estadio, ajustado por grado y ajustado por ambos.", 0.78, 3.55, 11.25, 0.5, size=11, color=MUTED)
    add_textbox(slide, "Archivo generado para revisión científica interna.", 0.78, 6.95, 10.0, 0.22, size=8, color=MUTED)


def add_methods_slide(prs: Presentation, skipped: list[dict[str, str]]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Diseño del análisis", "Criterios usados para todas las curvas y modelos")
    bullets = [
        "Cohorte: TCGA-KIRC, una muestra RNA-seq tumoral primaria por paciente.",
        "Endpoint: sobrevida global; Dead = evento, Alive = censura.",
        "Expresión: log2(TPM + 1); firma = promedio de genes incluidos.",
        "Kaplan-Meier: corte maxstat, curvas cuadradas, sin IC, sin grid y sin tabla inferior.",
        "Cox: High vs Low con cuatro modelos: univariado, +estadio, +grado, +estadio+grado.",
        "HLA-BQB2 del documento se trató como HLA-DQB2 por ausencia de HLA-BQB2 en la matriz.",
    ]
    y = 1.25
    for item in bullets:
        add_textbox(slide, f"- {item}", 0.82, y, 11.7, 0.32, size=13)
        y += 0.52
    if skipped:
        add_textbox(slide, "Análisis omitido", 0.82, 5.35, 11.4, 0.3, size=14, bold=True)
        skip_text = "; ".join(f"{row['label']}: maxstat sin corte válido" for row in skipped)
        add_textbox(slide, skip_text, 0.82, 5.78, 11.3, 0.45, size=11, color=MUTED)


def add_signature_summary(prs: Presentation, rows: list[dict[str, str]]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Resumen de firmas", "Lectura rápida: KM y Cox ajustado por estadio+grado")
    headers = ["Firma", "Genes", "KM", "HR ajustado", "p ajustado"]
    widths = [1.7, 5.15, 1.35, 1.45, 1.35]
    x0, y0 = 0.55, 1.22
    x = x0
    for header, width in zip(headers, widths):
        add_textbox(slide, header, x, y0, width, 0.25, size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        shape = slide.shapes[-1]
        shape.fill.solid()
        shape.fill.fore_color.rgb = INK
        x += width + 0.08
    y = y0 + 0.42
    for row in rows:
        x = x0
        fill = LIGHT if int((y - y0) / 0.48) % 2 == 0 else WHITE
        values = [
            row["analysis_id"],
            row["genes"].replace(";", ", "),
            pvalue(row.get("logrank_p_value")),
            fmt(row.get("stage_grade_adjusted_hr")),
            pvalue(row.get("stage_grade_adjusted_p_value")),
        ]
        for value, width in zip(values, widths):
            box = add_textbox(slide, value, x, y, width, 0.32, size=8, color=INK, align=PP_ALIGN.CENTER if width < 2 else PP_ALIGN.LEFT)
            box.fill.solid()
            box.fill.fore_color.rgb = fill
            box.line.color.rgb = WHITE
            x += width + 0.08
        y += 0.48
    add_textbox(slide, "Nota: HR < 1 indica menor riesgo en High vs Low despues del ajuste.", 0.62, 6.55, 11.6, 0.25, size=8, color=MUTED)


def add_group_slides(prs: Presentation, results_dir: Path, title: str, rows: list[dict[str, str]]) -> None:
    for page, offset in enumerate(range(0, len(rows), 2), start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        suffix = f" ({page}/{math.ceil(len(rows) / 2)})" if len(rows) > 2 else ""
        add_title(slide, f"{title}{suffix}", "Cada fila muestra KM y Cox univariado/multivariado del mismo marcador")
        add_analysis_row(slide, results_dir, rows[offset], 1.08)
        if offset + 1 < len(rows):
            add_analysis_row(slide, results_dir, rows[offset + 1], 4.06)


def main() -> int:
    args = parse_args()
    results_dir = args.results_dir
    summary_rows = read_csv(results_dir / "tables" / "summary_all_results.csv")
    skipped = read_csv(results_dir / "tables" / "missing_or_skipped.csv")
    by_id = {row["analysis_id"]: row for row in summary_rows}
    signature_rows = [by_id[analysis_id] for analysis_id in SIGNATURE_IDS if analysis_id in by_id]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    sample_n = signature_rows[0].get("n_patients", "529") if signature_rows else "529"
    events = signature_rows[0].get("n_events", "173") if signature_rows else "173"
    add_title_slide(prs, sample_n, events)
    add_methods_slide(prs, skipped)
    add_signature_summary(prs, signature_rows)

    for title, _, ids in GROUPS:
        rows = [by_id[item] for item in ids if item in by_id]
        if rows:
            add_group_slides(prs, results_dir, title, rows)

    total = len(prs.slides)
    for index, slide in enumerate(prs.slides, start=1):
        if index > 1:
            add_footer(slide, index, total)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
