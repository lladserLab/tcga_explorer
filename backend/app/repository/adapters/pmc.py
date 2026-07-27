from __future__ import annotations

import csv
from datetime import date, datetime
import gzip
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
from typing import Any
import urllib.parse
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from app.repository.adapters.cbioportal import download, write_tsv
from app.repository.storage import sha256_file


EUROPE_PMC_ADAPTER_VERSION = "europe_pmc_publication_v3"
DEFAULT_EUROPE_PMC_API = (
    "https://www.ebi.ac.uk/europepmc/webservices/rest"
)


def materialize_europe_pmc_source(
    spec: dict[str, Any],
    source_dir: Path,
) -> tuple[dict[str, Path], str, str | None, dict[str, Any]]:
    source_spec = spec["source"]
    pmcid = str(source_spec.get("pmcid") or "").strip().upper()
    api = str(
        source_spec.get("api_base") or DEFAULT_EUROPE_PMC_API
    ).rstrip("/")
    clinical_member = str(
        source_spec.get("clinical_member") or ""
    ).strip()
    expression_member = str(
        source_spec.get("expression_member") or ""
    ).strip()
    expression_url = str(
        source_spec.get("expression_url") or ""
    ).strip()
    expression_reference_url = str(
        source_spec.get("expression_reference_url") or ""
    ).strip()
    if (
        not re.fullmatch(r"PMC[0-9]+", pmcid)
        or not clinical_member
        or bool(expression_member) == bool(expression_url)
    ):
        raise ValueError(
            "Europe PMC sources require PMCID, clinical_member and exactly "
            "one of expression_member or expression_url."
        )

    full_text_url = f"{api}/{pmcid}/fullTextXML"
    supplement_url = f"{api}/{pmcid}/supplementaryFiles"
    full_text_path = source_dir / f"{pmcid}_full_text.xml"
    temporary_zip = source_dir / ".supplementary_files.zip"
    download(full_text_url, full_text_path)
    source_license = detect_creative_commons_license(full_text_path)
    download(supplement_url, temporary_zip)

    clinical_path = source_dir / Path(clinical_member).name
    with zipfile.ZipFile(temporary_zip) as archive:
        member_names = set(archive.namelist())
        required_members = {clinical_member}
        if expression_member:
            required_members.add(expression_member)
        missing = required_members - member_names
        if missing:
            raise ValueError(
                "Europe PMC supplementary archive is missing required "
                f"members: {sorted(missing)}"
            )
        extract_zip_member(archive, clinical_member, clinical_path)
        if expression_member:
            raw_expression_path = (
                source_dir / Path(expression_member).name
            )
            extract_zip_member(
                archive, expression_member, raw_expression_path
            )
    temporary_zip.unlink()
    if expression_url:
        raw_expression_path = (
            source_dir / download_filename(expression_url)
        )
        download(expression_url, raw_expression_path)
    expression_reference_path: Path | None = None
    if expression_reference_url:
        expression_reference_path = (
            source_dir / download_filename(expression_reference_url)
        )
        download(
            expression_reference_url,
            expression_reference_path,
        )

    clinical_spec = source_spec.get("clinical") or {}
    curated_records = clinical_spec.get("records")
    if curated_records is not None:
        if not isinstance(curated_records, list) or not curated_records:
            raise ValueError(
                "Europe PMC clinical records must be a non-empty list."
            )
        clinical_rows = normalize_curated_records(curated_records)
    else:
        if clinical_path.suffix.lower() != ".xlsx":
            raise ValueError(
                "A non-XLSX clinical supplement requires explicit curated "
                "clinical records in the study specification."
            )
        clinical_rows = read_xlsx_sheet(
            clinical_path,
            sheet_name=(
                str(clinical_spec.get("sheet_name") or "").strip()
                or None
            ),
            header_row=int(clinical_spec.get("header_row") or 1),
        )
    sample_id_column = str(
        clinical_spec.get("sample_id_column") or ""
    ).strip()
    patient_id_column = str(
        clinical_spec.get("patient_id_column") or sample_id_column
    ).strip()
    if not sample_id_column or not patient_id_column:
        raise ValueError(
            "Europe PMC clinical sample_id_column is required."
        )
    eligibility_rules = clinical_spec.get("eligibility_rules") or []
    eligible_rows = [
        row
        for row in clinical_rows
        if row_passes_rules(row, eligibility_rules)
    ]
    excluded_without_expression = 0
    if clinical_spec.get("require_expression_sample"):
        expression_sample_ids = {
            value.casefold()
            for value in publication_expression_sample_ids(
                raw_expression_path,
                source_spec.get("expression") or {
                    "layout": "simple",
                    "delimiter": str(
                        source_spec.get("expression_delimiter")
                        or "\t"
                    ),
                },
            )
        }
        rows_before_expression_filter = len(eligible_rows)
        eligible_rows = [
            row
            for row in eligible_rows
            if str(row.get(sample_id_column) or "")
            .strip()
            .casefold()
            in expression_sample_ids
        ]
        excluded_without_expression = (
            rows_before_expression_filter - len(eligible_rows)
        )
    eligible_rows.sort(
        key=lambda row: (
            str(row.get(patient_id_column) or ""),
            str(row.get(sample_id_column) or ""),
        )
    )
    if len(eligible_rows) < 10:
        raise ValueError(
            "Europe PMC source has fewer than 10 eligible clinical rows."
        )
    sample_ids = [
        str(row.get(sample_id_column) or "").strip()
        for row in eligible_rows
    ]
    patient_ids = [
        str(row.get(patient_id_column) or "").strip()
        for row in eligible_rows
    ]
    if (
        any(not value for value in sample_ids)
        or len(sample_ids) != len(set(sample_ids))
        or any(not value for value in patient_ids)
        or len(patient_ids) != len(set(patient_ids))
    ):
        raise ValueError(
            "Europe PMC eligible sample and patient IDs must be present "
            "and unique."
        )

    sample_type = str(
        clinical_spec.get("sample_type") or "Primary tumor"
    )
    patient_rows = []
    sample_rows = []
    for patient_id, sample_id, row in zip(
        patient_ids, sample_ids, eligible_rows, strict=True
    ):
        patient_rows.append(
            {
                "PATIENT_ID": patient_id,
                **row,
                "PMC_CLINICAL_METADATA_JSON": canonical_json(row),
            }
        )
        sample_rows.append(
            {
                "SAMPLE_ID": sample_id,
                "PATIENT_ID": patient_id,
                "SAMPLE_TYPE": sample_type,
                "PMC_CLINICAL_METADATA_JSON": canonical_json(row),
            }
        )

    patient_table = source_dir / "data_clinical_patient.txt"
    sample_table = source_dir / "data_clinical_sample.txt"
    patient_fields = [
        "PATIENT_ID",
        *[
            field
            for field in clinical_rows[0]
            if field != "PATIENT_ID"
        ],
        "PMC_CLINICAL_METADATA_JSON",
    ]
    write_tsv(patient_table, patient_rows, patient_fields)
    write_tsv(
        sample_table,
        sample_rows,
        [
            "SAMPLE_ID",
            "PATIENT_ID",
            "SAMPLE_TYPE",
            "PMC_CLINICAL_METADATA_JSON",
        ],
    )

    filtered_expression_path = (
        source_dir / "data_expression_selected.tsv"
    )
    expression_summary = materialize_publication_expression(
        raw_expression_path,
        filtered_expression_path,
        sample_ids,
        source_spec.get("expression") or {
            "layout": "simple",
            "delimiter": str(
                source_spec.get("expression_delimiter") or "\t"
            ),
        },
        expression_reference_path=expression_reference_path,
    )
    raw_source_hashes = {
        "full_text_xml": sha256_file(full_text_path),
        "clinical_supplement": sha256_file(clinical_path),
        "expression_source": sha256_file(raw_expression_path),
    }
    if expression_reference_path is not None:
        raw_source_hashes["expression_reference"] = sha256_file(
            expression_reference_path
        )
    source_snapshot = hashlib.sha256(
        canonical_json(raw_source_hashes).encode("utf-8")
    ).hexdigest()
    supplement_manifest = {
        "schema_version": "tcga-trace-pmc-supplement-manifest-v1",
        "pmcid": pmcid,
        "full_text_url": full_text_url,
        "supplement_url": supplement_url,
        "members": {
            "clinical": clinical_member,
            "expression": expression_member or None,
        },
        "external_urls": {
            "expression": expression_url or None,
            "expression_reference": expression_reference_url or None,
        },
        "clinical_record_mode": (
            "curated_from_supplement"
            if curated_records is not None
            else "parsed_xlsx"
        ),
        "source_sha256": raw_source_hashes,
    }
    supplement_manifest_path = (
        source_dir / "pmc_supplement_manifest.json"
    )
    supplement_manifest_path.write_text(
        canonical_json(supplement_manifest) + "\n",
        encoding="utf-8",
    )
    source_paths = {
        "expression": filtered_expression_path,
        "patients": patient_table,
        "samples": sample_table,
        "clinical_supplement": clinical_path,
        "expression_source": raw_expression_path,
        "supplement_manifest": supplement_manifest_path,
        "license": full_text_path,
    }
    if expression_reference_path is not None:
        source_paths["expression_reference"] = (
            expression_reference_path
        )
    return (
        source_paths,
        source_snapshot,
        str(source_spec.get("publication_date") or "") or None,
        {
            "source_api": api,
            "source_pmcid": pmcid,
            "source_license": source_license,
            "source_clinical_rows": len(clinical_rows),
            "source_eligible_rows": len(eligible_rows),
            "source_eligibility_rules": eligibility_rules,
            "source_clinical_rows_excluded_without_expression": (
                excluded_without_expression
            ),
            "source_clinical_record_mode": (
                "curated_from_supplement"
                if curated_records is not None
                else "parsed_xlsx"
            ),
            **expression_summary,
        },
    )


def detect_creative_commons_license(path: Path) -> str:
    root = ET.fromstring(path.read_bytes())
    licenses = [
        " ".join("".join(element.itertext()).split()).casefold()
        for element in root.iter()
        if local_name(element.tag) == "license"
    ]
    href_values = [
        str(value).casefold()
        for element in root.iter()
        if local_name(element.tag) == "license"
        for key, value in element.attrib.items()
        if local_name(key) == "href"
    ]
    combined = " ".join([*licenses, *href_values])
    if "creativecommons.org/licenses/by-nc-nd/4.0" in combined:
        return "CC-BY-NC-ND-4.0"
    if "creativecommons.org/licenses/by-nc/4.0" in combined:
        return "CC-BY-NC-4.0"
    if "creativecommons.org/licenses/by/4.0" in combined:
        return "CC-BY-4.0"
    raise ValueError(
        "Europe PMC article does not expose a supported CC BY, CC BY-NC "
        "or CC BY-NC-ND 4.0 license."
    )


def assert_cc_by_license(path: Path) -> None:
    detect_creative_commons_license(path)


def extract_zip_member(
    archive: zipfile.ZipFile,
    member: str,
    target: Path,
) -> None:
    if Path(member).name != member:
        raise ValueError(
            "Supplement member paths must be top-level filenames."
        )
    with archive.open(member) as source, target.open("wb") as destination:
        shutil.copyfileobj(source, destination)


def read_first_xlsx_sheet(path: Path) -> list[dict[str, str]]:
    return read_xlsx_sheet(path)


def read_xlsx_sheet(
    path: Path,
    *,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> list[dict[str, str]]:
    if header_row < 1:
        raise ValueError("Clinical XLSX header_row must be at least 1.")
    spreadsheet_ns = (
        "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    )
    office_rel_ns = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    package_rel_ns = (
        "http://schemas.openxmlformats.org/package/2006/relationships"
    )
    namespace = {"m": spreadsheet_ns, "r": office_rel_ns}
    with zipfile.ZipFile(path) as archive:
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(
                archive.read("xl/sharedStrings.xml")
            )
            shared = [
                "".join(element.itertext())
                for element in shared_root.findall("m:si", namespace)
            ]
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(
            archive.read("xl/_rels/workbook.xml.rels")
        )
        relationship_targets = {
            str(element.attrib.get("Id") or ""): str(
                element.attrib.get("Target") or ""
            )
            for element in relationships.findall(
                f"{{{package_rel_ns}}}Relationship"
            )
        }
        sheets = workbook.findall("m:sheets/m:sheet", namespace)
        selected = next(
            (
                element
                for element in sheets
                if sheet_name is None
                or element.attrib.get("name") == sheet_name
            ),
            None,
        )
        if selected is None:
            raise ValueError(
                f"Clinical XLSX sheet {sheet_name!r} was not found."
            )
        relationship_id = selected.attrib.get(
            f"{{{office_rel_ns}}}id"
        )
        target = relationship_targets.get(str(relationship_id))
        if not target:
            raise ValueError(
                "Clinical XLSX sheet relationship is missing."
            )
        normalized_target = (
            target.lstrip("/")
            if target.startswith("/")
            else f"xl/{target.lstrip('/')}"
        )
        sheet_root = ET.fromstring(archive.read(normalized_target))
    rows: list[tuple[int, dict[int, str]]] = []
    for row_element in sheet_root.findall(
        ".//m:sheetData/m:row", namespace
    ):
        row_number = int(row_element.attrib.get("r") or 0)
        row: dict[int, str] = {}
        for cell in row_element.findall("m:c", namespace):
            reference = str(cell.attrib.get("r") or "")
            match = re.match(r"([A-Z]+)[0-9]+", reference)
            if not match:
                continue
            cell_type = cell.attrib.get("t")
            if cell_type == "inlineStr":
                value = "".join(cell.itertext())
            else:
                value_element = cell.find("m:v", namespace)
                value = (
                    ""
                    if value_element is None
                    or value_element.text is None
                    else value_element.text
                )
            if cell_type == "s" and value:
                value = shared[int(value)]
            row[xlsx_column_index(match.group(1))] = value
        rows.append((row_number, row))
    if not rows:
        raise ValueError("Clinical XLSX contains no rows.")
    header_values = next(
        (row for number, row in rows if number == header_row),
        None,
    )
    if header_values is None:
        raise ValueError(
            f"Clinical XLSX header row {header_row} was not found."
        )
    max_column = max(header_values, default=-1)
    header = [
        header_values.get(index, "")
        for index in range(max_column + 1)
    ]
    if (
        any(not value for value in header)
        or len(header) != len(set(header))
    ):
        raise ValueError("Clinical XLSX header is missing or duplicated.")
    return [
        {
            field: row.get(index, "")
            for index, field in enumerate(header)
        }
        for number, row in rows
        if number > header_row
        if any(row.get(index, "") for index in range(len(header)))
    ]


def xlsx_column_index(label: str) -> int:
    value = 0
    for character in label:
        value = value * 26 + ord(character) - ord("A") + 1
    return value - 1


def download_filename(url: str) -> str:
    name = Path(urllib.parse.urlparse(url).path).name
    if not name or name in {".", ".."}:
        raise ValueError(f"Source URL has no safe filename: {url!r}")
    return name


def normalize_curated_records(
    records: list[dict[str, Any]],
) -> list[dict[str, str]]:
    if any(not isinstance(row, dict) or not row for row in records):
        raise ValueError(
            "Every curated clinical record must be a non-empty object."
        )
    fields = list(records[0])
    if any(list(row) != fields for row in records):
        raise ValueError(
            "Curated clinical records must use one stable field order."
        )
    return [
        {
            field: (
                ""
                if row[field] is None
                else str(row[field]).strip()
            )
            for field in fields
        }
        for row in records
    ]


def row_passes_rules(
    row: dict[str, str],
    rules: list[dict[str, Any]],
) -> bool:
    for rule in rules:
        column = str(rule.get("column") or "")
        value = str(row.get(column) or "").strip().casefold()
        included = {
            str(item).strip().casefold()
            for item in rule.get("include") or []
        }
        excluded = {
            str(item).strip().casefold()
            for item in rule.get("exclude") or []
        }
        if included and value not in included:
            return False
        if value in excluded:
            return False
        if rule.get("required") and not value:
            return False
    return True


def materialize_publication_expression(
    source: Path,
    target: Path,
    sample_ids: list[str],
    expression_spec: dict[str, Any],
    *,
    expression_reference_path: Path | None,
) -> dict[str, Any]:
    layout = str(
        expression_spec.get("layout") or "simple"
    ).strip()
    delimiter = str(expression_spec.get("delimiter") or "\t")
    if layout == "simple":
        return filter_expression_columns(
            source,
            target,
            sample_ids,
            delimiter=delimiter,
        )
    if layout == "gencode_ensembl_matrix":
        if expression_reference_path is None:
            raise ValueError(
                "gencode_ensembl_matrix requires "
                "expression_reference_url."
            )
        return materialize_gencode_ensembl_matrix(
            source,
            target,
            sample_ids,
            expression_reference_path,
            delimiter=delimiter,
            case_insensitive_samples=bool(
                expression_spec.get("case_insensitive_samples")
            ),
        )
    if layout == "geo_rpkm_two_row_header":
        return materialize_geo_rpkm_matrix(
            source,
            target,
            sample_ids,
            delimiter=delimiter,
            metadata_columns=int(
                expression_spec.get("metadata_columns") or 7
            ),
            ensembl_column=int(
                expression_spec.get("ensembl_column") or 2
            ),
            symbol_column=int(
                expression_spec.get("symbol_column") or 3
            ),
        )
    if layout == "xlsx_matrix":
        return materialize_xlsx_expression_matrix(
            source,
            target,
            sample_ids,
            sheet_name=str(
                expression_spec.get("sheet_name") or ""
            ).strip(),
            header_row=int(
                expression_spec.get("header_row") or 1
            ),
            feature_column=str(
                expression_spec.get("feature_column") or ""
            ).strip(),
            sample_start_column=int(
                expression_spec.get("sample_start_column") or 2
            ),
            feature_replacements=(
                expression_spec.get("feature_replacements") or {}
            ),
            case_insensitive_samples=bool(
                expression_spec.get("case_insensitive_samples")
            ),
        )
    raise ValueError(
        f"Unsupported publication expression layout: {layout!r}"
    )


def publication_expression_sample_ids(
    source: Path,
    expression_spec: dict[str, Any],
) -> list[str]:
    layout = str(
        expression_spec.get("layout") or "simple"
    ).strip()
    delimiter = str(expression_spec.get("delimiter") or "\t")
    if layout == "xlsx_matrix":
        header = read_xlsx_expression_header(
            source,
            sheet_name=str(
                expression_spec.get("sheet_name") or ""
            ).strip(),
            header_row=int(
                expression_spec.get("header_row") or 1
            ),
        )
        sample_start_index = int(
            expression_spec.get("sample_start_column") or 2
        ) - 1
        if sample_start_index < 0 or sample_start_index >= len(header):
            raise ValueError(
                "XLSX expression sample_start_column is outside the "
                "header."
            )
        return [
            value
            for value in header[sample_start_index:]
            if value
        ]
    with open_text_source(source) as input_handle:
        reader = csv.reader(input_handle, delimiter=delimiter)
        header = next(reader)
        if layout in {"simple", "gencode_ensembl_matrix"}:
            return [
                value.strip()
                for value in header[1:]
                if value.strip()
            ]
        if layout == "geo_rpkm_two_row_header":
            sample_labels = next(reader)
            metadata_columns = int(
                expression_spec.get("metadata_columns") or 7
            )
            return [
                value.strip()
                for value in sample_labels[metadata_columns:]
                if value.strip()
            ]
    raise ValueError(
        f"Unsupported publication expression layout: {layout!r}"
    )


def materialize_xlsx_expression_matrix(
    source: Path,
    target: Path,
    sample_ids: list[str],
    *,
    sheet_name: str,
    header_row: int,
    feature_column: str,
    sample_start_column: int,
    feature_replacements: dict[str, Any],
    case_insensitive_samples: bool,
) -> dict[str, Any]:
    if not sheet_name or not feature_column:
        raise ValueError(
            "xlsx_matrix requires sheet_name and feature_column."
        )
    if header_row < 1 or sample_start_column < 1:
        raise ValueError(
            "XLSX expression row and column indexes must be positive."
        )
    workbook = load_workbook(
        source,
        read_only=True,
        data_only=True,
    )
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"XLSX expression sheet {sheet_name!r} was not found."
            )
        worksheet = workbook[sheet_name]
        header = read_xlsx_expression_header(
            source,
            sheet_name=sheet_name,
            header_row=header_row,
            workbook=workbook,
        )
        try:
            feature_index = header.index(feature_column)
        except ValueError as error:
            raise ValueError(
                f"XLSX expression feature column {feature_column!r} "
                "was not found."
            ) from error
        sample_start_index = sample_start_column - 1
        if (
            sample_start_index <= feature_index
            or sample_start_index >= len(header)
        ):
            raise ValueError(
                "XLSX expression sample_start_column must follow the "
                "feature column and fall inside the header."
            )
        indexes = selected_sample_indexes(
            header,
            sample_ids,
            case_insensitive=case_insensitive_samples,
        )
        if any(index < sample_start_index for index in indexes):
            raise ValueError(
                "A requested XLSX expression sample resolved to a "
                "metadata column."
            )
        replacements = {
            str(source_value).strip(): str(target_value).strip()
            for source_value, target_value in feature_replacements.items()
        }
        if any(
            not source_value or not target_value
            for source_value, target_value in replacements.items()
        ):
            raise ValueError(
                "XLSX expression feature replacements must be non-empty."
            )
        source_gene_rows = 0
        replaced_features = 0
        with target.open(
            "w", newline="", encoding="utf-8"
        ) as output_handle:
            writer = csv.writer(
                output_handle,
                delimiter="\t",
                lineterminator="\n",
            )
            writer.writerow([feature_column, *sample_ids])
            for row in worksheet.iter_rows(
                min_row=header_row + 1,
                max_col=len(header),
                values_only=True,
            ):
                raw_feature = row[feature_index]
                if raw_feature is None:
                    continue
                feature = xlsx_cell_text(raw_feature)
                replacement = replacements.get(feature)
                if replacement:
                    feature = replacement
                    replaced_features += 1
                elif isinstance(raw_feature, (date, datetime)):
                    raise ValueError(
                        "XLSX expression contains a date-formatted feature "
                        f"{feature!r} without an explicit replacement."
                    )
                if not feature:
                    continue
                values = [
                    xlsx_cell_text(row[index])
                    for index in indexes
                ]
                if any(not value for value in values):
                    raise ValueError(
                        "XLSX expression contains a missing value for a "
                        f"selected sample at feature {feature!r}."
                    )
                try:
                    numeric_values = [float(value) for value in values]
                except ValueError as error:
                    raise ValueError(
                        "XLSX expression contains a non-numeric value for "
                        f"feature {feature!r}."
                    ) from error
                if any(
                    not math.isfinite(value)
                    for value in numeric_values
                ):
                    raise ValueError(
                        "XLSX expression contains a non-finite value for "
                        f"feature {feature!r}."
                    )
                writer.writerow([feature, *values])
                source_gene_rows += 1
    finally:
        workbook.close()
    if source_gene_rows < 10_000:
        raise ValueError(
            "XLSX expression matrix has fewer than 10,000 gene rows."
        )
    return {
        "source_expression_layout": "xlsx_matrix",
        "source_expression_columns": (
            len(header) - sample_start_index
        ),
        "selected_expression_columns": len(sample_ids),
        "source_expression_gene_rows": source_gene_rows,
        "source_expression_sheet": sheet_name,
        "source_expression_header_row": header_row,
        "source_expression_feature_column": feature_column,
        "source_expression_feature_replacements": replaced_features,
    }


def read_xlsx_expression_header(
    source: Path,
    *,
    sheet_name: str,
    header_row: int,
    workbook: Any | None = None,
) -> list[str]:
    if not sheet_name or header_row < 1:
        raise ValueError(
            "XLSX expression sheet_name and positive header_row are "
            "required."
        )
    owned_workbook = workbook is None
    active_workbook = workbook or load_workbook(
        source,
        read_only=True,
        data_only=True,
    )
    try:
        if sheet_name not in active_workbook.sheetnames:
            raise ValueError(
                f"XLSX expression sheet {sheet_name!r} was not found."
            )
        worksheet = active_workbook[sheet_name]
        values = next(
            worksheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                values_only=True,
            ),
            None,
        )
        if values is None:
            raise ValueError(
                f"XLSX expression header row {header_row} was not found."
            )
        header = [xlsx_cell_text(value) for value in values]
        while header and not header[-1]:
            header.pop()
        if (
            not header
            or any(not value for value in header)
            or len(header) != len(set(header))
        ):
            raise ValueError(
                "XLSX expression header is missing or duplicated."
            )
        return header
    finally:
        if owned_workbook:
            active_workbook.close()


def xlsx_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def filter_expression_columns(
    source: Path,
    target: Path,
    sample_ids: list[str],
    *,
    delimiter: str,
) -> dict[str, Any]:
    with open_text_source(source) as input_handle, target.open(
        "w", newline="", encoding="utf-8"
    ) as output_handle:
        reader = csv.reader(input_handle, delimiter=delimiter)
        writer = csv.writer(
            output_handle, delimiter="\t", lineterminator="\n"
        )
        header = next(reader)
        source_index = {
            sample_id: index
            for index, sample_id in enumerate(header)
        }
        missing = [
            sample_id
            for sample_id in sample_ids
            if sample_id not in source_index
        ]
        if missing:
            raise ValueError(
                "Clinical samples are absent from the supplementary "
                f"expression matrix: {missing[:5]}"
            )
        indexes = [source_index[sample_id] for sample_id in sample_ids]
        writer.writerow([header[0], *sample_ids])
        gene_rows = 0
        for row in reader:
            if not row:
                continue
            if len(row) != len(header):
                raise ValueError(
                    "Supplementary expression row width is inconsistent."
                )
            writer.writerow(
                [row[0], *(row[index] for index in indexes)]
            )
            gene_rows += 1
    if gene_rows < 10_000:
        raise ValueError(
            "Supplementary expression matrix has fewer than 10,000 genes."
        )
    return {
        "source_expression_layout": "simple",
        "source_expression_columns": len(header) - 1,
        "selected_expression_columns": len(sample_ids),
        "source_expression_gene_rows": gene_rows,
        "source_expression_delimiter": delimiter,
    }


def materialize_gencode_ensembl_matrix(
    source: Path,
    target: Path,
    sample_ids: list[str],
    gtf_path: Path,
    *,
    delimiter: str,
    case_insensitive_samples: bool,
) -> dict[str, Any]:
    gene_map = parse_gencode_gene_map(gtf_path)
    with open_text_source(source) as input_handle, target.open(
        "w", newline="", encoding="utf-8"
    ) as output_handle:
        reader = csv.reader(input_handle, delimiter=delimiter)
        writer = csv.writer(
            output_handle, delimiter="\t", lineterminator="\n"
        )
        header = next(reader)
        indexes = selected_sample_indexes(
            header,
            sample_ids,
            case_insensitive=case_insensitive_samples,
        )
        writer.writerow(
            ["Ensembl_Gene_Id|Hugo_Symbol", *sample_ids]
        )
        source_gene_rows = 0
        mapped_gene_rows = 0
        for row in reader:
            if not row:
                continue
            source_gene_rows += 1
            if len(row) != len(header):
                raise ValueError(
                    "GEO expression row width is inconsistent."
                )
            ensembl_id = row[0].strip()
            symbol = (
                gene_map.get(ensembl_id)
                or gene_map.get(ensembl_id.split(".", 1)[0])
            )
            if not symbol:
                continue
            writer.writerow(
                [
                    f"{ensembl_id}|{symbol}",
                    *(row[index] for index in indexes),
                ]
            )
            mapped_gene_rows += 1
    if mapped_gene_rows < 10_000:
        raise ValueError(
            "Fewer than 10,000 expression rows map to the pinned GENCODE "
            "reference."
        )
    return {
        "source_expression_layout": "gencode_ensembl_matrix",
        "source_expression_columns": len(header) - 1,
        "selected_expression_columns": len(sample_ids),
        "source_expression_gene_rows": source_gene_rows,
        "mapped_expression_gene_rows": mapped_gene_rows,
        "unmapped_expression_gene_rows": (
            source_gene_rows - mapped_gene_rows
        ),
        "source_expression_delimiter": delimiter,
        "expression_reference_genes": len(gene_map),
    }


def materialize_geo_rpkm_matrix(
    source: Path,
    target: Path,
    sample_ids: list[str],
    *,
    delimiter: str,
    metadata_columns: int,
    ensembl_column: int,
    symbol_column: int,
) -> dict[str, Any]:
    if metadata_columns < 1:
        raise ValueError("metadata_columns must be positive.")
    with open_text_source(source) as input_handle, target.open(
        "w", newline="", encoding="utf-8"
    ) as output_handle:
        reader = csv.reader(input_handle, delimiter=delimiter)
        writer = csv.writer(
            output_handle, delimiter="\t", lineterminator="\n"
        )
        header = next(reader)
        sample_labels = next(reader)
        if len(sample_labels) != len(header):
            raise ValueError(
                "GEO RPKM sample-label row width is inconsistent."
            )
        indexes = selected_sample_indexes(
            sample_labels,
            sample_ids,
            case_insensitive=False,
        )
        if any(index < metadata_columns for index in indexes):
            raise ValueError(
                "A requested sample resolved to a metadata column."
            )
        writer.writerow(
            ["Ensembl_Gene_Id|Hugo_Symbol", *sample_ids]
        )
        source_gene_rows = 0
        mapped_gene_rows = 0
        for row in reader:
            if not row:
                continue
            source_gene_rows += 1
            if len(row) != len(header):
                raise ValueError(
                    "GEO RPKM expression row width is inconsistent."
                )
            ensembl_id = row[ensembl_column].strip()
            symbol = row[symbol_column].strip()
            if not ensembl_id.startswith("ENSG") or not symbol:
                continue
            writer.writerow(
                [
                    f"{ensembl_id}|{symbol}",
                    *(row[index] for index in indexes),
                ]
            )
            mapped_gene_rows += 1
    if mapped_gene_rows < 10_000:
        raise ValueError(
            "GEO RPKM matrix has fewer than 10,000 Ensembl/HUGO rows."
        )
    return {
        "source_expression_layout": "geo_rpkm_two_row_header",
        "source_expression_columns": len(header) - metadata_columns,
        "selected_expression_columns": len(sample_ids),
        "source_expression_gene_rows": source_gene_rows,
        "mapped_expression_gene_rows": mapped_gene_rows,
        "unmapped_expression_gene_rows": (
            source_gene_rows - mapped_gene_rows
        ),
        "source_expression_delimiter": delimiter,
        "source_expression_metadata_columns": metadata_columns,
    }


def selected_sample_indexes(
    header: list[str],
    sample_ids: list[str],
    *,
    case_insensitive: bool,
) -> list[int]:
    def key(value: str) -> str:
        cleaned = str(value).strip()
        return cleaned.casefold() if case_insensitive else cleaned

    source_index: dict[str, int] = {}
    duplicates: set[str] = set()
    for index, value in enumerate(header):
        normalized = key(value)
        if not normalized:
            continue
        if normalized in source_index:
            duplicates.add(normalized)
        source_index[normalized] = index
    if duplicates:
        raise ValueError(
            "Expression matrix has duplicate sample labels after "
            f"normalization: {sorted(duplicates)[:5]}"
        )
    missing = [
        sample_id
        for sample_id in sample_ids
        if key(sample_id) not in source_index
    ]
    if missing:
        raise ValueError(
            "Clinical samples are absent from the expression matrix: "
            f"{missing[:5]}"
        )
    return [source_index[key(sample_id)] for sample_id in sample_ids]


def parse_gencode_gene_map(path: Path) -> dict[str, str]:
    mapping: dict[str, str] = {}
    conflicts: set[str] = set()
    with open_text_source(path) as handle:
        for line in handle:
            if not line or line.startswith("#"):
                continue
            fields = line.rstrip("\n").split("\t")
            if len(fields) != 9 or fields[2] != "gene":
                continue
            attributes = {
                key: value
                for key, value in re.findall(
                    r'([A-Za-z0-9_]+) "([^"]*)"', fields[8]
                )
            }
            gene_id = str(attributes.get("gene_id") or "").strip()
            symbol = str(attributes.get("gene_name") or "").strip()
            if not gene_id.startswith("ENSG") or not symbol:
                continue
            for key in {gene_id, gene_id.split(".", 1)[0]}:
                existing = mapping.get(key)
                if existing is not None and existing != symbol:
                    conflicts.add(key)
                else:
                    mapping[key] = symbol
    for key in conflicts:
        mapping.pop(key, None)
    return mapping


def open_text_source(path: Path):
    if path.suffix.lower() == ".gz":
        return gzip.open(
            path,
            mode="rt",
            newline="",
            encoding="utf-8-sig",
            errors="strict",
        )
    return path.open(
        newline="",
        encoding="utf-8-sig",
        errors="strict",
    )


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def canonical_json(value: Any) -> str:
    return json.dumps(
        value,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=True,
    )
