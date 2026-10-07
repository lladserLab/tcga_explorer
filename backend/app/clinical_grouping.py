from __future__ import annotations

from collections import Counter
import csv
from functools import lru_cache
from hashlib import sha256
import json
import math
from pathlib import Path
import re
import statistics
from typing import Any, Iterable
import unicodedata

from app.survival import sample_selection_key


CLINICAL_GROUPING_CATALOG_VERSION = "clinical-grouping-catalog-v1.5"
USER_CUSTOM_CLINICAL_SCHEMA_VERSION = "trace-user-custom-clinical-v1"
USER_CUSTOM_CLINICAL_METADATA_KEY = "user_custom_clinical"
USER_CUSTOM_CLINICAL_SOURCE = "User-declared private clinical metadata"
MAX_USER_CUSTOM_VARIABLES = 10
MIN_GROUPING_PATIENTS = 10
MIN_LEVEL_PATIENTS = 5
MAX_CATEGORICAL_LEVELS = 30
MAX_EXTERNAL_VARIABLES = 40
MAX_SERIALIZED_METADATA_BYTES = 2_000_000
LOW_COVERAGE_THRESHOLD = 0.5

MISSING_VALUES = {
    "",
    "#n/a",
    "-",
    "--",
    "[error]",
    "[not applicable]",
    "[not available]",
    "[not evaluated]",
    "[unknown]",
    "na",
    "n/a",
    "nan",
    "nd",
    "none",
    "missing",
    "n.a.",
    "not applicable",
    "not available",
    "not determined",
    "not determined or na",
    "not evaluated",
    "not evaluable",
    "not reported",
    "null",
    "unk",
    "unknown",
    "unknown or na",
    "unknown or not reported",
}

MISSING_KEYS = {
    re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()
    for value in MISSING_VALUES
}

PAM50_LEVEL_LABELS = {
    "luma": "Luminal A (PAM50)",
    "lumb": "Luminal B (PAM50)",
    "basal": "Basal-like (PAM50)",
    "basal like": "Basal-like (PAM50)",
    "her2": "HER2-enriched (PAM50)",
    "her2 enriched": "HER2-enriched (PAM50)",
    "normal": "Normal-like (PAM50)",
    "normal like": "Normal-like (PAM50)",
}


STANDARD_FIELDS: tuple[dict[str, Any], ...] = (
    {
        "id": "stage",
        "label": "Stage",
        "value_type": "categorical",
        "category": "standardized",
        "source": "Standardized dataset metadata",
        "description": "Pathologic or clinical stage normalized by the importer.",
    },
    {
        "id": "grade",
        "label": "Grade",
        "value_type": "categorical",
        "category": "standardized",
        "source": "Standardized dataset metadata",
        "description": "Histologic grade normalized by the importer.",
    },
    {
        "id": "gender",
        "label": "Gender",
        "value_type": "categorical",
        "category": "standardized",
        "source": "Standardized dataset metadata",
        "description": "Reported gender or sex field supplied by the dataset.",
    },
    {
        "id": "race",
        "label": "Race",
        "value_type": "categorical",
        "category": "standardized",
        "source": "Standardized dataset metadata",
        "description": "Reported race category supplied by the dataset.",
    },
    {
        "id": "age_at_index",
        "label": "Age at diagnosis",
        "value_type": "numeric",
        "category": "standardized",
        "source": "Standardized dataset metadata",
        "description": "Age in years at diagnosis or study index.",
        "unit": "years",
    },
)


TCGA_GDC_FIELDS: tuple[dict[str, Any], ...] = (
    {
        "id": "primary_diagnosis",
        "label": "Primary diagnosis",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {
        "id": "ajcc_pathologic_t",
        "label": "AJCC pathologic T",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {
        "id": "ajcc_pathologic_n",
        "label": "AJCC pathologic N",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {
        "id": "ajcc_pathologic_m",
        "label": "AJCC pathologic M",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {
        "id": "figo_stage",
        "label": "FIGO stage",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {
        "id": "laterality",
        "label": "Laterality",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {"id": "ethnicity", "label": "Ethnicity", "category": "clinical"},
    {
        "id": "classification_of_tumor",
        "label": "Tumor classification",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {
        "id": "metastasis_at_diagnosis",
        "label": "Metastasis at diagnosis",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {
        "id": "prior_malignancy",
        "label": "Prior malignancy",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {
        "id": "prior_treatment",
        "label": "Treatment recorded before specimen collection",
        "category": "clinical",
        "require_primary_diagnosis": True,
        "description": (
            "GDC diagnosis field indicating whether therapeutic agents were "
            "recorded before the body specimen was collected. A No value does "
            "not establish absence of later treatment."
        ),
        "timing": "before_specimen_collection",
        "level_labels": {
            "no": "No recorded prior treatment",
            "yes": "Recorded prior treatment",
        },
        "survival_eligible": False,
        "survival_unavailable_reason": (
            "Specimen-relative treatment timing is not guaranteed to precede "
            "the survival endpoint time origin."
        ),
        "analysis_note": (
            "Do not interpret No as definitively treatment-naive. Missing or "
            "unknown records remain unknown and are never combined with No."
        ),
    },
    {
        "id": "synchronous_malignancy",
        "label": "Synchronous malignancy",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {
        "id": "method_of_diagnosis",
        "label": "Method of diagnosis",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
    {
        "id": "tissue_or_organ_of_origin",
        "label": "Tissue or organ of origin",
        "category": "clinical",
        "require_primary_diagnosis": True,
    },
)


TCGA_COHORT_CLINICAL_FIELDS: dict[str, tuple[dict[str, Any], ...]] = {
    "TCGA-HNSC": (
        {
            "id": "tobacco_smoking_status",
            "label": "Tobacco smoking status",
        },
        {"id": "alcohol_history", "label": "Alcohol history"},
        {
            "id": "pack_years_smoked",
            "label": "Pack-years smoked",
            "value_type": "numeric",
            "unit": "pack-years",
        },
    ),
    "TCGA-LAML": (
        {"id": "calgb_risk_group", "label": "CALGB risk group"},
    ),
    "TCGA-LIHC": (
        {
            "id": "child_pugh_classification",
            "label": "Child–Pugh classification",
        },
    ),
    "TCGA-LUAD": (
        {
            "id": "tobacco_smoking_status",
            "label": "Tobacco smoking status",
        },
        {
            "id": "pack_years_smoked",
            "label": "Pack-years smoked",
            "value_type": "numeric",
            "unit": "pack-years",
        },
    ),
    "TCGA-LUSC": (
        {
            "id": "tobacco_smoking_status",
            "label": "Tobacco smoking status",
        },
        {
            "id": "pack_years_smoked",
            "label": "Pack-years smoked",
            "value_type": "numeric",
            "unit": "pack-years",
        },
    ),
}


TCGA_CDR_FIELDS: tuple[dict[str, Any], ...] = (
    {
        "id": "cdr.clinical_stage",
        "source_field": "clinical_stage",
        "label": "Clinical stage (TCGA-CDR)",
    },
    {
        "id": "cdr.histological_type",
        "source_field": "histological_type",
        "label": "Histological type (TCGA-CDR)",
    },
    {
        "id": "cdr.histological_grade",
        "source_field": "histological_grade",
        "label": "Histological grade (TCGA-CDR)",
    },
    {
        "id": "cdr.menopause_status",
        "source_field": "menopause_status",
        "label": "Menopause status (TCGA-CDR)",
    },
    {
        "id": "cdr.margin_status",
        "source_field": "margin_status",
        "label": "Margin status (TCGA-CDR)",
        "analysis_note": (
            "Margin status is a post-resection variable; comparisons are "
            "associational and must not be interpreted as baseline prognosis."
        ),
    },
    {
        "id": "cdr.residual_tumor",
        "source_field": "residual_tumor",
        "label": "Residual tumor (TCGA-CDR)",
        "analysis_note": (
            "Residual tumor is assessed after treatment or resection; comparisons "
            "are associational and not a baseline contrast."
        ),
    },
)


def _annotation(
    field: str,
    label: str,
    *,
    expression_derived: bool = False,
    variable_id: str | None = None,
    level_labels: dict[str, str] | None = None,
    method_summary: str | None = None,
    reference: dict[str, str] | None = None,
    comparability: str | None = None,
) -> dict[str, Any]:
    return {
        "id": variable_id or field,
        "source_field": field,
        "label": label,
        "expression_derived": expression_derived,
        "level_labels": level_labels or {},
        "method_summary": method_summary,
        "reference": reference,
        "comparability": comparability,
    }


# Deliberately curated. TCGA marker-paper tables also contain identifiers,
# endpoints, technical batches and thousands of molecular measurements that
# must not silently become analysis groupings.
TCGA_ANNOTATION_FIELDS: dict[str, tuple[dict[str, Any], ...]] = {
    "TCGA-ACC": (_annotation("paper_Histology", "Histologic subtype"),),
    "TCGA-BLCA": (
        _annotation("paper_Histologic.subtype", "Histologic subtype"),
        _annotation("paper_Histologic.grade", "Histologic grade"),
        _annotation(
            "paper_mRNA.cluster",
            "mRNA molecular subtype",
            expression_derived=True,
        ),
    ),
    "TCGA-BRCA": (
        _annotation(
            "paper_BRCA_Subtype_PAM50",
            "PAM50 intrinsic subtype",
            expression_derived=True,
            level_labels={
                "LumA": "Luminal A (PAM50)",
                "LumB": "Luminal B (PAM50)",
                "Basal": "Basal-like (PAM50)",
                "Her2": "HER2-enriched (PAM50)",
                "Normal": "Normal-like (PAM50)",
            },
            method_summary=(
                "Author-provided PAM50 call published with TCGA-BRCA; TRACE "
                "reads the reported call and does not rerun the classifier."
            ),
            reference={
                "label": "TCGA Breast Cancer, Nature 2012",
                "doi": "10.1038/nature11412",
                "url": "https://doi.org/10.1038/nature11412",
            },
            comparability=(
                "Method-dependent. Do not assume equivalence to receptor-defined "
                "subtype, a clinical assay, or a call from another expression platform."
            ),
        ),
        _annotation("paper_BRCA_Pathology", "Breast cancer pathology"),
    ),
    "TCGA-CESC": (
        _annotation("paper_CLIN.HPV_status", "HPV status"),
        _annotation("paper_CLIN.HPV_Hcall", "HPV type"),
    ),
    "TCGA-CHOL": (
        _annotation("paper_histologic.subtype", "Histologic subtype"),
        _annotation("paper_SUBTYPE..select.one.", "Pathology subtype"),
    ),
    "TCGA-COAD": (
        _annotation("paper_MSI_status", "Microsatellite instability status"),
        _annotation(
            "paper_expression_subtype",
            "Expression subtype",
            expression_derived=True,
        ),
    ),
    "TCGA-ESCA": (
        _annotation("paper_Histological.Type", "Histological type"),
        _annotation("paper_MSI.status", "Microsatellite instability status"),
        _annotation(
            "paper_ESCC.subtype",
            "Esophageal squamous subtype",
            expression_derived=True,
        ),
    ),
    "TCGA-GBM": (
        _annotation("paper_IDH.status", "IDH status"),
        _annotation("paper_IDH.codel.subtype", "IDH / 1p19q subtype"),
        _annotation("paper_MGMT.promoter.status", "MGMT promoter status"),
        _annotation("paper_ATRX.status", "ATRX status"),
        _annotation("paper_TERT.promoter.status", "TERT promoter status"),
        _annotation(
            "paper_Original.Subtype",
            "Original transcriptomic subtype",
            expression_derived=True,
        ),
        _annotation(
            "paper_Transcriptome.Subtype",
            "Transcriptome subtype",
            expression_derived=True,
        ),
    ),
    "TCGA-KIRC": (
        _annotation(
            "paper_mRNA_cluster",
            "mRNA molecular cluster",
            expression_derived=True,
        ),
    ),
    "TCGA-KIRP": (
        _annotation("paper_tumor_type.KIRP.path.", "Papillary RCC subtype"),
        _annotation(
            "paper_mRNA.clusters..3.group.NMF..Rathmell.group.",
            "mRNA molecular cluster",
            expression_derived=True,
        ),
    ),
    "TCGA-LGG": (
        _annotation("paper_Histology", "Histology"),
        _annotation("paper_IDH.status", "IDH status"),
        _annotation("paper_X1p.19q.codeletion", "1p/19q codeletion"),
        _annotation("paper_IDH.codel.subtype", "IDH / 1p19q subtype"),
        _annotation("paper_MGMT.promoter.status", "MGMT promoter status"),
        _annotation("paper_ATRX.status", "ATRX status"),
        _annotation("paper_TERT.promoter.status", "TERT promoter status"),
        _annotation(
            "paper_Transcriptome.Subtype",
            "Transcriptome subtype",
            expression_derived=True,
        ),
    ),
    "TCGA-LUAD": (
        _annotation(
            "paper_expression_subtype",
            "Expression subtype",
            expression_derived=True,
        ),
    ),
    "TCGA-LUSC": (
        _annotation(
            "paper_Expression.Subtype",
            "Expression subtype",
            expression_derived=True,
        ),
    ),
    "TCGA-PAAD": (
        _annotation("paper_Histological.type.by.DCC", "Histological type"),
        _annotation("paper_Grade", "Histological grade"),
        _annotation(
            "paper_mRNA.Bailey.Clusters..All.150.Samples..1squamous.2immunogenic.3progenitor.4ADEX",
            "Bailey expression subtype",
            expression_derived=True,
            variable_id="paper_PAAD_Bailey_expression_subtype",
        ),
        _annotation(
            "paper_mRNA.Moffitt.clusters..All.150.Samples..1basal..2classical",
            "Moffitt expression subtype",
            expression_derived=True,
            variable_id="paper_PAAD_Moffitt_expression_subtype",
        ),
        _annotation(
            "paper_mRNA.Collisson.clusters..All.150.Samples..1classical.2exocrine.3QM",
            "Collisson expression subtype",
            expression_derived=True,
            variable_id="paper_PAAD_Collisson_expression_subtype",
        ),
    ),
    "TCGA-PRAD": (
        _annotation("paper_Subtype", "Molecular subtype"),
    ),
    "TCGA-READ": (
        _annotation("paper_MSI_status", "Microsatellite instability status"),
        _annotation(
            "paper_expression_subtype",
            "Expression subtype",
            expression_derived=True,
        ),
    ),
    "TCGA-SARC": (
        _annotation("paper_histology", "Sarcoma histology"),
    ),
    "TCGA-SKCM": (
        _annotation("paper_MUTATIONSUBTYPES", "Mutation subtype"),
    ),
    "TCGA-STAD": (
        _annotation("paper_Molecular.Subtype", "Molecular subtype"),
        _annotation("paper_MSI.status", "Microsatellite instability status"),
    ),
    "TCGA-THCA": (
        _annotation("paper_histological_type", "Histological type"),
        _annotation("paper_BRAF_RAF_class", "BRAF / RAS-like class"),
    ),
    "TCGA-UCEC": (
        _annotation("paper_histology", "Histology"),
        _annotation("paper_histology_grade", "Histology / grade class"),
        _annotation(
            "paper_msi_status_7_marker_call",
            "Microsatellite instability status",
        ),
        _annotation(
            "paper_mrna_expression_cluster",
            "mRNA expression cluster",
            expression_derived=True,
        ),
    ),
    "TCGA-UCS": (
        _annotation("paper_Histologic.classification", "Histologic class"),
        _annotation("paper_histologic.subtype", "Histologic subtype"),
    ),
    "TCGA-UVM": (
        _annotation(
            "paper_mRNA.Cluster.No.",
            "mRNA molecular cluster",
            expression_derived=True,
        ),
    ),
}


EXTERNAL_ALLOWED_TOKENS = {
    "age": 60,
    "alcohol": 50,
    "cms": 90,
    "codel": 90,
    "ethnicity": 60,
    "gender": 60,
    "grade": 75,
    "her2": 95,
    "histologic": 85,
    "histology": 85,
    "hpv": 90,
    "idh": 90,
    "1p19q": 95,
    "menopaus": 80,
    "molecular": 85,
    "mgmt": 90,
    "msi": 90,
    "node": 70,
    "nodal": 70,
    "pam50": 100,
    "patholog": 75,
    "pgr": 90,
    "er_status": 95,
    "er_negpos": 100,
    "er_irs": 90,
    "estrogen": 90,
    "pr_status": 90,
    "pr_negpos": 100,
    "pr_irs": 90,
    "progesterone": 90,
    "race": 60,
    "receptor": 90,
    "response": 65,
    "sex": 60,
    "smoking": 55,
    "stage": 75,
    "subtype": 100,
    "treatment": 55,
}

EXTERNAL_NUMERIC_TOKENS = {
    "age": 80,
    "count": 70,
    "cycle": 70,
    "cycles": 70,
    "dose": 75,
    "er_irs": 90,
    "index": 75,
    "lymph_nodes_examined": 85,
    "lymph_nodes_positive": 85,
    "number": 70,
    "pack_year": 90,
    "percent": 85,
    "percentage": 85,
    "pr_irs": 90,
    "ratio": 80,
    "score": 75,
}
EXTERNAL_FORBIDDEN_TOKENS = {
    "accession",
    "barcode",
    "checksum",
    "days_to",
    "death",
    "event",
    "expression_value",
    "filename",
    "followup",
    "follow_up",
    "identifier",
    "overall_survival",
    "patient_id",
    "pfs_time",
    "release",
    "sample_id",
    "sha256",
    "source_file",
    "survival_time",
    "time_days",
    USER_CUSTOM_CLINICAL_METADATA_KEY,
}

FIELD_SPECIFIC_MISSING_VALUES = {
    "transcriptional_response": {"999"},
    "response": {"un"},
    "smoking_history": {"u"},
}


def _missing_key(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value).casefold().strip()
    return re.sub(r"[^a-z0-9]+", " ", normalized).strip()


def _is_missing_text(value: str, *, source_field: str = "") -> bool:
    key = _missing_key(value)
    if key in MISSING_KEYS or key.startswith("missing "):
        return True
    normalized_field = re.sub(
        r"[^a-z0-9]+",
        "_",
        unicodedata.normalize("NFKC", source_field).casefold(),
    ).strip("_")
    compact_value = key.replace(" ", "")
    return any(
        semantic_key in normalized_field and compact_value in sentinels
        for semantic_key, sentinels in FIELD_SPECIFIC_MISSING_VALUES.items()
    )


def _clean_value(
    value: Any,
    *,
    source_field: str = "",
) -> str | float | None:
    if value is None or isinstance(value, (dict, list, tuple, set)):
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        numeric = float(value)
        if _is_missing_text(f"{numeric:g}", source_field=source_field):
            return None
        return numeric if math.isfinite(numeric) else None
    text = str(value).strip()
    if _is_missing_text(text, source_field=source_field):
        return None
    return text


def _clean_category(value: Any, *, source_field: str = "") -> str | None:
    cleaned = _clean_value(value, source_field=source_field)
    if cleaned is None:
        return None
    if isinstance(cleaned, float):
        return f"{cleaned:g}"
    return re.sub(r"\s+", " ", cleaned).strip()


def _clean_numeric(value: Any, *, source_field: str = "") -> float | None:
    cleaned = _clean_value(value, source_field=source_field)
    if cleaned is None:
        return None
    try:
        numeric = float(str(cleaned).replace(",", "").strip())
    except (TypeError, ValueError):
        return None
    return numeric if math.isfinite(numeric) else None


def _normalized_field_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.casefold()).strip("_")


def _field_leaf(path: str) -> str:
    return _normalized_field_name(path.split(".")[-1])


def _is_gender_field(path: str) -> bool:
    normalized = _normalized_field_name(path)
    tokens = set(normalized.split("_"))
    return _field_leaf(path) in {"gender", "sex", "sex_at_birth"} or bool(
        tokens.intersection({"gender", "sex"})
    )


def _is_stage_field(path: str) -> bool:
    return "stage" in _normalized_field_name(path).split("_")


def _is_histology_field(path: str) -> bool:
    tokens = set(_normalized_field_name(path).split("_"))
    return bool(tokens.intersection({"histology", "histological", "histologic"}))


def _canonical_stage_value(value: str) -> str:
    normalized = re.sub(r"\s*-\s*", "-", value.strip())
    normalized = re.sub(
        r"(?i)\b(II|III|IV)\s+([A-D])\b",
        lambda match: f"{match.group(1).upper()}{match.group(2).upper()}",
        normalized,
    )
    normalized = re.sub(
        r"(?i)\b([pc]?T(?:is|x|[0-4][a-d]?))\s+([pc]?N(?:x|[0-3][a-c]?))",
        lambda match: f"{match.group(1)}{match.group(2)}",
        normalized,
    )
    normalized = re.sub(
        r"(?i)([pc]?N(?:x|[0-3][a-c]?))\s+([pc]?M(?:x|[0-1][a-c]?))",
        lambda match: f"{match.group(1)}{match.group(2)}",
        normalized,
    )
    return normalized


def _canonical_histology_value(value: str) -> str:
    normalized = _normalized_field_name(value)
    if normalized == "idc" or (
        ("ductal" in normalized or "duct" in normalized)
        and ("85003" in normalized or "invasive" in normalized or "infiltrating" in normalized)
    ):
        return "Invasive ductal carcinoma (IDC)"
    if normalized == "ilc" or (
        ("lobular" in normalized or "lobule" in normalized)
        and ("85203" in normalized or "invasive" in normalized or "infiltrating" in normalized)
    ):
        return "Invasive lobular carcinoma (ILC)"
    return value


def _canonical_category_mapping(
    values: dict[str, str | float | None],
    *,
    source_field: str,
    concept: str | None = None,
) -> dict[str, str | float | None]:
    cleaned = {
        barcode: _clean_category(value, source_field=source_field)
        for barcode, value in values.items()
    }
    if _is_gender_field(source_field) or concept == "gender":
        aliases = {
            "f": "Female",
            "female": "Female",
            "feminine": "Female",
            "woman": "Female",
            "w": "Female",
            "weiblich": "Female",
            "m": "Male",
            "male": "Male",
            "masculine": "Male",
            "man": "Male",
        }
        normalized_field = _normalized_field_name(source_field)
        if "0_male" in normalized_field and "1_female" in normalized_field:
            aliases.update({"0": "Male", "1": "Female"})
        elif "0_female" in normalized_field and "1_male" in normalized_field:
            aliases.update({"0": "Female", "1": "Male"})
        cleaned = {
            barcode: aliases.get(str(value).casefold(), value)
            if value is not None
            else None
            for barcode, value in cleaned.items()
        }
    if _is_stage_field(source_field) or concept == "stage":
        cleaned = {
            barcode: _canonical_stage_value(value) if value is not None else None
            for barcode, value in cleaned.items()
        }
    if _is_histology_field(source_field):
        cleaned = {
            barcode: _canonical_histology_value(value) if value is not None else None
            for barcode, value in cleaned.items()
        }
    cleaned = {
        barcode: "Other"
        if value is not None and value.casefold() in {"other", "others"}
        else value
        for barcode, value in cleaned.items()
    }

    variants: dict[str, Counter[str]] = {}
    for value in cleaned.values():
        if value is None:
            continue
        variants.setdefault(value.casefold(), Counter())[value] += 1
    canonical = {
        key: sorted(counts.items(), key=lambda item: (-item[1], item[0]))[0][0]
        for key, counts in variants.items()
    }
    return {
        barcode: canonical.get(value.casefold(), value)
        if value is not None
        else None
        for barcode, value in cleaned.items()
    }


def _representative_samples(samples: Iterable[Any]) -> list[Any]:
    retained: dict[str, Any] = {}
    for sample in sorted(samples, key=sample_selection_key):
        retained.setdefault(str(sample.patient_id), sample)
    return list(retained.values())


@lru_cache(maxsize=4)
def _read_delimited_rows(
    path_text: str,
    modified_ns: int,
) -> tuple[dict[str, str], ...]:
    del modified_ns
    path = Path(path_text)
    with path.open(newline="", encoding="utf-8", errors="replace") as handle:
        return tuple(csv.DictReader(handle, delimiter="\t"))


def _tcga_rows(path: Path) -> tuple[dict[str, str], ...]:
    if not path.is_file():
        return ()
    return _read_delimited_rows(str(path.resolve()), path.stat().st_mtime_ns)


@lru_cache(maxsize=8)
def _read_cdr_rows(
    path_text: str,
    modified_ns: int,
) -> tuple[dict[str, Any], ...]:
    del modified_ns
    path = Path(path_text)
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - installed in production
            raise RuntimeError("openpyxl is required to read TCGA-CDR metadata.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook["TCGA-CDR"] if "TCGA-CDR" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        return tuple(
            {
                headers[index]: value
                for index, value in enumerate(row)
                if index < len(headers) and headers[index]
            }
            for row in rows
            if any(value is not None and str(value).strip() for value in row)
        )
    delimiter = "\t" if path.suffix.lower() in {".tsv", ".txt"} else ","
    with path.open(newline="", encoding="utf-8", errors="replace") as handle:
        return tuple(csv.DictReader(handle, delimiter=delimiter))


def _cdr_rows(path: Path | None) -> tuple[dict[str, Any], ...]:
    if path is None or not path.is_file():
        return ()
    return _read_cdr_rows(str(path.resolve()), path.stat().st_mtime_ns)


def _tcga_row_indexes(
    rows: Iterable[dict[str, Any]],
) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    by_barcode: dict[str, dict[str, Any]] = {}
    by_patient: dict[str, dict[str, Any]] = {}
    for row in rows:
        barcode = _clean_category(row.get("") or row.get("barcode") or row.get("sample"))
        patient = _clean_category(
            row.get("patient_id")
            or row.get("patient")
            or row.get("bcr_patient_barcode")
            or (barcode[:12] if barcode else None)
        )
        if barcode:
            by_barcode[barcode] = row
        if patient:
            by_patient.setdefault(patient, row)
    return by_barcode, by_patient


def _cdr_index(
    rows: Iterable[dict[str, Any]], cohort: str
) -> dict[str, dict[str, Any]]:
    cohort_code = cohort.removeprefix("TCGA-").upper()
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        row_cohort = str(row.get("type") or "").strip().removeprefix("TCGA-").upper()
        patient = _clean_category(
            row.get("bcr_patient_barcode")
            or row.get("patient_id")
            or row.get("submitter_id")
        )
        if row_cohort == cohort_code and patient:
            result[patient] = row
    return result


def _flatten_metadata(
    value: Any,
    *,
    prefix: str = "",
    depth: int = 0,
) -> dict[str, Any]:
    if not isinstance(value, dict) or depth > 6:
        return {}
    result: dict[str, Any] = {}
    for raw_key, raw_value in value.items():
        key = str(raw_key).strip()
        if not key:
            continue
        path = f"{prefix}.{key}" if prefix else key
        decoded_value = raw_value
        if (
            isinstance(raw_value, str)
            and len(raw_value.encode("utf-8", errors="ignore"))
            <= MAX_SERIALIZED_METADATA_BYTES
            and raw_value.lstrip().startswith("{")
        ):
            try:
                candidate = json.loads(raw_value)
            except (TypeError, ValueError, json.JSONDecodeError):
                candidate = None
            if isinstance(candidate, dict):
                decoded_value = candidate
        if isinstance(decoded_value, dict):
            result.update(
                _flatten_metadata(decoded_value, prefix=path, depth=depth + 1)
            )
        elif not isinstance(decoded_value, (list, tuple, set)):
            result[path] = decoded_value
    return result


def _user_custom_payload(raw_metadata: Any) -> dict[str, Any] | None:
    """Return the validated wrapper persisted by private-dataset import.

    Repository samples wrap patient and sample metadata before they reach this
    module. Supporting both forms keeps older repository helpers and direct
    tests compatible without treating arbitrary flattened metadata as a
    trusted declaration.
    """

    if not isinstance(raw_metadata, dict):
        return None
    candidates = [raw_metadata.get(USER_CUSTOM_CLINICAL_METADATA_KEY)]
    for wrapper in ("patient", "sample"):
        wrapped = raw_metadata.get(wrapper)
        if isinstance(wrapped, dict):
            candidates.append(wrapped.get(USER_CUSTOM_CLINICAL_METADATA_KEY))
    for candidate in candidates:
        if not isinstance(candidate, dict):
            continue
        if candidate.get("schema_version") != USER_CUSTOM_CLINICAL_SCHEMA_VERSION:
            continue
        definitions = candidate.get("definitions")
        values = candidate.get("values")
        if isinstance(definitions, list) and isinstance(values, dict):
            return candidate
    return None


def _validated_user_custom_definition(raw: Any) -> dict[str, Any] | None:
    if not isinstance(raw, dict):
        return None
    variable_id = str(raw.get("id") or "").strip()
    label = str(raw.get("label") or "").strip()
    source_column = str(raw.get("source_column") or "").strip()
    value_type = str(raw.get("value_type") or "").strip()
    timing = str(raw.get("timing") or "unknown").strip().casefold().replace(
        "-", "_"
    )
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_.-]{0,63}", variable_id):
        return None
    if not label or len(label) > 100 or not source_column:
        return None
    if value_type not in {"categorical", "numeric"}:
        return None
    if timing not in {
        "baseline",
        "pre_treatment",
        "on_treatment",
        "post_treatment",
        "post_baseline",
        "outcome",
        "unknown",
    }:
        return None
    unit = str(raw.get("unit") or "").strip() or None
    description = str(raw.get("description") or "").strip() or None
    return {
        "id": variable_id,
        "label": label,
        "source_column": source_column,
        "value_type": value_type,
        "unit": unit[:64] if unit else None,
        "timing": timing,
        "expression_derived": bool(raw.get("expression_derived", False)),
        "description": description[:500] if description else None,
    }


def _user_custom_specs(samples: Iterable[Any]) -> tuple[dict[str, Any], ...]:
    registry: dict[str, dict[str, Any]] = {}
    for sample in samples:
        payload = _user_custom_payload(getattr(sample, "raw_metadata", None))
        if payload is None:
            continue
        definitions = payload.get("definitions") or []
        if len(definitions) > MAX_USER_CUSTOM_VARIABLES:
            raise ValueError(
                "Private clinical metadata contains more than 10 custom variables."
            )
        for raw_definition in definitions:
            definition = _validated_user_custom_definition(raw_definition)
            if definition is None:
                raise ValueError(
                    "Private clinical metadata contains an invalid custom variable declaration."
                )
            key = definition["id"].casefold()
            existing = registry.get(key)
            if existing is not None and existing != definition:
                raise ValueError(
                    f"Custom clinical variable {definition['id']!r} has inconsistent declarations."
                )
            registry[key] = definition
    ordered = sorted(registry.values(), key=lambda item: item["id"].casefold())
    return tuple(
        {
            "id": definition["id"],
            "source_field": (
                f"{USER_CUSTOM_CLINICAL_METADATA_KEY}.values."
                f"{definition['id']}"
            ),
            "declared_source_column": definition["source_column"],
            "label": definition["label"],
            "value_type": definition["value_type"],
            "category": "dataset_specific",
            "source": USER_CUSTOM_CLINICAL_SOURCE,
            "description": definition["description"],
            "unit": definition["unit"],
            "timing": definition["timing"],
            "expression_derived": definition["expression_derived"],
            "user_custom": True,
        }
        for definition in ordered
    )


def _user_custom_value(sample: Any, variable_id: str) -> Any:
    payload = _user_custom_payload(getattr(sample, "raw_metadata", None))
    if payload is None:
        return None
    return (payload.get("values") or {}).get(variable_id)


def _external_field_score(path: str) -> int | None:
    normalized = _normalized_field_name(path)
    if any(token in normalized for token in EXTERNAL_FORBIDDEN_TOKENS):
        return None
    scores = [
        score
        for token, score in EXTERNAL_ALLOWED_TOKENS.items()
        if token in normalized
    ]
    return max(scores) if scores else None


def _external_variable_id(path: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9_.-]+", "_", path).strip("._-")[:80]
    digest = sha256(path.encode("utf-8")).hexdigest()[:8]
    return f"metadata.{slug or 'field'}.{digest}"


def _is_metadata_wrapper(segment: str) -> bool:
    normalized = _normalized_field_name(segment)
    return normalized in {
        "characteristics",
        "clinical",
        "clinical_metadata",
        "curated_clinical_record",
        "metadata",
        "patient",
        "raw_metadata",
        "sample",
    } or normalized.endswith("_metadata_json")


def _external_semantic_key(path: str) -> str:
    parts = [
        normalized
        for raw in path.split(".")
        if not _is_metadata_wrapper(raw)
        if (normalized := _normalized_field_name(raw))
    ]
    return ".".join(parts)


def _humanize_field(path: str) -> str:
    semantic_key = _external_semantic_key(path) or _field_leaf(path)
    normalized = _normalized_field_name(semantic_key)
    special_labels = {
        "er_negpos": "ER status (observed)",
        "er_status": "ER status (observed)",
        "er_irs": "ER immunoreactive score (IRS)",
        "her2_negpos": "HER2 status (observed)",
        "her2_status": "HER2 status (observed)",
        "pgr_negpos": "PGR status (observed)",
        "pgr_status": "PGR status (observed)",
        "pr_negpos": "PR status (observed)",
        "pr_status": "PR status (observed)",
        "pr_irs": "PR immunoreactive score (IRS)",
    }
    if normalized.endswith("pam50_subtype"):
        return "PAM50 intrinsic subtype"
    for suffix, label in special_labels.items():
        if normalized.endswith(suffix):
            return label
    prediction = re.search(r"(?i)(er|her2|pgr|pr)_prediction_(mgc|sgc)$", normalized)
    if prediction:
        return f"{prediction.group(1).upper()} prediction ({prediction.group(2).upper()})"
    text = re.sub(r"[._-]+", " ", semantic_key).strip()
    text = re.sub(r"\s+", " ", text)
    replacements = {
        "ajcc": "AJCC",
        "bcg": "BCG",
        "cms": "CMS",
        "er": "ER",
        "ensat": "ENSAT",
        "fish": "FISH",
        "her2": "HER2",
        "hpv": "HPV",
        "idh": "IDH",
        "ihc": "IHC",
        "io": "IO",
        "irs": "IRS",
        "mgc": "MGC",
        "msi": "MSI",
        "pam50": "PAM50",
        "pgr": "PGR",
        "pr": "PR",
        "sgc": "SGC",
        "tnm": "TNM",
    }
    words = [replacements.get(word.casefold(), word) for word in text.split()]
    humanized = " ".join(words)
    return f"{humanized[:1].upper()}{humanized[1:]}" if humanized else path


def _external_standard_concept(path: str) -> str | None:
    semantic_key = _external_semantic_key(path)
    leaf = _field_leaf(semantic_key)
    normalized = _normalized_field_name(semantic_key)
    aliases = {
        "age": "age_at_index",
        "age_at_diagnosis": "age_at_index",
        "age_at_index": "age_at_index",
        "clinical_stage": "stage",
        "gender": "gender",
        "grade": "grade",
        "histologic_grade": "grade",
        "histological_grade": "grade",
        "nhg": "grade",
        "nuclear_grade": "grade",
        "pathologic_stage": "stage",
        "pathological_stage": "stage",
        "race": "race",
        "sex": "gender",
        "sex_at_birth": "gender",
        "stage": "stage",
        "tumor_grade": "grade",
        "tumour_grade": "grade",
    }
    concept = aliases.get(leaf) or aliases.get(normalized)
    if concept:
        return concept
    if normalized.endswith("patient_gender") or normalized.startswith("sex_0_"):
        return "gender"
    return None


def _external_level_labels(path: str) -> dict[str, str]:
    normalized = _normalized_field_name(_external_semantic_key(path))
    labels: dict[str, str] = {}
    if "pam50" in normalized:
        labels.update(PAM50_LEVEL_LABELS)
    if re.search(r"(?:^|_)(?:er|her2|pgr|pr)_(?:negpos|status|prediction)(?:_|$)", normalized):
        labels.update(
            {
                "0": "Negative",
                "1": "Positive",
                "n": "Negative",
                "negative": "Negative",
                "p": "Positive",
                "positive": "Positive",
            }
        )
    if "sex_0_male_1_female" in normalized:
        labels.update({"0": "Male", "1": "Female"})
    if "sex_0_female_1_male" in normalized:
        labels.update({"0": "Female", "1": "Male"})
    if any(token in normalized for token in ("therapy", "treatment")):
        labels.update({"n": "No", "y": "Yes"})
    return labels


def _external_numeric_field(
    path: str,
    raw_values: list[Any],
) -> bool:
    normalized = _normalized_field_name(_external_semantic_key(path))
    if not any(token in normalized for token in EXTERNAL_NUMERIC_TOKENS):
        return False
    excluded = {
        "class",
        "gender",
        "grade",
        "histology",
        "race",
        "sex",
        "stage",
        "status",
        "subtype",
        "type",
    }
    if set(normalized.split("_")).intersection(excluded):
        return False
    observed = [
        _clean_category(value, source_field=path)
        for value in raw_values
    ]
    observed = [value for value in observed if value is not None]
    if not observed:
        return False
    numeric = [
        _clean_numeric(value, source_field=path)
        for value in observed
    ]
    finite = [value for value in numeric if value is not None]
    if len(finite) / len(observed) < 0.9 or len(set(finite)) < 2:
        return False
    return len(set(finite)) > 2 or not set(finite).issubset({0.0, 1.0})


def _external_numeric_unit(path: str) -> str | None:
    normalized = _normalized_field_name(_external_semantic_key(path))
    if "pack_year" in normalized:
        return "pack-years"
    if "percent" in normalized or "percentage" in normalized:
        return "%"
    if "age" in normalized:
        return "years"
    if "dose" in normalized and "radi" in normalized:
        return "Gy (as reported)"
    return None


def _external_standard_like(path: str) -> bool:
    tokens = set(_normalized_field_name(_external_semantic_key(path)).split("_"))
    return bool(tokens.intersection({"age", "gender", "grade", "race", "sex", "stage"}))


def _external_value_signature(
    values: tuple[str | float | None, ...],
) -> tuple[str | float | None, ...]:
    return tuple(
        value.casefold() if isinstance(value, str) else value
        for value in values
    )


def _standard_comparison_signature(
    values: tuple[str | float | None, ...],
    concept: str,
) -> tuple[str | float | None, ...]:
    roman = {"0": "0", "1": "I", "2": "II", "3": "III", "4": "IV"}
    normalized: list[str | float | None] = []
    for value in values:
        if not isinstance(value, str):
            normalized.append(value)
            continue
        text = value.casefold().strip()
        if concept == "stage":
            text = re.sub(r"^(?:tnm\s+)?stage\s+", "", text)
            text = re.sub(r"\s+", "", text).upper()
            match = re.fullmatch(r"([0-4])([A-D]?)", text)
            if match:
                text = f"{roman[match.group(1)]}{match.group(2)}"
        elif concept == "grade":
            text = re.sub(r"^(?:histologic(?:al)?\s+)?grade\s*", "", text)
            text = re.sub(r"^g(?=[1-4]$)", "", text)
            text = text.upper()
        normalized.append(text)
    return tuple(normalized)


def _compatible_external_values(
    left: tuple[str | float | None, ...],
    right: tuple[str | float | None, ...],
) -> bool:
    for left_value, right_value in zip(left, right, strict=True):
        if left_value is None or right_value is None:
            continue
        if isinstance(left_value, str) and isinstance(right_value, str):
            if left_value.casefold() != right_value.casefold():
                return False
        elif left_value != right_value:
            return False
    return True


def _external_path_priority(path: str) -> tuple[int, int, int]:
    normalized = _normalized_field_name(path)
    return (
        0 if path.startswith("patient.") else 1,
        0 if "curated_clinical_record" in normalized else 1,
        len(path),
    )


def _metadata_path_is_mostly_numeric(
    path: str,
    flattened_metadata: dict[str, dict[str, Any]],
) -> bool:
    observed = [
        metadata.get(path)
        for metadata in flattened_metadata.values()
        if _clean_value(metadata.get(path), source_field=path) is not None
    ]
    return bool(observed) and (
        sum(
            _clean_numeric(value, source_field=path) is not None
            for value in observed
        )
        / len(observed)
        >= 0.9
    )


def _repository_standard_context(
    samples: list[Any],
    flattened_metadata: dict[str, dict[str, Any]],
) -> tuple[
    dict[str, dict[str, Any]],
    dict[str, str],
    set[str],
]:
    all_paths = sorted(
        {
            path
            for metadata in flattened_metadata.values()
            for path in metadata
        }
    )
    values_by_standard: dict[str, dict[str, Any]] = {}
    source_fields: dict[str, str] = {}
    claimed_paths: set[str] = set()
    for spec in STANDARD_FIELDS:
        variable_id = str(spec["id"])
        candidate_paths = [
            path
            for path in all_paths
            if _external_standard_concept(path) == variable_id
        ]
        if variable_id == "age_at_index":
            candidate_paths = [
                path
                for path in candidate_paths
                if _metadata_path_is_mostly_numeric(path, flattened_metadata)
            ]
        candidate_paths = sorted(
            candidate_paths,
            key=lambda path: (_external_path_priority(path), path),
        )
        claimed_paths.update(candidate_paths)
        direct_count = sum(
            _clean_value(
                getattr(sample, variable_id, None),
                source_field=variable_id,
            )
            is not None
            for sample in samples
        )
        resolved: dict[str, Any] = {}
        contributing_paths: Counter[str] = Counter()
        for sample in samples:
            barcode = str(sample.barcode)
            raw = getattr(sample, variable_id, None)
            if _clean_value(raw, source_field=variable_id) is None:
                metadata = flattened_metadata.get(barcode, {})
                for path in candidate_paths:
                    candidate = metadata.get(path)
                    if _clean_value(candidate, source_field=path) is not None:
                        raw = candidate
                        contributing_paths[path] += 1
                        break
            resolved[barcode] = raw
        values_by_standard[variable_id] = resolved
        if direct_count:
            source_fields[variable_id] = variable_id
        elif contributing_paths:
            source_fields[variable_id] = sorted(
                contributing_paths,
                key=lambda path: (
                    -contributing_paths[path],
                    _external_path_priority(path),
                    path,
                ),
            )[0]
        else:
            source_fields[variable_id] = variable_id
    return values_by_standard, source_fields, claimed_paths


def _external_specs(
    samples: list[Any],
    flattened_metadata: dict[str, dict[str, Any]] | None = None,
    *,
    claimed_paths: set[str] | None = None,
    standard_values: dict[str, dict[str, Any]] | None = None,
) -> tuple[dict[str, Any], ...]:
    representatives = _representative_samples(samples)
    counts_by_path: dict[str, Counter[str | float]] = {}
    values_by_path: dict[str, tuple[str | float | None, ...]] = {}
    value_types: dict[str, str] = {}
    scores: dict[str, int] = {}
    flattened_by_sample = [
        (flattened_metadata or {}).get(str(sample.barcode))
        or _flatten_metadata(getattr(sample, "raw_metadata", None) or {})
        for sample in representatives
    ]
    all_paths = sorted(
        {
            path
            for metadata in flattened_by_sample
            for path in metadata
        }
    )
    standard_vectors: dict[str, tuple[str | float | None, ...]] = {}
    for spec in STANDARD_FIELDS:
        variable_id = str(spec["id"])
        raw_mapping = (standard_values or {}).get(variable_id, {})
        raw_values = [
            raw_mapping.get(str(sample.barcode), getattr(sample, variable_id, None))
            for sample in representatives
        ]
        if spec.get("value_type") == "numeric":
            standard_vectors[variable_id] = tuple(
                _clean_numeric(value, source_field=variable_id)
                for value in raw_values
            )
        else:
            canonical = _canonical_category_mapping(
                {
                    str(index): value
                    for index, value in enumerate(raw_values)
                },
                source_field=variable_id,
                concept=variable_id,
            )
            standard_vectors[variable_id] = tuple(
                canonical[str(index)]
                if isinstance(canonical[str(index)], str)
                else None
                for index in range(len(raw_values))
            )
    for path in all_paths:
        if path in (claimed_paths or set()):
            continue
        raw_values = [metadata.get(path) for metadata in flattened_by_sample]
        if _external_numeric_field(path, raw_values):
            value_types[path] = "numeric"
            ordered_values = tuple(
                _clean_numeric(value, source_field=path)
                for value in raw_values
            )
        else:
            value_types[path] = "categorical"
            canonical_values = _canonical_category_mapping(
                {
                    str(index): value
                    for index, value in enumerate(raw_values)
                },
                source_field=path,
            )
            ordered_values = tuple(
                canonical_values[str(index)]
                if isinstance(canonical_values[str(index)], str)
                else None
                for index in range(len(raw_values))
            )
        values_by_path[path] = ordered_values
        for value in ordered_values:
            if value is not None:
                counts_by_path.setdefault(path, Counter())[value] += 1
        score = _external_field_score(path)
        if score is not None:
            scores[path] = score

    candidates: list[tuple[int, str]] = []
    for path, counts in counts_by_path.items():
        score = scores.get(path)
        if score is None:
            continue
        if value_types[path] == "numeric" or (
            2 <= len(counts) <= MAX_CATEGORICAL_LEVELS
        ):
            candidates.append((score, path))
    candidates.sort(
        key=lambda item: (
            _external_path_priority(item[1]),
            -item[0],
            _humanize_field(item[1]),
            item[1],
        )
    )
    merged: list[dict[str, Any]] = []
    for score, path in candidates:
        values = values_by_path[path]
        if _external_standard_like(path):
            exact_standard_duplicate = any(
                _external_value_signature(values)
                == _external_value_signature(standard_vector)
                for standard_vector in standard_vectors.values()
            )
            path_tokens = set(
                _normalized_field_name(_external_semantic_key(path)).split("_")
            )
            semantic_standard_duplicate = any(
                concept in path_tokens
                and _standard_comparison_signature(values, concept)
                == _standard_comparison_signature(
                    standard_vectors[concept], concept
                )
                for concept in ("stage", "grade")
            )
            if exact_standard_duplicate or semantic_standard_duplicate:
                continue
        semantic_key = _external_semantic_key(path)
        compatible = next(
            (
                item
                for item in merged
                if item["semantic_key"] == semantic_key
                and item["value_type"] == value_types[path]
                and _compatible_external_values(item["values"], values)
            ),
            None,
        )
        if compatible is not None:
            compatible["paths"].append(path)
            compatible["score"] = max(compatible["score"], score)
            compatible["values"] = tuple(
                left if left is not None else right
                for left, right in zip(compatible["values"], values, strict=True)
            )
            continue
        merged.append(
            {
                "score": score,
                "path": path,
                "paths": [path],
                "semantic_key": semantic_key,
                "value_type": value_types[path],
                "values": values,
            }
        )
    merged.sort(
        key=lambda item: (
            -item["score"],
            _humanize_field(item["path"]).casefold(),
            _external_path_priority(item["path"]),
            item["path"],
        )
    )
    retained = merged[:MAX_EXTERNAL_VARIABLES]
    specs = [
        {
            "id": _external_variable_id(item["path"]),
            "source_field": item["path"],
            "source_fields": tuple(item["paths"]),
            "label": _humanize_field(item["path"]),
            "value_type": item["value_type"],
            "category": "dataset_specific",
            "source": "Curated source clinical metadata",
            "description": (
                "Dataset-specific field retained from the source clinical table."
            ),
            "unit": (
                _external_numeric_unit(item["path"])
                if item["value_type"] == "numeric"
                else None
            ),
            "level_labels": _external_level_labels(item["path"]),
            "expression_derived": any(
                token in _normalized_field_name(item["semantic_key"])
                for token in ("pam50", "rna_subtype", "mrna", "expression_subtype")
            ),
            "analysis_note": (
                "This source field may reflect treatment or post-baseline response; "
                "interpret expression differences in that temporal context."
                if any(
                    token in item["semantic_key"].casefold()
                    for token in (
                        "cycle",
                        "dosage",
                        "dose",
                        "duration",
                        "pcr",
                        "recist",
                        "response",
                        "therapy",
                        "treatment",
                    )
                )
                else None
            ),
        }
        for item in retained
    ]
    duplicate_labels = Counter(spec["label"].casefold() for spec in specs)
    for spec in specs:
        if duplicate_labels[spec["label"].casefold()] > 1:
            spec["label"] = f"{spec['label']} · {spec['source_field']}"
    return tuple(specs)


def _append_analysis_note(base: str | None, addition: str | None) -> str | None:
    if not addition:
        return base
    return f"{base} {addition}" if base else addition


def _spec_interpretation_policy(
    spec: dict[str, Any],
) -> tuple[str | None, bool, str | None]:
    analysis_note = spec.get("analysis_note")
    timing = spec.get("timing")
    survival_eligible = bool(spec.get("survival_eligible", True))
    survival_unavailable_reason: str | None = spec.get(
        "survival_unavailable_reason"
    )
    if spec.get("user_custom"):
        if timing == "outcome":
            survival_eligible = False
            survival_unavailable_reason = (
                "Outcome-defined variables cannot be used as baseline survival "
                "filters, groupings or Cox covariates."
            )
            analysis_note = _append_analysis_note(
                analysis_note,
                "This variable is outcome-defined. Expression comparisons and "
                "GSEA are outcome-conditioned and must not be interpreted as "
                "baseline prognosis.",
            )
        elif timing in {"on_treatment", "post_treatment", "post_baseline"}:
            survival_eligible = False
            survival_unavailable_reason = (
                "Post-baseline variables cannot be used as baseline survival "
                "filters, groupings or Cox covariates."
            )
            analysis_note = _append_analysis_note(
                analysis_note,
                "This variable was measured after baseline or treatment. "
                "Expression comparisons and GSEA describe that temporal context.",
            )
        elif timing == "unknown":
            survival_eligible = False
            survival_unavailable_reason = (
                "Variable timing is unknown; declare baseline or pre-treatment "
                "timing before using it in survival analysis."
            )
            analysis_note = _append_analysis_note(
                analysis_note,
                "Measurement timing is unknown and must be resolved before "
                "survival modeling.",
            )
    if spec.get("expression_derived"):
        analysis_note = _append_analysis_note(
            analysis_note,
            "This annotation was derived from molecular expression data. "
            "Expression comparisons and GSEA may be circular, especially when "
            "tested genes contribute to its classifier; survival analyses "
            "treat it as an expression-derived stratum.",
        )
    return analysis_note, survival_eligible, survival_unavailable_reason


def _spec_provenance(spec: dict[str, Any]) -> dict[str, Any]:
    """Describe who supplied a variable and how safely it travels across studies."""

    category = str(spec.get("category") or "clinical")
    expression_derived = bool(spec.get("expression_derived"))
    user_custom = bool(spec.get("user_custom"))
    source = str(spec.get("source") or "Dataset metadata")

    if user_custom:
        origin = "user_declared"
        label = "User-declared variable"
        method_summary = (
            "Declared in the uploaded clinical file; TRACE does not derive or "
            "independently verify this variable."
        )
        comparability = (
            "Its definition is specific to this upload unless the same source "
            "coding and ascertainment are documented elsewhere."
        )
    elif category == "tumor_specific":
        origin = (
            "source_reported_molecular"
            if expression_derived
            else "source_reported_annotation"
        )
        label = (
            "Published molecular call"
            if expression_derived
            else "Published cohort annotation"
        )
        method_summary = (
            "Author-provided expression-derived classification; TRACE reads the "
            "reported call and does not rerun the classifier."
            if expression_derived
            else "Author-provided cohort annotation; TRACE reads the reported "
            "value and does not reconstruct it."
        )
        comparability = (
            "Method-dependent. Confirm classifier, normalization and assay platform "
            "before comparing this label across datasets."
            if expression_derived
            else "Cohort-specific. Confirm the source definition and ascertainment "
            "before comparing this label across datasets."
        )
    elif category == "dataset_specific":
        origin = (
            "source_reported_molecular"
            if expression_derived
            else "source_reported_metadata"
        )
        label = (
            "Source-reported molecular call"
            if expression_derived
            else "Source-reported metadata"
        )
        method_summary = (
            "The source dataset reports this expression-derived classification; "
            "TRACE does not rerun or validate the classifier."
            if expression_derived
            else "The source dataset reports this field; TRACE preserves its "
            "declared values without deriving a new classification."
        )
        comparability = (
            "Method-dependent. Confirm classifier, normalization and assay platform "
            "before comparing this label across datasets."
            if expression_derived
            else "Do not assume an identically named field has the same definition "
            "or timing in another dataset."
        )
    else:
        origin = "harmonized_clinical"
        label = "Harmonized clinical field"
        method_summary = (
            "Clinical or pathological metadata harmonized by the source or TRACE "
            "import contract; TRACE does not infer a molecular subtype from it."
        )
        comparability = (
            "Harmonized names improve alignment but do not guarantee identical "
            "ascertainment across studies."
        )

    return {
        "origin": origin,
        "label": label,
        "reported_by": source,
        "method_summary": spec.get("method_summary") or method_summary,
        "reference": spec.get("reference"),
        "expression_derived": expression_derived,
        "recomputed_by_trace": False,
        "comparability": spec.get("comparability") or comparability,
    }


def _categorical_catalog_entry(
    spec: dict[str, Any],
    values: list[str | None],
    total_patients: int,
) -> dict[str, Any]:
    counts = Counter(value for value in values if value is not None)
    non_missing = sum(counts.values())
    unavailable_reason: str | None = None
    if non_missing < MIN_GROUPING_PATIENTS:
        unavailable_reason = (
            f"Fewer than {MIN_GROUPING_PATIENTS} patients have this annotation."
        )
    elif len(counts) < 2:
        unavailable_reason = "Fewer than two observed levels are available."
    elif len(counts) > MAX_CATEGORICAL_LEVELS:
        unavailable_reason = (
            f"More than {MAX_CATEGORICAL_LEVELS} observed levels are not supported."
        )
    elif sum(count >= MIN_LEVEL_PATIENTS for count in counts.values()) < 2:
        unavailable_reason = (
            f"Fewer than two observed levels contain at least {MIN_LEVEL_PATIENTS} patients."
        )
    grade_tokens = {
        token
        for token in _normalized_field_name(
            " ".join(
                str(spec.get(key) or "")
                for key in ("id", "source_field", "label")
            )
        ).split("_")
        if token
    }
    if "grade" in grade_tokens and counts:
        t_like = sum(
            count
            for value, count in counts.items()
            if re.fullmatch(r"(?i)p?T(?:is|x|[0-4][a-d]?)", value.strip())
        )
        if t_like / non_missing >= 0.8:
            unavailable_reason = (
                "The source grade field resembles tumor T categories and failed semantic QC."
            )
    level_labels = {
        _missing_key(str(value)): label
        for value, label in (spec.get("level_labels") or {}).items()
    }
    levels = [
        {
            "value": value,
            "label": str(level_labels.get(_missing_key(value)) or value),
            "count": count,
            "analysis_eligible": count >= MIN_LEVEL_PATIENTS,
            "unavailable_reason": (
                None
                if count >= MIN_LEVEL_PATIENTS
                else f"Fewer than {MIN_LEVEL_PATIENTS} patients have this level."
            ),
        }
        for value, count in sorted(
            counts.items(), key=lambda item: (-item[1], item[0].casefold())
        )
    ]
    expression_derived = bool(spec.get("expression_derived"))
    (
        analysis_note,
        survival_eligible,
        survival_unavailable_reason,
    ) = _spec_interpretation_policy(spec)
    coverage = non_missing / total_patients if total_patients else 0.0
    if 0 < coverage < LOW_COVERAGE_THRESHOLD:
        coverage_note = (
            f"Coverage is limited to {coverage * 100:.1f}% of patients before "
            "eligibility filters; inspect missingness before interpretation."
        )
        analysis_note = (
            f"{analysis_note} {coverage_note}"
            if analysis_note
            else coverage_note
        )
    return {
        "id": spec["id"],
        "label": spec["label"],
        "value_type": "categorical",
        "category": spec.get("category") or "clinical",
        "source": spec.get("source") or "TCGA marker-paper annotation",
        "source_field": spec.get("source_field") or spec["id"],
        "declared_source_column": spec.get("declared_source_column"),
        "description": spec.get("description"),
        "catalog_version": CLINICAL_GROUPING_CATALOG_VERSION,
        "analysis_eligible": unavailable_reason is None,
        "unavailable_reason": unavailable_reason,
        "patient_count": total_patients,
        "non_missing_count": non_missing,
        "missing_count": max(0, total_patients - non_missing),
        "coverage": round(coverage, 6),
        "levels": levels,
        "numeric_summary": None,
        "expression_derived": expression_derived,
        "timing": spec.get("timing"),
        "survival_eligible": survival_eligible,
        "survival_unavailable_reason": survival_unavailable_reason,
        "analysis_note": analysis_note,
        "provenance": _spec_provenance(spec),
    }


def _numeric_catalog_entry(
    spec: dict[str, Any],
    values: list[float | None],
    total_patients: int,
) -> dict[str, Any]:
    finite = [value for value in values if value is not None and math.isfinite(value)]
    unavailable_reason = None
    if len(finite) < MIN_GROUPING_PATIENTS:
        unavailable_reason = (
            f"Fewer than {MIN_GROUPING_PATIENTS} patients have a finite value."
        )
    elif min(finite) == max(finite):
        unavailable_reason = "The observed values do not vary."
    coverage = len(finite) / total_patients if total_patients else 0.0
    (
        analysis_note,
        survival_eligible,
        survival_unavailable_reason,
    ) = _spec_interpretation_policy(spec)
    if 0 < coverage < LOW_COVERAGE_THRESHOLD:
        coverage_note = (
            f"Coverage is limited to {coverage * 100:.1f}% of patients before "
            "eligibility filters; inspect missingness before interpretation."
        )
        analysis_note = (
            f"{analysis_note} {coverage_note}"
            if analysis_note
            else coverage_note
        )
    return {
        "id": spec["id"],
        "label": spec["label"],
        "value_type": "numeric",
        "category": spec.get("category") or "standardized",
        "source": spec.get("source") or "Standardized dataset metadata",
        "source_field": spec.get("source_field") or spec["id"],
        "declared_source_column": spec.get("declared_source_column"),
        "description": spec.get("description"),
        "catalog_version": CLINICAL_GROUPING_CATALOG_VERSION,
        "analysis_eligible": unavailable_reason is None,
        "unavailable_reason": unavailable_reason,
        "patient_count": total_patients,
        "non_missing_count": len(finite),
        "missing_count": max(0, total_patients - len(finite)),
        "coverage": round(coverage, 6),
        "levels": [],
        "numeric_summary": {
            "min": min(finite) if finite else None,
            "max": max(finite) if finite else None,
            "median": statistics.median(finite) if finite else None,
            "unit": spec.get("unit"),
        },
        "expression_derived": bool(spec.get("expression_derived")),
        "timing": spec.get("timing"),
        "survival_eligible": survival_eligible,
        "survival_unavailable_reason": survival_unavailable_reason,
        "analysis_note": analysis_note,
        "provenance": _spec_provenance(spec),
    }


def clinical_grouping_context(
    samples: Iterable[Any],
    *,
    cohort: str,
    tcga_data_dir: Path | None = None,
    tcga_cdr_path: Path | None = None,
    repository: bool = False,
) -> tuple[list[dict[str, Any]], dict[str, dict[str, str | float | None]]]:
    sample_list = list(samples)
    representatives = _representative_samples(sample_list)
    total_patients = len(representatives)
    repository_metadata = (
        {
            str(sample.barcode): _flatten_metadata(
                getattr(sample, "raw_metadata", None) or {}
            )
            for sample in sample_list
        }
        if repository
        else {}
    )

    tcga_by_barcode: dict[str, dict[str, Any]] = {}
    tcga_by_patient: dict[str, dict[str, Any]] = {}
    cdr_by_patient: dict[str, dict[str, Any]] = {}
    if not repository:
        if tcga_data_dir is not None:
            tcga_by_barcode, tcga_by_patient = _tcga_row_indexes(
                _tcga_rows(tcga_data_dir / cohort / "col_data.tsv")
            )
        cdr_by_patient = _cdr_index(_cdr_rows(tcga_cdr_path), cohort)

    repository_standard_values: dict[str, dict[str, Any]] = {}
    repository_standard_source_fields: dict[str, str] = {}
    claimed_repository_paths: set[str] = set()
    if repository:
        (
            repository_standard_values,
            repository_standard_source_fields,
            claimed_repository_paths,
        ) = _repository_standard_context(sample_list, repository_metadata)
    specs: list[dict[str, Any]] = [
        {
            **spec,
            "source_field": repository_standard_source_fields.get(
                str(spec["id"]), str(spec["id"])
            ),
        }
        for spec in STANDARD_FIELDS
    ]
    if repository:
        if any(getattr(sample, "study_arm", None) for sample in sample_list):
            specs.append(
                {
                    "id": "study_arm",
                    "label": "Treatment arm",
                    "value_type": "categorical",
                    "category": "dataset_specific",
                    "source": "Curated randomized treatment assignment",
                    "source_field": "study_arm",
                    "declared_source_column": "patient.ICI_RX + sample.NON_ICI_RX",
                    "sample_attribute": "study_arm",
                    "description": (
                        "Prespecified treatment arm reconstructed from the source "
                        "trial assignment fields, never from response or outcome."
                    ),
                    "timing": "baseline_randomized_assignment",
                    "analysis_note": (
                        "Select exactly one treatment arm for IMmotion150 PFS. "
                        "TRACE does not pool PFS across randomized arms."
                    ),
                }
            )
        specs.extend(_user_custom_specs(sample_list))
        specs.extend(
            _external_specs(
                sample_list,
                repository_metadata,
                claimed_paths=claimed_repository_paths,
                standard_values=repository_standard_values,
            )
        )
    else:
        specs.extend(
            {
                **spec,
                "value_type": "categorical",
                "source": "GDC harmonized clinical metadata",
                "source_field": spec["id"],
                "description": spec.get("description") or (
                    "Clinical field from the primary-diagnosis row in the "
                    "harmonized TCGA sample metadata table."
                    if spec.get("require_primary_diagnosis")
                    else "Clinical field from harmonized TCGA sample metadata."
                ),
            }
            for spec in TCGA_GDC_FIELDS
        )
        specs.extend(
            {
                **spec,
                "value_type": spec.get("value_type") or "categorical",
                "category": "clinical",
                "source": "GDC harmonized clinical metadata",
                "source_field": spec["id"],
                "description": "Cohort-relevant field from harmonized TCGA clinical metadata.",
            }
            for spec in TCGA_COHORT_CLINICAL_FIELDS.get(cohort, ())
        )
        specs.extend(
            {
                **spec,
                "value_type": "categorical",
                "category": "clinical",
                "source": "TCGA Clinical Data Resource",
                "description": "Curated field from the TCGA Pan-Cancer Clinical Data Resource.",
            }
            for spec in TCGA_CDR_FIELDS
        )
        specs.extend(
            {
                **spec,
                "value_type": "categorical",
                "category": "tumor_specific",
                "source": "TCGA marker-paper annotation",
                "description": "Cohort-specific annotation published with the TCGA marker paper.",
            }
            for spec in TCGA_ANNOTATION_FIELDS.get(cohort, ())
        )

    values_by_variable: dict[str, dict[str, str | float | None]] = {}
    catalog: list[dict[str, Any]] = []
    for spec in specs:
        variable_id = str(spec["id"])
        source_field = str(spec.get("source_field") or variable_id)
        values: dict[str, str | float | None] = {}
        for sample in sample_list:
            raw: Any = None
            if spec.get("sample_attribute"):
                raw = getattr(sample, str(spec["sample_attribute"]), None)
            elif spec.get("user_custom"):
                raw = _user_custom_value(sample, variable_id)
            elif variable_id in {field["id"] for field in STANDARD_FIELDS}:
                raw = repository_standard_values.get(variable_id, {}).get(
                    str(sample.barcode),
                    getattr(sample, variable_id, None),
                )
            elif repository:
                metadata = repository_metadata.get(str(sample.barcode), {})
                for candidate_field in spec.get("source_fields") or (source_field,):
                    candidate = metadata.get(candidate_field)
                    if _clean_value(
                        candidate,
                        source_field=candidate_field,
                    ) is not None:
                        raw = candidate
                        break
            elif variable_id.startswith("cdr."):
                raw = cdr_by_patient.get(str(sample.patient_id), {}).get(source_field)
            else:
                row = tcga_by_barcode.get(str(sample.barcode)) or tcga_by_patient.get(
                    str(sample.patient_id), {}
                )
                diagnosis_primary = str(
                    row.get("diagnosis_is_primary_disease") or ""
                ).strip().casefold() in {"1", "true", "yes"}
                raw = (
                    None
                    if spec.get("require_primary_diagnosis")
                    and not diagnosis_primary
                    else row.get(source_field)
                )
                if raw is None and (
                    not spec.get("require_primary_diagnosis")
                    or diagnosis_primary
                ):
                    raw = (getattr(sample, "raw_metadata", None) or {}).get(
                        source_field
                    )
            if spec.get("value_type") == "numeric":
                value = _clean_numeric(raw, source_field=source_field)
            else:
                value = _clean_category(raw, source_field=source_field)
            values[str(sample.barcode)] = value
        if spec.get("value_type") != "numeric":
            values = _canonical_category_mapping(
                values,
                source_field=source_field,
                concept=(
                    variable_id
                    if variable_id in {field["id"] for field in STANDARD_FIELDS}
                    else None
                ),
            )
        values_by_variable[variable_id] = values
        representative_values = [
            values.get(str(sample.barcode)) for sample in representatives
        ]
        if spec.get("value_type") == "numeric":
            catalog.append(
                _numeric_catalog_entry(
                    spec,
                    [
                        float(value) if isinstance(value, (int, float)) else None
                        for value in representative_values
                    ],
                    total_patients,
                )
            )
        else:
            catalog.append(
                _categorical_catalog_entry(
                    spec,
                    [value if isinstance(value, str) else None for value in representative_values],
                    total_patients,
                )
            )

    category_order = {
        "standardized": 0,
        "clinical": 1,
        "tumor_specific": 2,
        "dataset_specific": 3,
    }
    catalog.sort(
        key=lambda item: (
            category_order.get(item["category"], 9),
            not item["analysis_eligible"],
            item["label"].casefold(),
            item["id"],
        )
    )
    return catalog, values_by_variable


def clinical_variable_analysis_eligibility(
    definition: dict[str, Any],
    *,
    analysis_context: str | None = None,
) -> dict[str, Any]:
    context = str(analysis_context or "grouping").strip().casefold()
    if context not in {"grouping", "expression", "gsea", "survival"}:
        raise ValueError(f"Unsupported clinical-variable context: {context!r}.")
    if not definition.get("analysis_eligible", False):
        return {
            "eligible": False,
            "reason": definition.get("unavailable_reason")
            or "The variable is not eligible for analysis.",
            "warnings": [],
        }
    if context == "survival" and not definition.get(
        "survival_eligible", True
    ):
        return {
            "eligible": False,
            "reason": definition.get("survival_unavailable_reason")
            or "The variable timing is not eligible for survival analysis.",
            "warnings": [],
        }
    warnings = (
        [str(definition["analysis_note"])]
        if definition.get("analysis_note")
        else []
    )
    return {"eligible": True, "reason": None, "warnings": warnings}


def filter_samples_by_clinical_variable(
    samples: Iterable[Any],
    definition: dict[str, Any],
    values: dict[str, str | float | None],
    *,
    categorical_levels: Iterable[str] | None = None,
    numeric_min: float | None = None,
    numeric_max: float | None = None,
    analysis_context: str | None = None,
) -> tuple[list[Any], dict[str, Any], list[str]]:
    """Apply one declared categorical-level or numeric-range filter.

    This helper intentionally does not construct adjustment terms. It provides
    a reusable backend contract for a future request schema while keeping
    custom upload metadata out of automatic Cox adjustment.
    """

    policy = clinical_variable_analysis_eligibility(
        definition,
        analysis_context=analysis_context,
    )
    if not policy["eligible"]:
        raise ValueError(
            f"Clinical variable {definition.get('label')!r} cannot be used: "
            f"{policy['reason']}"
        )
    sample_list = list(samples)
    selected: list[Any] = []
    value_type = definition.get("value_type")
    filter_definition: dict[str, Any]
    if value_type == "categorical":
        requested = [
            str(value).strip()
            for value in (categorical_levels or [])
            if str(value).strip()
        ]
        requested = list(dict.fromkeys(requested))
        if not requested:
            raise ValueError(
                "Categorical clinical filtering requires at least one level."
            )
        if numeric_min is not None or numeric_max is not None:
            raise ValueError(
                "Numeric ranges cannot be applied to categorical variables."
            )
        observed = {
            str(level["value"]).casefold(): str(level["value"])
            for level in definition.get("levels") or []
        }
        unknown = sorted(
            value for value in requested if value.casefold() not in observed
        )
        if unknown:
            raise ValueError(
                "Clinical filter contains unobserved levels: "
                + ", ".join(unknown)
                + "."
            )
        allowed = {value.casefold() for value in requested}
        selected = [
            sample
            for sample in sample_list
            if isinstance(values.get(str(sample.barcode)), str)
            and str(values[str(sample.barcode)]).casefold() in allowed
        ]
        filter_definition = {"levels": requested}
    elif value_type == "numeric":
        if categorical_levels:
            raise ValueError(
                "Categorical levels cannot be applied to numeric variables."
            )
        lower = float(numeric_min) if numeric_min is not None else None
        upper = float(numeric_max) if numeric_max is not None else None
        if lower is not None and not math.isfinite(lower):
            raise ValueError("Numeric clinical filter minimum must be finite.")
        if upper is not None and not math.isfinite(upper):
            raise ValueError("Numeric clinical filter maximum must be finite.")
        if lower is None and upper is None:
            raise ValueError(
                "Numeric clinical filtering requires a minimum or maximum."
            )
        if lower is not None and upper is not None and lower > upper:
            raise ValueError(
                "Numeric clinical filter minimum cannot exceed its maximum."
            )

        def in_range(value: Any) -> bool:
            if not isinstance(value, (int, float)):
                return False
            numeric = float(value)
            return (
                math.isfinite(numeric)
                and (lower is None or numeric >= lower)
                and (upper is None or numeric <= upper)
            )

        selected = [
            sample
            for sample in sample_list
            if in_range(values.get(str(sample.barcode)))
        ]
        filter_definition = {"min": lower, "max": upper}
    else:
        raise ValueError("Unsupported clinical variable type for filtering.")

    def patient_keys(rows: Iterable[Any]) -> set[str]:
        return {
            str(
                getattr(row, "patient_id", None)
                or getattr(row, "barcode", "")
            )
            for row in rows
        }

    considered_patients = patient_keys(sample_list)
    retained_patients = patient_keys(selected)
    audit = {
        "variable_id": definition.get("id"),
        "variable_label": definition.get("label"),
        "value_type": value_type,
        "category": definition.get("category"),
        "source": definition.get("source"),
        "source_field": definition.get("source_field"),
        "expression_derived": bool(definition.get("expression_derived")),
        "provenance": definition.get("provenance"),
        "timing": definition.get("timing"),
        "analysis_context": analysis_context or "grouping",
        "filter": filter_definition,
        "patients_considered": len(considered_patients),
        "patients_retained": len(retained_patients),
        "patients_excluded": len(considered_patients - retained_patients),
        "samples_considered": len(sample_list),
        "samples_retained": len(selected),
        "samples_excluded": len(sample_list) - len(selected),
        "automatic_cox_adjustment": False,
    }
    return selected, audit, list(policy["warnings"])


def apply_custom_clinical_filters(
    samples: Iterable[Any],
    custom_filters: Iterable[Any],
    *,
    analysis_context: str,
    cohort: str | None = None,
    tcga_data_dir: Path | None = None,
    tcga_cdr_path: Path | None = None,
    repository: bool | None = None,
) -> tuple[list[Any], list[dict[str, Any]], list[str]]:
    """Apply curated clinical filters with AND semantics.

    Only variables exposed by the dataset-aware clinical catalog can be
    selected. Standardized fields keep their dedicated filter controls;
    cohort-specific clinical, tumor and dataset annotations can be combined
    here. Survival calls enforce timing eligibility, while Expression and GSEA
    retain interpretive warnings. No Cox adjustment terms are made.
    """

    selected = list(samples)
    requested_filters = list(custom_filters)
    if not requested_filters:
        return selected, [], []
    observed_cohorts = {
        str(getattr(sample, "cohort", "") or "").strip()
        for sample in selected
        if str(getattr(sample, "cohort", "") or "").strip()
    }
    resolved_cohort = str(cohort or "").strip()
    if not resolved_cohort:
        if len(observed_cohorts) == 1:
            resolved_cohort = next(iter(observed_cohorts))
        elif repository is True:
            resolved_cohort = "PRIVATE"
        else:
            raise ValueError(
                "Clinical filtering requires exactly one dataset cohort."
            )
    if repository is None:
        repository = not resolved_cohort.startswith("TCGA-")
    catalog, values_by_variable = clinical_grouping_context(
        selected,
        cohort=resolved_cohort,
        tcga_data_dir=None if repository else tcga_data_dir,
        tcga_cdr_path=None if repository else tcga_cdr_path,
        repository=repository,
    )
    definitions = {
        str(item["id"]).casefold(): item
        for item in catalog
        if item.get("category") in {
            "clinical",
            "tumor_specific",
            "dataset_specific",
        }
    }
    audits: list[dict[str, Any]] = []
    warnings: list[str] = []
    for requested in requested_filters:
        if isinstance(requested, dict):
            variable_id = str(requested.get("variable_id") or "").strip()
            categorical_levels = requested.get("categorical_levels") or []
            numeric_min = requested.get("numeric_min")
            numeric_max = requested.get("numeric_max")
        else:
            variable_id = str(
                getattr(requested, "variable_id", "") or ""
            ).strip()
            categorical_levels = (
                getattr(requested, "categorical_levels", None) or []
            )
            numeric_min = getattr(requested, "numeric_min", None)
            numeric_max = getattr(requested, "numeric_max", None)
        definition = definitions.get(variable_id.casefold())
        if definition is None:
            raise ValueError(
                f"Clinical filter variable {variable_id!r} is not declared "
                "as an eligible cohort-specific filter for this dataset."
            )
        canonical_id = str(definition["id"])
        selected, audit, filter_warnings = (
            filter_samples_by_clinical_variable(
                selected,
                definition,
                values_by_variable[canonical_id],
                categorical_levels=categorical_levels,
                numeric_min=numeric_min,
                numeric_max=numeric_max,
                analysis_context=analysis_context,
            )
        )
        audits.append(audit)
        warnings.extend(filter_warnings)
    return selected, audits, list(dict.fromkeys(warnings))


def resolve_clinical_grouping_variable(
    samples: Iterable[Any],
    variable_id: str,
    *,
    cohort: str,
    tcga_data_dir: Path | None = None,
    tcga_cdr_path: Path | None = None,
    repository: bool = False,
    analysis_context: str | None = None,
) -> tuple[dict[str, Any], dict[str, str | float | None]]:
    catalog, values = clinical_grouping_context(
        samples,
        cohort=cohort,
        tcga_data_dir=tcga_data_dir,
        tcga_cdr_path=tcga_cdr_path,
        repository=repository,
    )
    definition = next(
        (item for item in catalog if item["id"] == variable_id), None
    )
    if definition is None:
        raise ValueError(
            f"Clinical grouping variable {variable_id!r} is not declared for this dataset."
        )
    if not definition["analysis_eligible"]:
        raise ValueError(
            f"Clinical grouping variable {definition['label']!r} is unavailable: "
            f"{definition['unavailable_reason']}"
        )
    policy = clinical_variable_analysis_eligibility(
        definition,
        analysis_context=analysis_context,
    )
    if not policy["eligible"]:
        raise ValueError(
            f"Clinical grouping variable {definition['label']!r} is unavailable "
            f"for {analysis_context or 'grouping'}: {policy['reason']}"
        )
    return definition, values[variable_id]
