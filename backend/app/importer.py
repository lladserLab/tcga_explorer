import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.models import ClinicalEndpoint, Cohort, DataSource, GeneIndex, Sample

MISSING = {"", "NA", "N/A", "NULL", "null", "None", "not reported", "Not Reported", "--"}
TCGA_CDR_SOURCE_ID = "tcga_cdr"
TCGA_RNA_SOURCE_ID = "tcga_rna"
TCGA_CDR_URL = "https://gdc.cancer.gov/about-data/publications/PanCan-Clinical-2018"
ENDPOINT_COLUMNS = {
    "OS": ("OS", "OS.time"),
    "PFI": ("PFI", "PFI.time"),
    "DFI": ("DFI", "DFI.time"),
    "DSS": ("DSS", "DSS.time"),
}


def clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return None if text in MISSING else text


def parse_float(value: Any) -> float | None:
    text = clean(value)
    if text is None:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_int(value: Any) -> int | None:
    parsed = parse_float(value)
    return None if parsed is None else int(parsed)


def normalize_stage(row: dict[str, str]) -> str | None:
    stage = clean(row.get("ajcc_pathologic_stage")) or clean(row.get("paper_pathologic_stage")) or clean(row.get("figo_stage"))
    if stage is None:
        return None
    return stage.replace("_", " ")


def derive_os(row: dict[str, str]) -> tuple[float | None, int | None]:
    vital = (clean(row.get("vital_status")) or "").lower()
    days_to_death = parse_float(row.get("days_to_death"))
    days_to_follow_up = parse_float(row.get("days_to_last_follow_up"))
    days_to_last_known = parse_float(row.get("days_to_last_known_disease_status"))

    if vital == "dead":
        event = 1
        time_days = days_to_death or days_to_follow_up or days_to_last_known
    elif vital == "alive":
        event = 0
        time_days = days_to_follow_up or days_to_last_known
    else:
        event = None
        time_days = days_to_death or days_to_follow_up or days_to_last_known

    if time_days is None or time_days <= 0:
        return None, None
    return time_days, event


def read_summary(summary_path: Path) -> list[dict[str, str]]:
    with summary_path.open(newline="", encoding="utf-8", errors="replace") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def import_cohorts_and_samples(db: Session, tcga_data_dir: Path, force: bool = False) -> dict[str, int]:
    if force:
        db.execute(delete(Sample))
        db.execute(delete(GeneIndex))
        db.execute(delete(Cohort))
        db.commit()

    summary_path = tcga_data_dir / "summary_table.tsv"
    if not summary_path.exists():
        raise FileNotFoundError(f"summary_table.tsv not found under {tcga_data_dir}")

    existing = db.scalar(select(func.count()).select_from(Cohort))
    source = db.get(DataSource, TCGA_RNA_SOURCE_ID)
    summary_modified = modified_datetime(summary_path)
    if existing and not force:
        if source and source.source_file_modified_at == summary_modified:
            register_tcga_rna_source(db, tcga_data_dir)
            return {"cohorts": int(existing), "samples": int(db.scalar(select(func.count()).select_from(Sample)) or 0)}
        db.execute(delete(Sample))
        db.execute(delete(GeneIndex))
        db.execute(delete(Cohort))
        db.commit()

    cohorts = read_summary(summary_path)
    sample_total = 0
    for item in cohorts:
        cohort_id = item["cohort"]
        cohort_dir = tcga_data_dir / cohort_id
        if not cohort_dir.exists():
            continue

        cohort = Cohort(
            id=cohort_id,
            disease_type=clean(item.get("disease_type")),
            primary_site=clean(item.get("primary_site")),
            n_samples_paired=parse_int(item.get("n_samples_paired")),
            n_patients_paired=parse_int(item.get("n_patients_paired")),
            n_primary_tumor=parse_int(item.get("n_primary_tumor")),
            n_solid_normal=parse_int(item.get("n_solid_normal")),
            n_other_samples=parse_int(item.get("n_other_samples")),
            n_genes=parse_int(item.get("n_genes")),
            design_formula=clean(item.get("design_formula")),
            status=clean(item.get("status")),
            data_path=str(cohort_dir),
        )
        db.add(cohort)
        db.flush()
        sample_total += import_samples_for_cohort(db, cohort_id, cohort_dir)

    db.commit()
    register_tcga_rna_source(db, tcga_data_dir)
    return {"cohorts": len(cohorts), "samples": sample_total}


def register_tcga_rna_source(db: Session, tcga_data_dir: Path) -> DataSource:
    source = get_or_create_data_source(
        db,
        source_id=TCGA_RNA_SOURCE_ID,
        label="TCGA RNA-seq cohort data",
        kind="rna_expression",
    )
    summary_path = tcga_data_dir / "summary_table.tsv"
    source.status = "ready" if summary_path.exists() else "missing"
    source.source_url = "https://portal.gdc.cancer.gov/"
    source.source_path = str(tcga_data_dir)
    source.source_file_modified_at = modified_datetime(summary_path)
    previous_metadata = source.metadata_json or {}
    source.metadata_json = {
        "description": "RNA-seq expression and sample metadata imported from TCGA cancer cohort files.",
        "summary_file": str(summary_path),
        **{
            key: previous_metadata[key]
            for key in ("manifest_hash", "data_release", "data_through_date", "manifest_path")
            if key in previous_metadata
        },
    }
    source.imported_at = datetime.utcnow()
    db.commit()
    return source


def import_tcga_cdr(db: Session, cdr_path: Path, force: bool = False) -> dict[str, int | str]:
    source = get_or_create_data_source(
        db,
        source_id=TCGA_CDR_SOURCE_ID,
        label="TCGA Clinical Data Resource",
        kind="clinical_endpoint",
    )
    source.source_url = TCGA_CDR_URL
    source.source_path = str(cdr_path)
    source.metadata_json = {
        "description": "TCGA-CDR standardized clinical outcome endpoint resource.",
        "citation": "Liu et al., Cell 2018, doi:10.1016/j.cell.2018.02.052",
        "supported_endpoints": sorted(ENDPOINT_COLUMNS),
    }

    if not cdr_path.exists():
        existing = int(
            db.scalar(
                select(func.count())
                .select_from(ClinicalEndpoint)
                .where(ClinicalEndpoint.source_id == TCGA_CDR_SOURCE_ID)
            )
            or 0
        )
        source.status = "ready_cached" if existing else "missing"
        source.source_file_modified_at = None
        source.imported_at = datetime.utcnow()
        db.commit()
        return {"status": source.status, "endpoints": existing}

    modified_at = modified_datetime(cdr_path)
    existing = int(
        db.scalar(
            select(func.count())
            .select_from(ClinicalEndpoint)
            .where(ClinicalEndpoint.source_id == TCGA_CDR_SOURCE_ID)
        )
        or 0
    )
    if existing and not force and source.source_file_modified_at == modified_at and source.status in {"ready", "ready_cached"}:
        return {"status": source.status, "endpoints": existing}

    rows = read_tcga_cdr_rows(cdr_path)
    endpoint_rows = tcga_cdr_endpoint_rows(rows)
    db.execute(delete(ClinicalEndpoint).where(ClinicalEndpoint.source_id == TCGA_CDR_SOURCE_ID))
    db.flush()
    if endpoint_rows:
        db.add_all(endpoint_rows)
    source.status = "ready" if endpoint_rows else "empty"
    source.source_file_modified_at = modified_at
    source.imported_at = datetime.utcnow()
    source.metadata_json = {
        **(source.metadata_json or {}),
        "row_count": len(rows),
        "endpoint_record_count": len(endpoint_rows),
    }
    db.commit()
    return {"status": source.status, "endpoints": len(endpoint_rows)}


def get_or_create_data_source(db: Session, source_id: str, label: str, kind: str) -> DataSource:
    source = db.get(DataSource, source_id)
    if source is None:
        source = DataSource(id=source_id, label=label, kind=kind, status="missing")
        db.add(source)
        db.flush()
    else:
        source.label = label
        source.kind = kind
    return source


def read_tcga_cdr_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency is installed in Docker
            raise RuntimeError("openpyxl is required to read TCGA-CDR XLSX files.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        return [
            {headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]}
            for row in rows
            if any(value is not None and str(value).strip() for value in row)
        ]

    delimiter = "\t" if suffix in {".tsv", ".txt"} else ","
    with path.open(newline="", encoding="utf-8", errors="replace") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def tcga_cdr_endpoint_rows(rows: list[dict[str, Any]]) -> list[ClinicalEndpoint]:
    endpoints: list[ClinicalEndpoint] = []
    seen: set[tuple[str, str]] = set()
    for row in rows:
        normalized = {str(key).strip(): value for key, value in row.items()}
        patient_id = clean(
            normalized.get("bcr_patient_barcode")
            or normalized.get("patient_id")
            or normalized.get("submitter_id")
        )
        cohort = normalize_cdr_cohort(normalized.get("type") or normalized.get("cohort") or normalized.get("project_id"))
        if not patient_id or not cohort:
            continue
        for endpoint, (event_col, time_col) in ENDPOINT_COLUMNS.items():
            event = parse_int(normalized.get(event_col))
            time_days = parse_float(normalized.get(time_col))
            if event not in {0, 1} or time_days is None or time_days <= 0:
                continue
            key = (patient_id, endpoint)
            if key in seen:
                continue
            seen.add(key)
            endpoints.append(
                ClinicalEndpoint(
                    source_id=TCGA_CDR_SOURCE_ID,
                    cohort=cohort,
                    patient_id=patient_id,
                    endpoint=endpoint,
                    time_days=time_days,
                    event=event,
                    raw_metadata={
                        "type": clean(normalized.get("type")),
                        "source_event_column": event_col,
                        "source_time_column": time_col,
                    },
                )
            )
    return endpoints


def normalize_cdr_cohort(value: Any) -> str | None:
    text = clean(value)
    if text is None:
        return None
    text = text.upper()
    if text.startswith("TCGA-"):
        return text
    if len(text) <= 5:
        return f"TCGA-{text}"
    return text


def modified_datetime(path: Path) -> datetime | None:
    if not path.exists():
        return None
    return datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).replace(tzinfo=None)


def import_samples_for_cohort(db: Session, cohort_id: str, cohort_dir: Path) -> int:
    col_data_path = cohort_dir / "col_data.tsv"
    if not col_data_path.exists():
        return 0

    count = 0
    batch: list[Sample] = []
    with col_data_path.open(newline="", encoding="utf-8", errors="replace") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        for row in reader:
            barcode = clean(row.get("")) or clean(row.get("barcode"))
            patient_id = clean(row.get("patient_id")) or clean(row.get("patient")) or (barcode[:12] if barcode else None)
            if barcode is None or patient_id is None:
                continue
            os_time_days, os_event = derive_os(row)
            age = parse_float(row.get("age_at_index"))
            sample = Sample(
                cohort=cohort_id,
                patient_id=patient_id,
                barcode=barcode,
                sample_type=clean(row.get("sample_type")),
                stage=normalize_stage(row),
                gender=clean(row.get("gender")) or clean(row.get("sex_at_birth")),
                race=clean(row.get("race")),
                age_at_index=age,
                vital_status=clean(row.get("vital_status")),
                os_time_days=os_time_days,
                os_event=os_event,
                raw_metadata={
                    "primary_diagnosis": clean(row.get("primary_diagnosis")),
                    "paper_BRCA_Subtype_PAM50": clean(row.get("paper_BRCA_Subtype_PAM50")),
                    "progression_or_recurrence": clean(row.get("progression_or_recurrence")),
                },
            )
            batch.append(sample)
            count += 1
            if len(batch) >= 1000:
                db.add_all(batch)
                db.flush()
                batch.clear()

    if batch:
        db.add_all(batch)
        db.flush()
    return count


def ensure_gene_index(db: Session, tcga_data_dir: Path, cohort_id: str) -> int:
    existing = db.scalar(select(func.count()).select_from(GeneIndex).where(GeneIndex.cohort == cohort_id))
    if existing:
        return int(existing)

    matrix_path = tcga_data_dir / cohort_id / "count_matrix.tsv"
    if not matrix_path.exists():
        raise FileNotFoundError(f"count_matrix.tsv not found for {cohort_id}")

    genes: list[GeneIndex] = []
    seen_symbols: set[str] = set()
    with matrix_path.open(newline="", encoding="utf-8", errors="replace") as handle:
        reader = csv.reader(handle, delimiter="\t")
        next(reader, None)
        for idx, row in enumerate(reader, start=1):
            if not row:
                continue
            symbol = clean(row[0])
            if symbol is None:
                continue
            normalized = symbol.upper()
            if normalized in seen_symbols:
                continue
            seen_symbols.add(normalized)
            genes.append(GeneIndex(cohort=cohort_id, gene_symbol=normalized, row_number=idx))
            if len(genes) >= 5000:
                db.add_all(genes)
                db.flush()
                genes.clear()

    if genes:
        db.add_all(genes)
        db.flush()
    db.commit()
    return int(db.scalar(select(func.count()).select_from(GeneIndex).where(GeneIndex.cohort == cohort_id)) or 0)
